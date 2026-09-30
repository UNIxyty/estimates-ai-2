"""Build realistic fixture files for the ingest tests (openpyxl / python-docx / reportlab / pypdf).

Nothing binary is committed: `build_all(dir)` writes everything into a temp dir once per test session.
Run directly (`python tests/fixtures/make_fixtures.py OUT_DIR`) to inspect the files by hand.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

BOLD = Font(bold=True)

# (text, unit, qty, norm_h, rate, material, mechanisms)
LV_SHEET1 = [
    ("section", "1.", "Kabeļi un caurules"),
    ("item", "Kabelis NYM-J 3x1,5 mm²", "m", 120, 0.08, 12.5, 0.65, 0),
    ("item", "Kabelis NYM-J 3x2,5 mm²", "m", 80, 0.09, 12.5, 0.95, 0),
    ("item", "Kabelis NYM-J 5x2,5 mm²", "m", 40, 0.12, 12.5, 1.60, 0),
    ("item", "Gofrētā caurule d20", "m", 100, 0.05, 12.5, 0.30, 0),
    ("subtotal", "Kopā 1. sadaļā"),
    ("section", "2.", "Elektroierīces"),
    ("item", "Kontaktligzda 2-vietīga zemapmetuma IP20", "gab.", 24, 0.35, 12.5, "6,80", 0),
    ("item", "Slēdzis 1-polīgs zemapmetuma", "gab.", 12, 0.30, 12.5, 4.20, 0),
    ("item", "Gaismeklis LED 36W IP44", "gab.", 16, 0.60, 12.5, 38.00, 0.5),
    ("item", "Sadales skapis 36 moduļi", "kompl.", 1, 5.5, 12.5, 145.00, 0),
    ("subtotal", "Kopā 2. sadaļā"),
]
LV_SHEET2 = [
    ("section", "1.", "Datu tīkls un signalizācija"),
    ("item", "Datu rozete RJ45 2-vietīga", "gab.", 10, 0.40, 14.0, 9.50, 0),
    ("item", "Kabelis UTP cat6 4x2x0.5", "m", 200, 0.04, 14.0, 0.45, 0),
    ("item", "Dūmu detektors", "gab.", 8, 0.50, 14.0, 22.00, 0),
    ("item", "Kabeļu trase 100x60", "m", 30, 0.20, 14.0, 7.80, 0.3),
    ("subtotal", "Kopā 1. sadaļā"),
    ("section", "2.", "Zemējums"),
    ("item", "Zemējuma elektrods 1,5 m", "gab.", 4, 1.20, 15.0, 18.00, 1.0),
    ("item", "Potenciālu izlīdzināšanas kopne", "gab.", 2, 0.80, 15.0, 12.00, 0),
    ("subtotal", "Kopā 2. sadaļā"),
]


def _lv_sheet(ws, rows, *, blank: bool, rate_label: float | None, formulas_in_blank: bool = True):
    ws["A1"] = "Lokālā tāme Nr. 1"
    ws["A1"].font = BOLD
    ws.merge_cells("A1:O1")
    ws["A2"] = "Objekts: Dzīvojamā māja, Rīga, Brīvības iela 1"
    if rate_label is not None:
        ws["A3"] = "Vidējā stundas likme, EUR/h"
        ws["D3"] = rate_label
    # two-row header: merged group "Vienības izmaksas" over UNLABELLED unit columns E..J
    for col, text in (("A", "Nr."), ("B", "Darba nosaukums"), ("C", "Mērv."), ("D", "Daudz.")):
        ws[f"{col}5"] = text
        ws[f"{col}5"].font = BOLD
        ws.merge_cells(f"{col}5:{col}6")
    ws["E5"] = "Vienības izmaksas"
    ws.merge_cells("E5:J5")
    ws["K5"] = "Kopā uz visu apjomu"
    ws.merge_cells("K5:O5")
    for col, text in (("K", "Darbietilpība (c/h)"), ("L", "Darba alga (EUR)"), ("M", "Materiāli (EUR)"),
                      ("N", "Mehānismi (EUR)"), ("O", "Summa (EUR)")):
        ws[f"{col}6"] = text
        ws[f"{col}6"].alignment = Alignment(wrap_text=True)
    for i in range(1, 16):
        ws.cell(row=7, column=i, value=i)
    r = 8
    sec_first = None
    subtotal_rows = []
    for spec in rows:
        kind = spec[0]
        if kind == "section":
            ws[f"A{r}"] = spec[1]
            ws[f"B{r}"] = spec[2]
            ws[f"A{r}"].font = BOLD
            ws[f"B{r}"].font = BOLD
            sec_first = r + 1
        elif kind == "item":
            _, text, unit, qty, norm, rate, mat, mech = spec
            ws[f"B{r}"] = text
            ws[f"C{r}"] = unit
            ws[f"D{r}"] = qty
            if not blank:
                ws[f"E{r}"] = norm
                ws[f"F{r}"] = rate
                ws[f"G{r}"] = f"=ROUND(E{r}*F{r},2)"
                ws[f"H{r}"] = mat
                ws[f"I{r}"] = mech
                ws[f"J{r}"] = f"=G{r}+H{r}+I{r}"
            if not blank or formulas_in_blank:
                ws[f"K{r}"] = f"=D{r}*E{r}"
                ws[f"L{r}"] = f"=D{r}*G{r}"
                ws[f"M{r}"] = f"=D{r}*H{r}"
                ws[f"N{r}"] = f"=D{r}*I{r}"
                ws[f"O{r}"] = f"=L{r}+M{r}+N{r}"
        elif kind == "subtotal":
            ws[f"B{r}"] = spec[1]
            ws[f"B{r}"].font = BOLD
            if not blank or formulas_in_blank:
                for col in "KLMNO":
                    ws[f"{col}{r}"] = f"=SUM({col}{sec_first}:{col}{r - 1})"
            subtotal_rows.append(r)
        r += 1
    # grand total + markups
    tot = r
    ws[f"B{tot}"] = "Kopā"
    ws[f"B{tot}"].font = BOLD
    if not blank or formulas_in_blank:
        for col in "KLMNO":
            ws[f"{col}{tot}"] = "=" + "+".join(f"{col}{x}" for x in subtotal_rows)
    ws[f"B{tot + 1}"] = "Virsizdevumi 8 %"
    ws[f"B{tot + 2}"] = "Peļņa 5 %"
    ws[f"B{tot + 3}"] = "Darba devēja sociālais nodoklis 23,59 %"
    ws[f"B{tot + 4}"] = "Transporta izdevumi"
    ws[f"N{tot + 4}"] = 0.03
    ws[f"N{tot + 4}"].number_format = "0%"
    ws[f"B{tot + 5}"] = "Pavisam kopā"
    ws[f"B{tot + 5}"].font = BOLD
    if not blank or formulas_in_blank:
        ws[f"O{tot + 1}"] = f"=O{tot}*0.08"
        ws[f"O{tot + 2}"] = f"=O{tot}*5%"
        ws[f"O{tot + 3}"] = f"=L{tot}*0.2359"
        ws[f"O{tot + 4}"] = f"=M{tot}*N{tot + 4}"
        ws[f"O{tot + 5}"] = f"=SUM(O{tot}:O{tot + 4})"
    ws[f"B{tot + 7}"] = "Piezīme: cenas norādītas bez PVN."
    for row in ws.iter_rows(min_row=8, max_row=tot + 5, min_col=5, max_col=15):
        for c in row:
            c.number_format = '#,##0.00 "€"' if c.column > 6 else "0.00"
    ws.column_dimensions["B"].width = 45
    return tot


def make_lv_estimate(path: str, *, blank: bool = False) -> str:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Elektroinstalācija"
    t1 = _lv_sheet(ws1, LV_SHEET1, blank=blank, rate_label=12.5)
    ws2 = wb.create_sheet("Vājstrāvas")
    t2 = _lv_sheet(ws2, LV_SHEET2, blank=blank, rate_label=None, formulas_in_blank=False)
    ws3 = wb.create_sheet("Kopsavilkums")
    ws3["A1"] = "Kopsavilkums"
    ws3["A3"] = "Tāme"
    ws3["B3"] = "Summa, EUR"
    ws3["A4"] = "Elektroinstalācija"
    ws3["B4"] = f"='Elektroinstalācija'!O{t1 + 5}"
    ws3["A5"] = "Vājstrāvas"
    ws3["B5"] = f"='Vājstrāvas'!O{t2 + 5}"
    ws3["A6"] = "Kopā"
    ws3["B6"] = "=B4+B5"
    wb.save(path)
    return path


def make_da_estimate(path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tilbud"
    ws["A1"] = "Tilbud på el-installation, Parcelhus"
    ws["A2"] = "Timepris: 450 kr."
    headers = ["Nr", "Tekst", "Antal", "Enhed", "Timer", "Materialer", "Timepris", "Arbejdsløn i alt",
               "Materialer i alt", "I alt", "Kilde"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h).font = BOLD
    rows = [
        ("section", "Stærkstrøm"),
        ("Stikkontakt, dobbelt, LK Fuga", 20, "stk", 0.4, 85, "Solar prisliste 2024"),
        ("Installationskabel PVT 3G1,5", 150, "lbm", 0.05, 6.5, "Solar prisliste 2024"),
        ("Afbryder 1-pol", 10, "stk", 0.3, 60, "AO 2024"),
        ("Gruppetavle 36 modul", 1, "stk", 6.0, 2500, "Tilbud fra grossist"),
        ("LED downlight IP44", 18, "stk", 0.5, 220, "Lemvigh-Müller"),
        ("section", "Svagstrøm"),
        ("Dataudtag RJ45, dobbelt", 6, "stk", 0.5, 140, "Solar prisliste 2024"),
        ("Røgalarm 230V", 5, "stk", 0.4, 310, "AO 2024"),
    ]
    r = 5
    first = r
    for spec in rows:
        if spec[0] == "section":
            ws.cell(row=r, column=2, value=spec[1]).font = BOLD
        else:
            text, qty, unit, hours, mat, src = spec
            ws.cell(row=r, column=1, value=r - 4)
            ws.cell(row=r, column=2, value=text)
            ws.cell(row=r, column=3, value=qty)
            ws.cell(row=r, column=4, value=unit)
            ws.cell(row=r, column=5, value=hours)
            ws.cell(row=r, column=6, value=mat)
            ws.cell(row=r, column=7, value=450)
            ws.cell(row=r, column=8, value=f"=C{r}*E{r}*G{r}")
            ws.cell(row=r, column=9, value=f"=C{r}*F{r}")
            ws.cell(row=r, column=10, value=f"=H{r}+I{r}")
            ws.cell(row=r, column=11, value=src)
        r += 1
    last = r - 1
    ws.cell(row=r, column=2, value="I alt").font = BOLD
    for col in "HIJ":
        ws[f"{col}{r}"] = f"=SUM({col}{first}:{col}{last})"
    ws.cell(row=r + 1, column=2, value="Avance 10%")
    ws[f"J{r + 1}"] = f"=J{r}*0.1"
    ws.cell(row=r + 2, column=2, value="Moms 25%")
    ws[f"J{r + 2}"] = f"=(J{r}+J{r + 1})*0.25"
    ws.cell(row=r + 3, column=2, value="Total inkl. moms").font = BOLD
    ws[f"J{r + 3}"] = f"=SUM(J{r}:J{r + 2})"
    for row in ws.iter_rows(min_row=5, max_row=r + 3, min_col=6, max_col=10):
        for c in row:
            c.number_format = '#,##0.00 "kr."'
    wb.save(path)
    return path


def make_norms(path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Laika normas"
    ws["A1"] = "Elektromontāžas darbu laika normas"
    for i, h in enumerate(["Kategorija", "Darba nosaukums", "Mērv.", "Laika norma, c/h"], start=1):
        ws.cell(row=3, column=i, value=h).font = BOLD
    rows = [
        ("Kabeļi", None, "m", 0.08),
        ("Kabeļi", "Kabelis NYM 3x1,5 mm² ievilkšana caurulē", "m", 0.07),
        ("Kabeļi", "Kabelis NYM 5x2,5 mm² ievilkšana caurulē", "m", 0.10),
        ("Ierīces", "Kontaktligzda zemapmetuma montāža", "gab.", 0.35),
        ("Ierīces", "Slēdzis zemapmetuma montāža", "gab.", "0,30"),
        ("Sadales", "Sadales skapis 36 moduļi", "gab.", 5.5),
        ("Sadales", "Sadales skapis 24 moduļi", "gab.", 4.0),
        ("Sadales", None, "gab.", 3.0),
    ]
    for i, (cat, item, unit, h) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=cat)
        ws.cell(row=i, column=2, value=item)
        ws.cell(row=i, column=3, value=unit)
        ws.cell(row=i, column=4, value=h)
    ws2 = wb.create_sheet("Norms EN")
    for i, h in enumerate(["Item", "Unit", "Hours per unit"], start=1):
        ws2.cell(row=1, column=i, value=h).font = BOLD
    for i, (item, unit, h) in enumerate([
        ("Distribution board 36 modules", "pcs", 5.5),
        ("Cable laying in cable tray", "m", 0.06),
        ("Socket outlet, flush mounted", "pcs", 0.35),
    ], start=2):
        ws2.cell(row=i, column=1, value=item)
        ws2.cell(row=i, column=2, value=unit)
        ws2.cell(row=i, column=3, value=h)
    wb.save(path)
    return path


DOCX_ROWS = [
    ("1", "Kabelis NYM-J 3x1,5 mm²", "m", "120", "1,65", "198,00"),
    ("2", "Kontaktligzda 2-vietīga", "gab.", "24", "11,18", "268,32"),
    ("3", "Gaismeklis LED 36W IP44", "gab.", "16", "45,50", "728,00"),
]


def make_docx(path: str) -> str:
    import docx
    d = docx.Document()
    d.add_heading("Cenu piedāvājums elektroinstalācijas darbiem", level=1)
    d.add_paragraph("Cenas norādītas EUR bez PVN. Stundas likme 12,50 EUR/h.")
    t = d.add_table(rows=1, cols=6)
    for i, h in enumerate(["Nr.", "Nosaukums", "Mērv.", "Daudz.", "Cena, EUR", "Summa, EUR"]):
        t.rows[0].cells[i].text = h
    for row in DOCX_ROWS:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    d.save(path)
    return path


def make_pdf(path: str) -> str | None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        return None
    doc = SimpleDocTemplate(path, pagesize=A4)
    styles = getSampleStyleSheet()
    data = [["No.", "Description", "Unit", "Qty", "Unit price, EUR", "Total, EUR"]]
    data += [["1", "Cable NYM-J 3x1.5 mm2", "m", "120", "1.65", "198.00"],
             ["2", "Double socket outlet IP20", "pcs", "24", "11.18", "268.32"],
             ["3", "LED luminaire 36W IP44", "pcs", "16", "45.50", "728.00"]]
    tbl = Table(data)
    tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([Paragraph("Price offer for electrical installation works", styles["Title"]), tbl])
    return path


def make_encrypted_pdf(src_pdf: str, path: str) -> str | None:
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError:
        return None
    reader = PdfReader(src_pdf)
    writer = PdfWriter()
    for p in reader.pages:
        writer.add_page(p)
    writer.encrypt(user_password="secret", owner_password="owner")
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def make_scanned_pdf(path: str) -> str | None:
    """A PDF with only a drawn rectangle (no text layer) - stands in for a scan."""
    try:
        from reportlab.pdfgen import canvas
    except ImportError:
        return None
    c = canvas.Canvas(path)
    c.rect(50, 50, 400, 600, fill=1)
    c.showPage()
    c.save()
    return path


def make_big(path: str, rows: int = 5000, cols: int = 15) -> str:
    """Large LV-style estimate for performance/RSS checks (write_only for speed)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Tāme")
    ws.append(["Nr.", "Darba nosaukums", "Mērv.", "Daudz.", "Laika norma c/h", "Stundas likme", "Darba alga",
               "Materiāli", "Mehānismi", "Kopā", "Darbietilpība c/h", "Darba alga kopā", "Materiāli kopā",
               "Mehānismi kopā", "Summa"][:cols])
    names = ["Kabelis NYM-J 3x1,5 mm²", "Kabelis NYM-J 3x2,5 mm²", "Kontaktligzda 2-vietīga", "Slēdzis 1-polīgs",
             "Gaismeklis LED 36W IP44", "Gofrētā caurule d20", "Kabeļu trase 100x60", "Datu rozete RJ45"]
    r = 2
    for i in range(rows):
        if i % 50 == 0:
            ws.append([None, f"{i // 50 + 1}. sadaļa"])
            r += 1
            continue
        q = (i % 37) + 1
        n = round(0.05 + (i % 7) * 0.05, 2)
        m = round(0.5 + (i % 11) * 1.3, 2)
        ul = round(n * 12.5, 2)
        ws.append([i, f"{names[i % len(names)]} poz. {i}", "m" if i % 3 else "gab.", q, n, 12.5, ul, m, 0,
                   ul + m, q * n, q * ul, q * m, 0, q * (ul + m)][:cols])
        r += 1
    wb.save(path)
    return path


def soffice_convert(src: str, fmt: str, out_dir: str) -> str | None:
    soffice = shutil.which("soffice") or "/usr/bin/soffice"
    if not os.path.exists(soffice):
        return None
    profile = os.path.join(tempfile.gettempdir(), f"lo_fix_{uuid.uuid4().hex}")
    try:
        subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", fmt,
                        "--outdir", out_dir, src], capture_output=True, timeout=120, check=False,
                       env={**os.environ, "HOME": profile})
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + "." + fmt.split(":")[0])
    return out if os.path.exists(out) else None


_BUILT: dict[str, dict[str, str | None]] = {}


def build_all(out_dir: str) -> dict[str, str | None]:
    if out_dir in _BUILT:
        return _BUILT[out_dir]
    os.makedirs(out_dir, exist_ok=True)
    p = lambda name: os.path.join(out_dir, name)  # noqa: E731
    files: dict[str, str | None] = {
        "lv": make_lv_estimate(p("lv_estimate.xlsx")),
        "lv_blank": make_lv_estimate(p("lv_blank.xlsx"), blank=True),
        "da": make_da_estimate(p("da_estimate.xlsx")),
        "norms": make_norms(p("norms.xlsx")),
        "docx": make_docx(p("price_offer.docx")),
        "pdf": make_pdf(p("price_offer.pdf")),
        "scanned_pdf": make_scanned_pdf(p("scanned.pdf")),
    }
    files["encrypted_pdf"] = make_encrypted_pdf(files["pdf"], p("encrypted.pdf")) if files["pdf"] else None
    _BUILT[out_dir] = files
    return files


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="ingest_fixtures_")
    for k, v in build_all(out).items():
        print(f"{k:14s} {v}")
