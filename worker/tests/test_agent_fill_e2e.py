"""End-to-end: user attaches an xlsx blank → run_agent → priced copy, provenance, cards, permission
flow (Allow → resume re-prices parked rows), viewer edit marks EDITED. Deterministic path (no model)."""
import agent_fixtures as fx  # noqa: I001  (sets env first)

import os

import openpyxl
import pytest
from openpyxl.styles import Font

from app import db, jobs
from app.agent import run as agent_run
from app.agent import documents
from app.llm import client as llm


@pytest.fixture(autouse=True)
def _env():
    fx.migrate()
    object.__setattr__(llm.settings, "llm_enabled", False)
    created: list[str] = []
    yield created
    object.__setattr__(llm.settings, "llm_enabled", True)
    fx.cleanup_files(created)
    db.execute("DELETE FROM jobs")


def make_blank(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elektroinstalācija"
    ws["A1"] = "Lokālā tāme Nr. 3 — Elektroinstalācijas darbi"
    ws["A2"] = "Stundas likme: 14,00 EUR/h"
    hdr = ["Nr.", "Darba nosaukums", "Mērv.", "Daudz.", "Laika norma c/h", "Darba alga EUR", "Materiāli EUR",
           "Kopā EUR", "Darbietilpība c/h", "Darba alga EUR", "Materiāli EUR", "Kopā EUR"]
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.font = Font(bold=True)
    ws.append([None, "Kabeļu līnijas"])
    ws["B5"].font = Font(bold=True)
    ws.append([1, "Kabelis NYM-J 3x1,5 mm2 guldīšana", "m", 120])
    ws.append([2, "Kabelis NYM-J 5x2,5 mm2 guldīšana", "m", 40])
    ws.append([None, "Ierīces"])
    ws["B8"].font = Font(bold=True)
    ws.append([3, "Kontaktligzda 2-vietīga IP44 montāža", "gab.", 12])
    ws.append([4, "Sadales skapis 36 moduļi montāža", "gab.", 1])
    ws.append([5, "Ugunsdrošības vārsts DN200 pieslēgšana", "gab.", 2])
    for r in range(6, 12):
        if ws[f"D{r}"].value is not None:
            ws[f"L{r}"] = f"=J{r}+K{r}"
    ws.append([None, "Kopā", None, None, None, None, None, None, None, "=SUM(J6:J11)", "=SUM(K6:K11)", "=SUM(L6:L11)"])
    wb.save(path)


def test_fill_blank_end_to_end_with_permission_flow(_env):
    ref = fx.make_file("Tame_ref_2025.xlsx", items=[
        {"text": "Kabelis NYM-J 3x1,5 mm2 guldīšana", "unit": "m", "norm_h": 0.1, "labour": 1.2, "material": 0.85, "rate": 12},
        {"text": "Kontaktligzda 2-vietīga IP44 montāža", "unit": "gab.", "norm_h": 0.5, "labour": 6.0, "material": 9.5, "rate": 12},
    ])
    other = fx.make_file("Tame_Riga_2024.xlsx", items=[
        {"text": "Kabelis NYM-J 5x2,5 mm2 guldīšana", "unit": "m", "norm_h": 0.14, "labour": 1.68, "material": 2.1, "rate": 12}])
    norms = fx.make_file("Normas.xlsx", tag="hourly_norms", norms=[
        {"text": "Sadales skapis 36 moduļi", "unit": "gab.", "hours": 5.5, "specificity": "parameterised",
         "params": {"modules": 36}}])
    _env.extend([ref, other, norms])

    u = fx.make_user()
    conv = fx.make_conversation(u)
    up_dir = os.path.join(os.environ["DATA_DIR"], "uploads", fx.uid())
    os.makedirs(up_dir)
    blank = os.path.join(up_dir, "original.xlsx")
    make_blank(blank)
    upload = db.fetchone("""INSERT INTO uploads(user_id, conversation_id, original_name, ext, stored_path)
                            VALUES (%s,%s,'Tame_blank.xlsx','xlsx',%s) RETURNING *""", (u["id"], conv["id"], blank))
    run = fx.make_run(conv, allowed=[ref, norms], selected=[ref, norms], text="Aizpildi šo tāmi",
                      attachments=[str(upload["id"])])
    rid = str(run["id"])

    agent_run.run_agent({"run_id": rid})

    r = db.fetchone("SELECT status, kind, state FROM runs WHERE id=%s", (rid,))
    assert r["kind"] == "fill_blank"
    assert r["status"] == "waiting", r  # one row parked on a permission card
    doc = db.fetchone("SELECT * FROM documents WHERE run_id=%s", (rid,))
    rows = {x["row_idx"]: x for x in db.fetchall("SELECT * FROM estimate_rows WHERE document_id=%s", (doc["id"],))}
    assert set(rows) == {6, 7, 9, 10, 11}
    assert rows[6]["price_source"] == "exact" and float(rows[6]["hourly_rate"]) == 14.0   # blank's own rate
    assert float(rows[6]["unit_labour"]) == pytest.approx(1.4)                            # 0.1 h × 14
    assert rows[7]["price_source"] == "pending_permission"                               # best match unselected
    assert rows[10]["price_source"] == "norm" and float(rows[10]["norm_h_per_unit"]) == 5.5
    assert rows[11]["price_source"] == "none" and "NO PRICE" in rows[11]["flags"]

    wb = openpyxl.load_workbook(doc["stored_path"])
    ws = wb["Elektroinstalācija"]
    assert ws["E6"].value == 0.1 and ws["F6"].value == 1.4 and ws["G6"].value == 0.85
    assert ws["J6"].value == pytest.approx(168.0) and ws["K6"].value == pytest.approx(102.0)
    assert ws["L6"].value == "=J6+K6" and ws["J12"].value == "=SUM(J6:J11)"               # formulas intact
    assert ws["F7"].value is None                                                         # parked row empty
    assert openpyxl.load_workbook(blank)["Elektroinstalācija"]["F6"].value is None        # blank untouched

    cards = {c["kind"]: c for c in db.fetchall("SELECT * FROM cards WHERE run_id=%s", (rid,))}
    assert "document" in cards and cards["permission"]["payload"]["file_name"] == "Tame_Riga_2024.xlsx"
    assert cards["permission"]["payload"]["rows"][0]["row"] == 7
    usage = db.fetchone("SELECT COUNT(*) AS n FROM file_usage WHERE run_id=%s AND file_id=%s", (rid, ref))
    assert usage["n"] == 1
    ev_types = [e["type"] for e in db.fetchall("SELECT type FROM run_events WHERE run_id=%s ORDER BY seq", (rid,))]
    assert ev_types[0] == "run.status" and "step.progress" in ev_types and "document.ready" in ev_types
    assert "card.created" in ev_types and ev_types[-1] == "run.status"

    # --- web approves (state-guarded update + allowed set) and the resume job runs
    perm = cards["permission"]
    db.execute("""UPDATE cards SET status='approved', decision='{"action":"allow","data":{}}', decided_at=now()
                  WHERE id=%s AND status='pending'""", (perm["id"],))
    db.execute("UPDATE runs SET allowed_file_ids = allowed_file_ids || %s::uuid WHERE id=%s", (other, rid))
    agent_run.resume_run({"run_id": rid, "card_id": str(perm["id"])})

    r = db.fetchone("SELECT status FROM runs WHERE id=%s", (rid,))
    assert r["status"] == "done"
    doc2 = db.fetchone("SELECT * FROM documents WHERE id=%s", (doc["id"],))
    assert doc2["version"] == 2
    row7 = db.fetchone("SELECT * FROM estimate_rows WHERE document_id=%s AND row_idx=7", (doc["id"],))
    assert row7["price_source"] == "exact" and row7["matched"][0]["file_id"] == other
    ws2 = openpyxl.load_workbook(doc2["stored_path"])["Elektroinstalācija"]
    assert ws2["F7"].value == pytest.approx(0.14 * 14) and ws2["F6"].value == 1.4

    # --- viewer inspector edit → EDITED, recalculated, new workbook version, visible to later Q&A
    out = documents.update_row(str(doc["id"]), "Elektroinstalācija", 9, {"unit_material": 11.0}, str(u["id"]))
    assert "EDITED" in out["row"]["flags"] and float(out["row"]["total_material"]) == pytest.approx(132.0)
    ws3 = openpyxl.load_workbook(out["document"] and db.fetchone(
        "SELECT stored_path FROM documents WHERE id=%s", (doc["id"],))["stored_path"])["Elektroinstalācija"]
    assert ws3["G9"].value == 11.0 and ws3["K9"].value == pytest.approx(132.0)
    # re-pricing never overwrites an edited row
    documents._upsert_rows  # noqa: B018
    assert db.fetchone("SELECT price_source FROM estimate_rows WHERE document_id=%s AND row_idx=9",
                       (doc["id"],))["price_source"] == "edited"


def test_cancel_before_start_marks_cancelled(_env):
    u = fx.make_user()
    conv = fx.make_conversation(u)
    run = fx.make_run(conv, text="hello")
    db.execute("UPDATE runs SET cancel_requested=true WHERE id=%s", (run["id"],))
    agent_run.run_agent({"run_id": str(run["id"])})
    assert db.fetchone("SELECT status FROM runs WHERE id=%s", (run["id"],))["status"] == "cancelled"
