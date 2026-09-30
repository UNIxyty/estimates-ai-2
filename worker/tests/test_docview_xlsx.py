"""xlsx -> JSON conversion, number formats, formulas, row kinds, paging/filters, caching (no DB)."""
from __future__ import annotations

import datetime as dt
import sys
import threading
import time
import types

import openpyxl
import pytest

from app.docview import cache, xlsx_json
from app.docview.formula import Evaluator
from app.docview.numfmt import format_value, guess_locale
from test_docview_helpers import make_big_xlsx, make_estimate_xlsx, patch_settings, write_garbage


@pytest.fixture(autouse=True)
def _env(tmp_path, monkeypatch):
    patch_settings(monkeypatch, data_dir=str(tmp_path / "data"))
    monkeypatch.setenv("DOCVIEW_USE_INGEST", "0")
    cache.clear_memory()
    yield
    cache.clear_memory()


def _cells(row):
    return {c[0]: c for c in row["cells"]}


# ------------------------------------------------------------------ number formats

@pytest.mark.parametrize("value,fmt,en,eu", [
    (1234.5, "#,##0.00", "1,234.50", "1 234,50"),
    (1234.5, "0.00", "1234.50", "1234,50"),
    (0.1234, "0.0%", "12.3%", "12,3%"),
    (0.5, "0%", "50%", "50%"),
    (-1234.5, "#,##0.00;(#,##0.00)", "(1,234.50)", "(1 234,50)"),
    (-1234.5, "#,##0.00;[Red]-#,##0.00", "-1,234.50", "-1 234,50"),
    (1234.5, '#,##0.00\\ [$€-426]', "1,234.50 €", "1 234,50 €"),
    (1234.5, '#,##0.00 "EUR"', "1,234.50 EUR", "1 234,50 EUR"),
    (123.456, "$#,##0.00", "$123.46", "$123,46"),
    (12.0, "General", "12", "12"),
    (0.1 + 0.2, "General", "0.3", "0,3"),
    (1234567, '#,##0,"k"', "1,235k", "1 235k"),
    (1.5, "0.0#", "1.5", "1,5"),
    (7, "000", "007", "007"),
])
def test_number_formats(value, fmt, en, eu):
    assert format_value(value, fmt, "en") == en
    assert format_value(value, fmt, "eu") == eu


def test_date_and_misc_formats():
    d = dt.datetime(2026, 3, 5, 14, 7, 9)
    assert format_value(d, "dd.mm.yyyy") == "05.03.2026"
    assert format_value(d, "yyyy-mm-dd h:mm") == "2026-03-05 14:07"
    assert format_value(d, "mmm d, yyyy") == "Mar 5, 2026"
    assert format_value(d, "hh:mm:ss") == "14:07:09"
    assert format_value(46086, "dd.mm.yyyy") == "05.03.2026"  # serial number with a date format
    assert format_value(True, "General") == "TRUE"
    assert format_value("teksts", "#,##0.00") == "teksts"
    assert format_value(None, "0.00") == ""
    assert format_value(0, '_-* #,##0.00\\ "€"_-;\\-* #,##0.00\\ "€"_-;_-* "-"??\\ "€"_-;_-@_-').strip() == "- €"
    assert guess_locale({'#,##0.00\\ [$€-426]'}) == "eu"
    assert guess_locale({"#,##0.00", "0%"}) == "en"


# ------------------------------------------------------------------ formula fallback

def test_formula_evaluator():
    cells = {
        "S": {(1, 1): 2, (1, 2): 3, (2, 1): "=A1*B1", (3, 1): "=SUM(A1:B2)", (4, 1): "=IFERROR(1/0,7)",
              (5, 1): "=ROUND(A2/7,2)", (6, 1): '=IF(A1>1,"big","small")&"!"', (7, 1): "='O t'!A1+10%",
              (8, 1): "=A8+1", (9, 1): "=SUMPRODUCT(A1:A2,B1:B2)", (10, 1): "=VLOOKUP(1,A1:B2,2)",
              (11, 1): "=$A$1^2-(-B1)"},
        "O t": {(1, 1): 5},
    }
    ev = Evaluator(cells, {"S": {(2, 1): 99}})  # cached value wins over evaluation
    assert ev.value("S", 2, 1) == 99
    ev = Evaluator(cells, {})
    assert ev.value("S", 2, 1) == 6
    assert ev.value("S", 3, 1) == 11
    assert ev.value("S", 4, 1) == 7
    assert ev.value("S", 5, 1) == 0.86
    assert ev.value("S", 6, 1) == "big!"
    assert ev.value("S", 7, 1) == pytest.approx(5.1)
    assert ev.value("S", 8, 1) is None  # circular
    assert ev.value("S", 9, 1) == 6
    assert ev.value("S", 10, 1) is None  # unsupported function
    assert ev.value("S", 11, 1) == 7


# ------------------------------------------------------------------ conversion

def test_convert_estimate(tmp_path):
    p = make_estimate_xlsx(str(tmp_path / "est.xlsx"))
    m = xlsx_json.convert_workbook(p)
    assert m["locale"] == "eu"  # [$€-426] (lv-LV) in a format
    s = m["sheets"][0]
    assert s["name"] == "Tāme" and s["idx"] == 0
    assert s["row_count"] == 10 and s["col_count"] == 7
    assert s["frozen"] == {"rows": 3, "cols": 2}
    assert s["merges"] == [{"r1": 1, "c1": 1, "r2": 1, "c2": 7}]
    assert s["hidden_rows"] == [2] and s["hidden_cols"] == [5]
    assert s["col_widths"][1] == 285 and s["col_widths"][0] == 64
    rows = {r["r"]: r for r in s["rows"]}
    assert len(s["rows"]) == 10
    assert rows[3]["h"] == 40
    assert rows[2].get("hidden") is True

    title = _cells(rows[1])[1]
    assert "m" in title[3] and "b" in title[3] and "c" in title[3]

    kinds = {r: rows[r]["kind"] for r in rows}
    assert kinds == {1: "note", 2: "note", 3: "header", 4: "section", 5: "item", 6: "item", 7: "item",
                     8: "subtotal", 9: "blank", 10: "total"}

    item = _cells(rows[5])
    assert item[5][1] == 1.25 and item[5][2] == "1,25"
    assert item[6][1] == 150 and item[6][2] == "150,00" and "f" in item[6][3] and item[6][4] == "=D5*E5"
    assert item[7][2] == "12,5%"
    assert "r" in item[4][3]  # numbers right-align by default
    assert "bg=FFF2CC" in _cells(rows[4])[2][3]
    sub = _cells(rows[8])[6]
    assert sub[1] == 392 and sub[2] == "392,00 €" and "fg=C00000" in sub[3] and sub[3].startswith("bf")
    assert _cells(rows[10])[6][1] == pytest.approx(474.32)
    assert "w" in _cells(rows[3])[1][3].split(";")[0]
    # cross-sheet formula without cached value
    assert _cells(m["sheets"][1]["rows"][0])[2][1] == 392
    assert _cells(rows[2])[2][2] == "05.03.2026"


def test_deep_formula_chain(tmp_path):
    """Running totals thousands of rows deep must not hit the recursion limit."""
    p = str(tmp_path / "chain.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = 1, "=A1"
    for r in range(2, 4001):
        ws.cell(r, 1, 1)
        ws.cell(r, 2, f"=B{r - 1}+A{r}")
    ws["C1"] = "=B4000*2"  # forward reference to the end of the chain
    wb.save(p)
    rows = xlsx_json.convert_workbook(p)["sheets"][0]["rows"]
    assert _cells(rows[3999])[2][1] == 4000
    assert _cells(rows[0])[3][1] == 8000


def test_cached_values_preferred(tmp_path, monkeypatch):
    """When Excel/LibreOffice already computed a value, the cached one is shown (not re-evaluated)."""
    p = str(tmp_path / "c.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = 2, "=A1*10"
    wb.save(p)
    import app.docview.xlsx_json as xj

    real = xj._load

    def fake_load(path, **kw):
        wb = real(path, **kw)
        if kw.get("data_only"):
            class _Cell:
                def __init__(s, r, c, v):
                    s.row, s.column, s.value = r, c, v

            class _WS:
                title = "Sheet"

                def reset_dimensions(s):
                    pass

                def iter_rows(s, max_row=None):
                    return [[_Cell(1, 1, 2)], [_Cell(2, 1, 12345)]]

            class _WB:
                worksheets = [_WS()]

                def close(s):
                    pass
            return _WB()
        return wb

    monkeypatch.setattr(xj, "_load", fake_load)
    m = xj.convert_workbook(p)
    assert _cells(m["sheets"][0]["rows"][1])[1][1] == 12345


def test_ingest_structure_preferred(tmp_path, monkeypatch):
    p = make_estimate_xlsx(str(tmp_path / "est.xlsx"))
    mod = types.ModuleType("app.ingest.structure")

    class RowInfo:
        def __init__(self, row, kind):
            self.row, self.kind = row, kind

    def analyse_workbook(path):
        return types.SimpleNamespace(sheets=[{"name": "Tāme", "rows": [RowInfo(1, "header"), RowInfo(5, "section"),
                                                                        RowInfo(6, "bogus")]}])

    mod.analyse_workbook = analyse_workbook
    monkeypatch.setitem(sys.modules, "app.ingest.structure", mod)
    monkeypatch.setenv("DOCVIEW_USE_INGEST", "1")
    m = xlsx_json.convert_workbook(p)
    rows = {r["r"]: r for r in m["sheets"][0]["rows"]}
    assert rows[1]["kind"] == "header" and rows[5]["kind"] == "section"
    assert rows[6]["kind"] == "item"  # unknown kind -> heuristic kept
    assert m["sheets"][0]["kind_source"] == "ingest" and m["sheets"][1]["kind_source"] == "heuristic"

    def broken(path):
        raise RuntimeError("boom")

    mod.analyse_workbook = broken
    m = xlsx_json.convert_workbook(p)
    assert m["sheets"][0]["kind_source"] == "heuristic"
    assert {r["r"]: r["kind"] for r in m["sheets"][0]["rows"]}[5] == "item"


def test_corrupt_workbook(tmp_path):
    p = write_garbage(str(tmp_path / "bad.xlsx"))
    with pytest.raises(cache.Unreadable) as e:
        xlsx_json.load(p)
    assert "zip" in e.value.reason
    n = cache.stats["conversions"]
    with pytest.raises(cache.Unreadable):
        xlsx_json.load(p)  # negative result is cached: no second parse
    assert cache.stats["conversions"] == n


# ------------------------------------------------------------------ paging / filters / find

@pytest.fixture
def big(tmp_path):
    return xlsx_json.load(make_big_xlsx(str(tmp_path / "big.xlsx"), 1500))


def test_pagination(big):
    s = big.model["sheets"][0]
    assert s["row_count"] == 1500 and s["col_count"] == 15
    assert s["hidden_rows"] == [500, 1000, 1500]
    assert s["frozen"] == {"rows": 1, "cols": 0}
    res = xlsx_json.query_rows(big, 0, offset=0, limit=200)
    assert res["total"] == 1500 and len(res["rows"]) == 200 and res["rows"][0]["r"] == 1
    assert "matches" not in res
    res = xlsx_json.query_rows(big, 0, offset=1400, limit=200)
    assert len(res["rows"]) == 100 and res["rows"][-1]["r"] == 1500
    res = xlsx_json.query_rows(big, 0, offset=0, limit=10_000)
    assert res["limit"] == 500 and len(res["rows"]) == 500
    res = xlsx_json.query_rows(big, 0, offset=5000, limit=100)
    assert res["rows"] == [] and res["total"] == 1500
    kinds = {r["kind"] for r in xlsx_json.query_rows(big, 0, limit=500)["rows"]}
    assert {"header", "section", "item"} <= kinds


def test_around(big):
    res = xlsx_json.query_rows(big, 0, limit=200, around=777)
    assert res["offset"] == 600 and any(r["r"] == 777 for r in res["rows"])
    res = xlsx_json.query_rows(big, 0, limit=200, around=99999)
    assert res["offset"] == 1400
    res = xlsx_json.query_rows(big, 0, limit=50, around=1)
    assert res["offset"] == 0


def test_search_and_matches(big):
    res = xlsx_json.query_rows(big, 0, limit=5, q="RINDA 77")
    # "rinda 77" and "rinda 770".."779" (sections carry no "rinda")
    expected = [r for r in range(2, 1501) if "rinda 77" in f"rinda {r}" and (r - 2) % 50 != 0]
    assert res["total"] == len(expected)
    assert res["matches"] == expected
    assert [r["r"] for r in res["rows"]] == expected[:5]
    # around within a filtered result jumps to the page containing that match
    res = xlsx_json.query_rows(big, 0, limit=5, q="rinda 77", around=776)
    assert res["offset"] == 5 and 776 in [r["r"] for r in res["rows"]]
    many = xlsx_json.query_rows(big, 0, limit=1, q="kabelis")
    assert many["total"] > 500 and len(many["matches"]) == 500
    none = xlsx_json.query_rows(big, 0, q="nav tāda")
    assert none["total"] == 0 and none["rows"] == [] and none["matches"] == []
    # display search: formatted numbers are searchable ("249.60")
    assert xlsx_json.query_rows(big, 0, q="249.60")["total"] >= 1


def test_marker_filters_on_plain_workbook(big):
    # knowledge files have no provenance: marker filters return nothing
    for f in ("web", "check", "no_price", "edited", "flagged"):
        assert xlsx_json.query_rows(big, 0, filter_=f, markers=None)["total"] == 0


def test_markers_merge_and_filters(big):
    def m(marker, flags, source="exact"):
        return {"marker": marker, "flags": flags, "price_source": source, "confidence": "low",
                "confidence_pct": 40, "estimate_row_id": f"id-{marker}", "_eff": set(flags)}

    mk = {3: m("WEB", ["WEB"]), 4: m("CHECK", ["CHECK"]), 5: m("EDITED", ["EDITED", "WEB"]), 6: m(None, [])}
    res = xlsx_json.query_rows(big, 0, filter_="web", markers=mk)
    assert [r["r"] for r in res["rows"]] == [3, 5] and res["total"] == 2
    assert xlsx_json.query_rows(big, 0, filter_="flagged", markers=mk)["total"] == 3
    assert xlsx_json.query_rows(big, 0, filter_="edited", markers=mk)["total"] == 1
    page = xlsx_json.query_rows(big, 0, limit=10, markers=mk)["rows"]
    r3 = next(r for r in page if r["r"] == 3)
    assert r3["marker"] == "WEB" and r3["estimate_row_id"] == "id-WEB" and "_eff" not in r3
    assert "marker" not in next(r for r in page if r["r"] == 7)
    # merging must not mutate the cached model
    assert "marker" not in big.model["sheets"][0]["rows"][2]


# ------------------------------------------------------------------ caching

def test_cache_memory_disk_and_single_conversion(tmp_path):
    p = make_big_xlsx(str(tmp_path / "c.xlsx"), 1500)
    c0 = cache.stats["conversions"]
    t = time.perf_counter()
    wb1 = xlsx_json.load(p)
    first = time.perf_counter() - t
    assert cache.stats["conversions"] == c0 + 1

    t = time.perf_counter()
    wb2 = xlsx_json.load(p)
    mem = time.perf_counter() - t
    assert wb2 is wb1 and cache.stats["conversions"] == c0 + 1

    cache.clear_memory()
    d0 = cache.stats["disk_hits"]
    t = time.perf_counter()
    wb3 = xlsx_json.load(p)
    disk = time.perf_counter() - t
    assert cache.stats["conversions"] == c0 + 1 and cache.stats["disk_hits"] == d0 + 1
    assert wb3.model == wb1.model
    assert disk < first and mem < 0.01
    print(f"\n[docview cache 1500 rows] first={first * 1000:.0f}ms disk={disk * 1000:.0f}ms "
          f"memory={mem * 1e6:.0f}us")

    # a modified file gets a new key and is reconverted
    wb = openpyxl.load_workbook(p)
    wb.active["B3"] = "Changed text"
    wb.save(p)
    wb4 = xlsx_json.load(p)
    assert cache.stats["conversions"] == c0 + 2
    assert _cells(wb4.model["sheets"][0]["rows"][2])[2][1] == "Changed text"


def test_concurrent_first_requests_convert_once(tmp_path):
    p = make_big_xlsx(str(tmp_path / "cc.xlsx"), 800)
    c0 = cache.stats["conversions"]
    results = []
    threads = [threading.Thread(target=lambda: results.append(xlsx_json.load(p))) for _ in range(6)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert cache.stats["conversions"] == c0 + 1
    assert len(results) == 6 and all(r is results[0] for r in results)
