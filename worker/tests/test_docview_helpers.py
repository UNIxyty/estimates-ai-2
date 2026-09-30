"""Shared fixture generators for the docview tests (no tests in here)."""
from __future__ import annotations

import base64
import datetime as dt
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

TEST_DB_URL = os.environ.get("TEST_DATABASE_URL", "postgresql://estimates:dev@127.0.0.1:5433/estimates")

# 1x1 transparent PNG
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def make_estimate_xlsx(path: str) -> str:
    """A small, realistic Latvian estimate: header, sections, items, subtotals, total, merges, formulas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tāme"
    ws["A1"] = "Lokālā tāme Nr. 1"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:G1")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Datums"
    ws["B2"] = dt.datetime(2026, 3, 5)
    ws["B2"].number_format = "dd.mm.yyyy"
    ws["A3"] = "Nr."
    ws["B3"] = "Nosaukums"
    ws["C3"] = "Mērv."
    ws["D3"] = "Daudz."
    ws["E3"] = "Cena"
    ws["F3"] = "Summa"
    ws["G3"] = "Uzcenojums"
    for c in "ABCDEFG":
        ws[f"{c}3"].font = Font(bold=True)
        ws[f"{c}3"].alignment = Alignment(wrap_text=True)
    ws["A4"] = 1
    ws["B4"] = "Kabeļu līnijas"
    ws["B4"].font = Font(bold=True)
    ws["B4"].fill = PatternFill("solid", fgColor="FFF2CC")
    items = [("Kabelis NYM 3x1.5", "m", 120, 1.25), ("Kabelis NYM 3x2.5", "m", 80, 1.9),
             ("Rozete 2-vietīga IP44", "gab", 12, 7.5)]
    r = 5
    for name, unit, qty, price in items:
        ws.cell(r, 2, name)
        ws.cell(r, 3, unit)
        ws.cell(r, 4, qty)
        ws.cell(r, 5, price).number_format = "#,##0.00"
        ws.cell(r, 6, f"=D{r}*E{r}").number_format = "#,##0.00"
        ws.cell(r, 7, 0.125).number_format = "0.0%"
        r += 1
    ws.cell(r, 2, "Kopā sadaļā").font = Font(bold=True)
    ws.cell(r, 6, f"=SUM(F5:F{r - 1})").number_format = '#,##0.00\\ [$€-426]'
    ws.cell(r, 6).font = Font(bold=True, color="FFC00000")
    sub_row = r
    r += 2  # blank row between
    ws.cell(r, 2, "Pavisam kopā").font = Font(bold=True)
    ws.cell(r, 6, f"=F{sub_row}*1.21").number_format = "#,##0.00"
    ws.cell(r, 2).alignment = Alignment(horizontal="right")
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["E"].hidden = True
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "C4"
    ws2 = wb.create_sheet("Kopsavilkums")
    ws2["A1"] = "Kopā"
    ws2["B1"] = "='Tāme'!F" + str(sub_row)
    wb.save(path)
    return path


def make_big_xlsx(path: str, rows: int, cols: int = 15, *, section_every: int = 50) -> str:
    """Large sheet: header, section row every ``section_every`` rows, items with formulas, hidden rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Big"
    headers = ["Nr.", "Nosaukums", "Mērv.", "Daudz.", "Norma h", "Darbs", "Materiāls", "Summa"] + [
        f"Kol {i}" for i in range(9, cols + 1)]
    bold = Font(bold=True)
    for c, h in enumerate(headers[:cols], 1):
        ws.cell(1, c, h).font = bold
    ws.freeze_panes = "A2"
    for r in range(2, rows + 1):
        if (r - 2) % section_every == 0:
            ws.cell(r, 2, f"Sadaļa {r // section_every + 1}").font = bold
            continue
        ws.cell(r, 1, r)
        ws.cell(r, 2, f"Kabelis NYM 3x{(r % 6) + 1}.5 rinda {r}")
        ws.cell(r, 3, "m")
        ws.cell(r, 4, (r % 97) + 1)
        ws.cell(r, 5, 0.05 + (r % 7) / 100).number_format = "0.000"
        ws.cell(r, 6, 12.5 * ((r % 5) + 1)).number_format = "#,##0.00"
        ws.cell(r, 7, 3.1 * ((r % 11) + 1)).number_format = "#,##0.00"
        ws.cell(r, 8, f"=D{r}*(F{r}+G{r})").number_format = "#,##0.00"
        for c in range(9, cols + 1):
            ws.cell(r, c, (r * c) % 1000 / 10)
        if r % 500 == 0:
            ws.row_dimensions[r].hidden = True
    wb.save(path)
    return path


def make_docx(path: str) -> str:
    import docx
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    d = docx.Document()
    d.add_heading("Darbu apraksts", level=1)
    d.add_paragraph("Ievads <script>alert('x')</script> & teksts")
    p = d.add_paragraph("Bold part ")
    p.add_run("bold").bold = True

    # hyperlinks: one https, one javascript:
    def add_link(par, url, text):
        part = par.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        h.append(run)
        par._p.append(h)

    lp = d.add_paragraph("Links: ")
    add_link(lp, "https://example.com/spec", "spec")
    add_link(lp, "javascript:alert(1)", "evil")

    t = d.add_table(rows=2, cols=3)
    for i, txt in enumerate(["Pozīcija", "Daudz.", "Cena"]):
        t.cell(0, i).text = txt
    for i, txt in enumerate(["Rozete", "12", "7,50"]):
        t.cell(1, i).text = txt
    img = os.path.join(os.path.dirname(path), "px.png")
    with open(img, "wb") as f:
        f.write(PNG_1PX)
    d.add_picture(img)
    d.save(path)
    return path


def write_garbage(path: str, data: bytes = b"this is not a zip file at all") -> str:
    with open(path, "wb") as f:
        f.write(data)
    return path


def patch_settings(mp, **changes):
    """Replace the frozen settings object seen by app.config / app.db / docview (call-time lookups)."""
    import dataclasses

    from app import config

    new = dataclasses.replace(config.settings, **changes)
    mp.setattr(config, "settings", new)
    try:
        from app import db

        mp.setattr(db, "settings", new)
    except Exception:  # noqa: BLE001
        pass
    return new
