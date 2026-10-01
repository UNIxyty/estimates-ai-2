"""Which sheets to price: prefer sheets with price + quantity columns and a Tāme layout over specification sheets
whose wording is merely "more electrical" (the Eiženijas iela blank: EL-1…4 are specifications, Tame MājaA/B are
the estimate)."""
import agent_fixtures as fx  # noqa: I001  (sets env first)

import glob
import os
import shutil
import tempfile

import pytest
from openpyxl import Workbook

from app.agent.fill import SUGGEST_AT, rank_sheets
from app.ingest.structure import analyse_workbook

SPEC_ITEMS = ["Grupu sadalnes komplekts 36 moduļi", "Zemapmetuma slēdzis 1-polīgs", "Kontaktligzda 2P+E 16A IP44",
              "Kabelis NYM-J 3x2,5 mm2", "Automātslēdzis 16A 1P B", "Noplūdes strāvas aizsargiekārta 40A 30mA"]
TAME_ITEMS = ["Tips 1 (S1, S10) UK636, individuālas", "Tips 2 (S2, S11) UK636, individuālas",
              "Tips 3 (S3, S12) UK648, individuālas", "Montāžas palīgmateriāli", "Iezemēšanas kontūrs"]
TAME_HEAD = ["Nr.p.k.", "Darba nosaukums", "Mērvienība", "Daudzums", "Laika norma (c/h)", "Darba samaksas likme (€/h)",
             "Darba alga (€)", "Būvizstrādājumi (€)", "Mehānismi (€)", "Kopā (€)", "Darbietilpība (c/h)",
             "Darba alga (€)", "Būvizstrādājumi (€)", "Mehānismi (€)", "Kopā (€)"]


def _eizenijas_like(path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for n in range(1, 5):
        ws = wb.create_sheet(f"EL-{n}")
        ws.append(["Daudzums", "Poz.", "Nosaukums", "Mērv.", "Ražotājs", "Piezīmes"])
        for i, t in enumerate(SPEC_ITEMS):
            ws.append([i + 2, f"E{i}", t, "gab.", "ABB", ""])
    for name in ("Tame MājaA", "Tame MājaB"):
        ws = wb.create_sheet(name)
        ws.append(["Lokālā tāme Nr. 1"])
        ws.append([None, None, None, None, "Vienības izmaksas", None, None, None, None, None, "Kopā uz visu apjomu"])
        ws.append(TAME_HEAD)
        ws.append([None, "Elektroapgāde"])
        for i, t in enumerate(TAME_ITEMS):
            r = ws.max_row + 1
            ws.append([i + 1, t, "kompl.", 3, None, None, None, None, None, f"=G{r}+H{r}+I{r}", f"=D{r}*E{r}",
                       f"=D{r}*G{r}", f"=D{r}*H{r}", f"=D{r}*I{r}", f"=L{r}+M{r}+N{r}"])
    wb.save(path)


def _suggested(path: str) -> tuple[list[str], dict]:
    ws = analyse_workbook(path)
    candidates = [s for s in ws.sheets if any(r.kind == "item" for r in s.rows)]
    ranked = rank_sheets(candidates)
    conf = {s.name: (c, why) for s, c, why in ranked}
    return [s.name for s, c, _ in ranked if c >= SUGGEST_AT], conf


def test_tame_sheets_with_price_columns_beat_electrical_specifications(tmp_path):
    p = str(tmp_path / "eizenijas_like.xlsx")
    _eizenijas_like(p)
    suggested, conf = _suggested(p)
    assert set(suggested) == {"Tame MājaA", "Tame MājaB"}, conf
    assert all(conf[f"EL-{n}"][0] <= 0.3 and "no price columns" in conf[f"EL-{n}"][1] for n in range(1, 5))


def _real_eizenijas() -> str | None:
    root = os.environ.get("GATE_ESTIMATES_DIR") or os.path.join(os.path.dirname(__file__), "..", "..", "gate-estimates")
    hits = glob.glob(os.path.join(root, "EL-T*Eiž*nijas*.xls")) + glob.glob(os.path.join(root, "EL-T*Eiz*nijas*.xls"))
    return hits[0] if hits else None


@pytest.mark.skipif(not _real_eizenijas() or not (shutil.which("soffice") or shutil.which("libreoffice")),
                    reason="real Eiženijas iela estimate (gitignored client file) or LibreOffice not available")
def test_real_eizenijas_blank_suggests_the_tame_sheets():
    from app.agent.common import ensure_xlsx
    d = tempfile.mkdtemp()
    src = os.path.join(d, "eizenijas.xls")
    shutil.copyfile(_real_eizenijas(), src)
    suggested, conf = _suggested(ensure_xlsx(src))
    assert set(suggested) == {"Tame MājaA", "Tame MājaB"}, conf
