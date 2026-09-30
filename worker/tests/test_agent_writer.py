import agent_fixtures as fx  # noqa: F401,I001  (sets env first)

import os
import tempfile

import openpyxl
from openpyxl.styles import Font, PatternFill

from app.agent import writer


def _blank(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elektro"
    ws.append(["Nr.", "Darba nosaukums", "Mērv.", "Daudz.", "Norma", "Darbs", "Materiāli", "Kopā", "Avots"])
    ws.append([1, "Kabelis NYM 3x1,5", "m", 120, None, None, None, "=D2*(F2+G2)", None])
    ws.append([2, "Kontaktligzda", "gab.", 4, None, 0, None, None, "manuāli"])
    ws.merge_cells("F4:G4")
    ws.append([3, "Merged row", "gab.", 1, None, None, None, None, None])
    ws["G2"].font = Font(bold=True, color="FF0000")
    ws["G2"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["G2"].number_format = "#,##0.00"
    ws.column_dimensions["B"].width = 42
    wb.save(path)


def _row(r, **kw):
    base = {"sheet_name": "Elektro", "row_idx": r, "qty": None, "norm_h_per_unit": None, "hourly_rate": None,
            "unit_labour": None, "unit_material": None, "price_source": "exact",
            "matched": [{"file_name": "ref.xlsx", "row": 7}], "norm_ref": None, "web": None}
    return {**base, **kw}


def test_writes_only_value_cells_and_keeps_formulas_styles_merges_widths():
    d = tempfile.mkdtemp()
    src, dst = os.path.join(d, "blank.xlsx"), os.path.join(d, "out.xlsx")
    _blank(src)
    lay = writer.SheetLayout("Elektro", {"item": "B", "unit": "C", "qty": "D", "norm_h": "E", "unit_labour": "F",
                                         "unit_material": "G", "total": "H", "source": "I"})
    rows = [_row(2, qty=120, norm_h_per_unit=0.1, hourly_rate=12, unit_labour=1.2, unit_material=0.85),
            _row(3, qty=4, unit_labour=6.0, unit_material=9.5, price_source="norm",
                 norm_ref={"file_name": "normas.xlsx"}),
            _row(4, qty=1, unit_labour=5.0, unit_material=1.0)]
    rep = writer.write_rows(src, dst, {"Elektro": lay}, rows, "LV")
    wb = openpyxl.load_workbook(dst)
    ws = wb["Elektro"]
    assert ws["H2"].value == "=D2*(F2+G2)"                    # formula untouched
    assert ws["E2"].value == 0.1 and ws["F2"].value == 1.2 and ws["G2"].value == 0.85
    assert ws["G2"].font.bold and ws["G2"].fill.fgColor.rgb.endswith("FFFF00") and ws["G2"].number_format == "#,##0.00"
    assert ws["I2"].value == "Atsauce: ref.xlsx r. 7"         # source column in the blank's language
    assert ws["F3"].value == 6.0                              # a 0 placeholder in a price cell is filled
    assert ws["I3"].value == "manuāli"                        # existing text never overwritten
    assert ws["H3"].value == 62.0                             # empty non-formula row-total filled
    assert ws["F4"].value == 5.0 and ws["G4"].value is None   # merged non-anchor cell skipped
    assert "F4:G4" in {str(r) for r in ws.merged_cells.ranges}
    assert ws.column_dimensions["B"].width == 42
    assert rep.skipped_formula >= 1 and rep.skipped_merged == 1
    assert openpyxl.load_workbook(src)["Elektro"]["F2"].value is None  # original blank unchanged


def test_source_labels_follow_blank_language():
    r = _row(2, price_source="web", web={"url": "https://x"})
    assert writer.source_label(r, "DA") == "Web: https://x"
    assert writer.source_label({"price_source": "none"}, "LV") == "NAV CENAS"
    assert writer.source_label({"price_source": "none"}, "DA") == "INGEN PRIS"
