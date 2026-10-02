"""Unit-rate blank filling (2.4): plan, cell writing, recalculation, and the chat run path."""
import agent_fixtures as fx  # noqa: I001  (sets env first)

import os
from copy import copy
import shutil

import openpyxl
import pytest

from app import db
from app.unitrate.boq import read_boq
from app.unitrate.fill import Planner, Setup, extras_from_workbook, verify, write_plan
from app.unitrate.gate import make_blank
from app.unitrate.match import RefMeta, UnitRateKB, Want, rows_from_workbook
from unitrate_fixtures import goodman_like

REF_NAME = "3733 - Goodman AMS01 - Electrical Services Installation - BOQ.xlsx"


@pytest.fixture(scope="module")
def ref(tmp_path_factory):
    p = str(tmp_path_factory.mktemp("urf") / "ref.xlsx")
    goodman_like(p)
    return p


def _blank(ref: str, tmp_path, edit=None) -> str:
    """The reference with every M.G.S. input cleared, plus optional edits (a new line, a changed quantity …)."""
    b = str(tmp_path / "blank.xlsx")
    make_blank(ref, b)
    if edit:
        book = openpyxl.load_workbook(b)
        edit(book)
        book.save(b)
    return b


def _plan(ref: str, blank: str, setup: Setup | None = None):
    rwb = read_boq(ref, file_name=REF_NAME)
    meta = RefMeta("ref", REF_NAME, market=rwb.market, client=rwb.client, end_client=rwb.end_client,
                   package=rwb.package, doc_date=rwb.doc_date)
    bwb = read_boq(blank, file_name="blank.xlsx")
    kb = UnitRateKB(rows_from_workbook(rwb, meta), Want(market=bwb.market, client=bwb.client, package=bwb.package,
                                                        end_client=bwb.end_client))
    setup = setup or Setup(phase_weeks={"Programme": 78}, offer_date="2026-10-02", validity_date="2026-12-02")
    plan = Planner(bwb, kb, setup, extras=extras_from_workbook(rwb, meta),
                   ref_ratios=[(meta, rr) for rr in rwb.ratios]).plan()
    return rwb, bwb, plan


def _line(plan, sheet, row):
    return next(ln for ln in plan.lines if ln.sheet == sheet and ln.row == row)


def test_blank_gets_the_reference_rates_and_keeps_formulas(ref, tmp_path):
    blank = _blank(ref, tmp_path)
    rwb, bwb, plan = _plan(ref, blank)
    truth = {(sh.name, r.row): r for sh in rwb.pricing_sheets for r in sh.rows if r.kind == "item"}
    for ln in plan.lines:
        if ln.kind != "item":
            continue
        t = truth[(ln.sheet, ln.row)]
        if t.install_rate is not None and t.rate_basis != "supply_only":
            assert ln.install_rate == t.install_rate, (ln.sheet, ln.row, ln.description)
    out = str(tmp_path / "filled.xlsx")
    rep = write_plan(blank, out, plan, bwb)
    ws = openpyxl.load_workbook(out)["Containment Install"]
    assert ws["G15"].value == 25 and ws["H15"].value == "=E15*G15"     # rate written, formula total intact
    assert ws["E15"].value == "=SUM(C15:D15)"                           # quantities untouched
    assert rep.skipped_formula and rep.written > 20


def test_free_issue_and_supply_scope(ref, tmp_path):
    blank = _blank(ref, tmp_path)
    _, bwb, plan = _plan(ref, blank)
    fitting = _line(plan, "Lighting Install", 15)            # STANDARD FITTINGS (Free Issued for Install)
    assert fitting.install_rate == 30 and fitting.supply_rate is None
    assert not any(w["role"] == "supply" for w in fitting.writes)
    bracket = _line(plan, "Lighting Install", 19)            # STANDARD CABLING (Supply & Install)
    assert (bracket.install_rate, bracket.supply_rate) == (10, 7)
    # material supply not in our scope: no supply anywhere
    _, _, plan2 = _plan(ref, blank, Setup(phase_weeks={"Programme": 78}, material_supply=False))
    assert all(ln.supply_rate is None for ln in plan2.lines)


def test_not_participating_and_unticked_packages_are_left_empty(ref, tmp_path):
    blank = _blank(ref, tmp_path)
    setup = Setup(phase_weeks={"Programme": 78},
                  participating={"Lighting Installation": False, "Lighting Installation supply of the materials": False})
    _, bwb, plan = _plan(ref, blank, setup)
    assert not [ln for ln in plan.lines if ln.sheet == "Lighting Install"]
    assert any(s["sheet"] == "Lighting Install" and "unticked" in s["reason"] for s in plan.skipped_sheets)
    assert any(s["sheet"] == "Site Attendance" for s in plan.skipped_sheets)


def test_only_input_cells_are_written(ref, tmp_path):
    def unblue(book):
        book["Containment Install"]["G16"].fill = openpyxl.styles.PatternFill(fill_type=None)
    blank = _blank(ref, tmp_path, unblue)
    _, bwb, plan = _plan(ref, blank)
    out = str(tmp_path / "filled.xlsx")
    rep = write_plan(blank, out, plan, bwb)
    ws = openpyxl.load_workbook(out)["Containment Install"]
    assert ws["G16"].value is None and "Containment Install!G16" in rep.skipped_not_input
    assert ws["G15"].value == 25


def test_a_rate_column_without_the_input_fill_is_still_the_input_column(ref, tmp_path):
    # FR12X / Small power: some cells are light-blue, the Rate column never is
    def unblue_rates(book):
        ws = book["Containment Install"]
        for r in range(10, 26):   # item rows lose the fill; the contractor items below keep theirs
            ws[f"G{r}"].fill = openpyxl.styles.PatternFill(fill_type=None)
    blank = _blank(ref, tmp_path, unblue_rates)
    _, bwb, plan = _plan(ref, blank)
    out = str(tmp_path / "filled.xlsx")
    rep = write_plan(blank, out, plan, bwb)
    ws = openpyxl.load_workbook(out)["Containment Install"]
    assert ws["G15"].value == 25 and ws["G16"].value == 45
    assert not [c for c in rep.skipped_not_input if c.startswith("Containment Install!G")]


def test_contractor_items_delivery_prelims_and_summary(ref, tmp_path):
    blank = _blank(ref, tmp_path)
    setup = Setup(phase_weeks={"Programme": 78}, delivery_pct=12, offer_date="2026-10-02",
                  validity_date="2026-12-02", subcontractor="M.G.S. IT LLC", prelims_weeks=48)
    _, bwb, plan = _plan(ref, blank, setup)
    acc = next(ln for ln in plan.lines if ln.kind == "contractor" and "Accom" in ln.description
               and ln.sheet == "Containment Install")
    assert acc.qty == 18 and acc.install_rate == 20000                 # 78 weeks → 18 months × €20,000
    deliv = next(ln for ln in plan.lines if ln.kind == "delivery")
    supply = sum(ln.supply_total for ln in plan.lines if ln.sheet == deliv.sheet and ln.kind == "item")
    assert deliv.amount == pytest.approx(0.12 * supply, abs=0.01)
    assert {w["role"] for w in deliv.writes} >= {"qty", "supply_total"}
    sm = next(ln for ln in plan.lines if ln.kind == "prelim" and ln.description == "Site Manager")
    assert sm.install_rate == 2200 and sm.qty == 48
    assert not any(ln.kind == "prelim" and ln.description == "Offices" for ln in plan.lines)  # never priced: empty
    out = str(tmp_path / "filled.xlsx")
    write_plan(blank, out, plan, bwb)
    s = openpyxl.load_workbook(out)["Summary"]
    assert s["E9"].value == "M.G.S. IT LLC" and "2026" in str(s["E11"].value) and "2026" in str(s["E13"].value)


def test_unknown_item_is_not_priced_and_listed(ref, tmp_path):
    def add(book):
        ws = book["Containment Install"]
        ws["B24"], ws["C24"], ws["E24"], ws["F24"] = "Widget frobnicator for the roof", 3, "=C24", "no"
        ws["G24"].fill = copy(ws["G15"].fill)
    blank = _blank(ref, tmp_path, add)
    _, bwb, plan = _plan(ref, blank)
    ln = _line(plan, "Containment Install", 24)
    assert ln.install_rate is None and "NO PRICE" in ln.flags and "not priced, check" in ln.note
    assert ln in plan.not_priced and not ln.writes


def test_web_prices_only_supply_never_install(ref, tmp_path):
    def add(book):
        ws = book["Lighting Install"]
        ws["B26"], ws["C26"], ws["E26"], ws["F26"] = "Widget frobnicator for the roof", 3, "=C26", "Nr"
    blank = _blank(ref, tmp_path, add)
    rwb = read_boq(ref, file_name=REF_NAME)
    meta = RefMeta("ref", REF_NAME, market=rwb.market, client=rwb.client)
    bwb = read_boq(blank, file_name="blank.xlsx")
    asked = []

    def web(row):
        asked.append(row.description)
        return {"unit_price": 12.5, "currency": "EUR", "url": "https://www.elektrika.lv/x", "product": "Widget"}
    plan = Planner(bwb, UnitRateKB(rows_from_workbook(rwb, meta), Want()), Setup(phase_weeks={"Programme": 78}),
                   extras=extras_from_workbook(rwb, meta), web_supply=web).plan()
    ln = _line(plan, "Lighting Install", 26)
    assert ln.supply_rate == 12.5 and ln.install_rate is None and "WEB" in ln.flags
    assert ln.supply.method == "web" and asked == ["Widget frobnicator for the roof"]
    fitting = _line(plan, "Lighting Install", 15)       # free issue: no supply lookup at all
    assert fitting.supply_rate is None


def test_band_rule_marks_check(ref, tmp_path):
    def add(book):
        ws = book["Containment Install"]
        ws["B24"], ws["C24"], ws["E24"], ws["F24"] = "400 x 100mm Cable Ladder", 50, "=C24", "m"
    blank = _blank(ref, tmp_path, add)
    _, _, plan = _plan(ref, blank)
    ln = _line(plan, "Containment Install", 24)
    assert ln.install.method == "band" and "CHECK" in ln.flags and "interpolated from band 300–500" in ln.note


def test_accessory_ratio_off_by_more_than_5_percent_is_checked_not_changed(ref, tmp_path):
    def glands(book):
        book["Lighting Install"]["C22"] = 600   # the reference has 4 glands per luminaire bracket (=C23*4 … )
    blank = _blank(ref, tmp_path, glands)
    _, bwb, plan = _plan(ref, blank)
    assert plan.checks, "a ratio far from the reference must be reported"
    chk = plan.checks[0]
    ln = _line(plan, chk["sheet"], chk["row"])
    assert "CHECK" in ln.flags and "quantity kept as issued" in ln.note
    out = str(tmp_path / "filled.xlsx")
    write_plan(blank, out, plan, bwb)
    assert openpyxl.load_workbook(out)["Lighting Install"]["C22"].value == 600


@pytest.mark.skipif(not shutil.which("soffice") and not os.environ.get("SOFFICE_BIN"),
                    reason="LibreOffice not installed here (it is in the worker image)")
def test_recalculation_check(ref, tmp_path):
    blank = _blank(ref, tmp_path)
    _, bwb, plan = _plan(ref, blank)
    out = str(tmp_path / "filled.xlsx")
    write_plan(blank, out, plan, bwb)
    v = verify(out, plan, bwb)
    assert v["recalculated"] and v["ok"], v["problems"]
    assert v["totals"]["Containment Install"]["total"] > 0


# ------------------------------------------------------------------ chat run path (database)


@pytest.fixture
def _db():
    fx.migrate()
    created: list[str] = []
    yield created
    fx.cleanup_files(created)
    db.execute("DELETE FROM jobs")


def _ingest(path: str, name: str) -> str:
    from app.ingest import pipeline
    fid = fx.uid()
    d = os.path.join(os.environ["DATA_DIR"], "knowledge", fid)
    os.makedirs(d, exist_ok=True)
    dst = os.path.join(d, "original.xlsx")
    shutil.copyfile(path, dst)
    db.execute("""INSERT INTO files(id, original_name, ext, stored_path, tag, size_bytes)
                  VALUES (%s,%s,'xlsx',%s,'reference_estimate',%s)""", (fid, name, dst, os.path.getsize(dst)))
    pipeline.ingest_file({"file_id": fid})
    return fid


def test_run_path_unit_rate_blank(ref, tmp_path, _db):
    from app.agent import run as agent_run
    from app.agent.knowledge import Knowledge
    rid = _ingest(ref, REF_NAME)
    _db.append(rid)
    lv = fx.make_file("Tame_ref_2025.xlsx", items=[{"text": "Kabelis NYM-J 3x1,5 mm2 guldīšana", "unit": "m",
                                                     "norm_h": 0.1, "labour": 1.2, "material": 0.85, "rate": 12}])
    _db.append(lv)
    # the hourly knowledge never contains unit-rate items (no mixing of pricing models)
    kb = Knowledge.load()
    assert rid not in kb.files and not any(i.file_id == rid for i in kb.items) and lv in kb.files

    blank = _blank(ref, tmp_path)
    u = fx.make_user()
    conv = fx.make_conversation(u)
    up = fx.uid()
    d = os.path.join(os.environ["DATA_DIR"], "uploads", up)
    os.makedirs(d, exist_ok=True)
    shutil.copyfile(blank, os.path.join(d, "original.xlsx"))
    db.execute("""INSERT INTO uploads(id, user_id, conversation_id, original_name, ext, stored_path)
                  VALUES (%s,%s,%s,'Goodman blank.xlsx','xlsx',%s)""", (up, u["id"], conv["id"],
                                                                          os.path.join(d, "original.xlsx")))
    run = fx.make_run(conv, allowed=[rid], text="Fill this blank", attachments=[up])
    agent_run.run_agent({"run_id": str(run["id"])})
    card = db.fetchone("SELECT * FROM cards WHERE run_id=%s AND kind='clarify' AND status='pending'", (run["id"],))
    assert card and card["payload"]["purpose"] == "unit_rate_setup"
    ids = {q["id"] for q in card["payload"]["questions"]}
    assert "hourly_rate" not in ids and {"material_supply", "delivery_pct", "offer_date", "validity_date"} <= ids
    pk = next(q for q in card["payload"]["questions"] if q["id"] == "packages")
    assert "Cable Installation" not in pk["suggested"]          # Not participating: pre-unticked
    answers = {"weeks:Programme": 78, "packages": pk["suggested"], "material_supply": "Yes", "delivery_pct": 12,
               "offer_date": "2026-10-02", "validity_date": "2026-12-02", "subcontractor": "M.G.S. IT LLC"}
    db.execute("UPDATE cards SET status='answered', decision=%s WHERE id=%s",
               (db.jsonb({"action": "answer", "data": {"answers": answers}}), card["id"]))
    agent_run.resume_run({"run_id": str(run["id"]), "card_id": str(card["id"])})
    doc = db.fetchone("SELECT * FROM documents WHERE run_id=%s", (run["id"],))
    assert doc and doc["pricing_model"] == "unit_rate"
    rows = db.fetchall("SELECT * FROM estimate_rows WHERE document_id=%s", (doc["id"],))
    ladder = next(r for r in rows if r["sheet_name"] == "Containment Install" and r["row_idx"] == 15)
    assert float(ladder["unit_labour"]) == 25 and ladder["norm_h_per_unit"] is None and ladder["hourly_rate"] is None
    assert ladder["row_kind"] == "item" and ladder["install_method"] in ("exact", "item")
    assert ladder["matched"] and ladder["matched"][0]["file_id"] == rid
    assert doc["totals"]["pricing_model"] == "unit_rate" and doc["totals"]["install"] > 0
    assert db.fetchone("SELECT status FROM runs WHERE id=%s", (run["id"],))["status"] == "done"


def test_run_path_only_hourly_references_asks_for_a_different_pricing_model(ref, tmp_path, _db):
    from app.agent import run as agent_run
    lv = fx.make_file("Tame_ref_2025.xlsx", items=[{"text": "300 x 100mm Cable Ladder", "unit": "m",
                                                     "norm_h": 2, "labour": 24, "material": 0, "rate": 12}])
    _db.append(lv)
    blank = _blank(ref, tmp_path)
    u = fx.make_user()
    conv = fx.make_conversation(u)
    up = fx.uid()
    d = os.path.join(os.environ["DATA_DIR"], "uploads", up)
    os.makedirs(d, exist_ok=True)
    shutil.copyfile(blank, os.path.join(d, "original.xlsx"))
    db.execute("""INSERT INTO uploads(id, user_id, conversation_id, original_name, ext, stored_path)
                  VALUES (%s,%s,%s,'blank.xlsx','xlsx',%s)""", (up, u["id"], conv["id"], os.path.join(d, "original.xlsx")))
    run = fx.make_run(conv, allowed=[lv], text="Fill this blank", attachments=[up])
    rid = str(run["id"])
    agent_run.run_agent({"run_id": rid})
    card = db.fetchone("SELECT * FROM cards WHERE run_id=%s AND kind='clarify' AND status='pending'", (rid,))
    db.execute("UPDATE cards SET status='answered', decision=%s WHERE id=%s",
               (db.jsonb({"action": "answer", "data": {"answers": {"weeks:Programme": 78}}}), card["id"]))
    agent_run.resume_run({"run_id": rid, "card_id": str(card["id"])})
    others = db.fetchall("SELECT id FROM files WHERE pricing_model='unit_rate' AND deleted_at IS NULL "
                         "AND status='analysed'")
    perm = db.fetchall("SELECT * FROM cards WHERE run_id=%s AND kind='permission'", (rid,))
    if others:   # other unit-rate references exist in this database: those are offered first
        assert perm and not perm[0]["payload"].get("cross_model")
        return
    assert perm and perm[0]["payload"]["cross_model"] and perm[0]["payload"]["ref_file_id"] == lv
    assert perm[0]["payload"]["title"] == "Different pricing model"
    db.execute("UPDATE cards SET status='approved', decision=%s WHERE id=%s",
               (db.jsonb({"action": "allow"}), perm[0]["id"]))
    agent_run.resume_run({"run_id": rid, "card_id": str(perm[0]["id"])})
    doc = db.fetchone("SELECT * FROM documents WHERE run_id=%s", (rid,))
    row = db.fetchone("SELECT * FROM estimate_rows WHERE document_id=%s AND sheet_name='Containment Install' "
                      "AND row_idx=15", (doc["id"],))
    assert float(row["unit_labour"]) == 24 and "CHECK" in row["flags"] and "different pricing model" in row["note"]
