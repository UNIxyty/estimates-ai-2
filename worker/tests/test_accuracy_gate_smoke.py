"""Smoke test of the accuracy-gate machinery on synthetic workbooks. This is NOT the gate result:
real numbers come from running app.accuracy_gate on real hand-priced estimates."""
import agent_fixtures as fx  # noqa: I001  (sets env first)

import os
import tempfile

import openpyxl
import pytest
from openpyxl.styles import Font

from app import accuracy_gate, db
from app.llm import client as llm


def _estimate(path: str, rate: float, rows: list[tuple[str, str, float, float, float]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elektro"
    ws["A1"] = "Lokālā tāme"
    ws["A2"] = f"Stundas likme: {rate} EUR/h"
    ws.append([])
    ws.append(["Nr.", "Darba nosaukums", "Mērv.", "Daudz.", "Laika norma c/h", "Darba alga EUR", "Materiāli EUR",
               "Kopā EUR", "Darbietilpība c/h", "Darba alga EUR", "Materiāli EUR", "Kopā EUR"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for i, (text, unit, qty, norm, mat) in enumerate(rows, start=5):
        ws.append([i - 4, text, unit, qty, norm, round(norm * rate, 2), mat, f"=F{i}+G{i}", f"=D{i}*E{i}",
                   f"=D{i}*F{i}", f"=D{i}*G{i}", f"=J{i}+K{i}"])
    wb.save(path)


ITEMS = [("Kabelis NYM-J 3x1,5 mm2 guldīšana", "m", 100, 0.1, 0.85),
         ("Kabelis NYM-J 5x2,5 mm2 guldīšana", "m", 40, 0.14, 2.1),
         ("Kontaktligzda 2-vietīga IP44 montāža", "gab.", 10, 0.5, 9.5),
         ("Slēdzis 1-polīgs zemapmetuma", "gab.", 8, 0.3, 4.2)]


def test_gate_leave_one_out_runs_and_scores(tmp_path):
    fx.migrate()
    object.__setattr__(llm.settings, "llm_enabled", False)
    try:
        d = tmp_path / "refs"
        d.mkdir()
        _estimate(str(d / "a.xlsx"), 12, ITEMS)
        _estimate(str(d / "b.xlsx"), 12, [(t, u, q * 2, n, m * 1.05) for t, u, q, n, m in ITEMS])
        _estimate(str(d / "c.xlsx"), 12, ITEMS[:3])
        out = tmp_path / "report.md"
        assert accuracy_gate.main(["--loo", str(d), "--out", str(out), "--cleanup"]) == 0
        report = out.read_text()
        assert "Labour hours within ±15%" in report and "a.xlsx" in report and "c.xlsx" in report
        assert "## Knowledge ingested" in report and "Median agent/hand ratio" in report
        import json
        res = json.loads((tmp_path / "report.json").read_text())["results"]
        a = next(r for r in res if r["name"] == "a.xlsx")["score"]
        assert a["hours"]["rows"] == 4 and a["hours"]["within_15pct"] >= 3
        assert a["hours"]["median_ratio"] == pytest.approx(1.0, abs=0.05)
        assert a["labour"]["rows"] == 4 and a["labour"]["within_15pct"] >= 3
        assert a["how"]["direct"] >= 3 and a["how"]["model"] == 0
        assert a["share_without_model_call"] == 1.0 and a["cost_usd"] == 0
    finally:
        object.__setattr__(llm.settings, "llm_enabled", True)
        db.execute("DELETE FROM jobs")


def test_gate_knowledge_dir_is_ingested_with_types_and_reported(tmp_path):
    fx.migrate()
    object.__setattr__(llm.settings, "llm_enabled", False)
    try:
        d = tmp_path / "refs"
        d.mkdir()
        k = tmp_path / "knowledge" / "estimates"
        k.mkdir(parents=True)
        _estimate(str(d / "a.xlsx"), 12, ITEMS)
        _estimate(str(d / "b.xlsx"), 12, ITEMS[:2])
        _estimate(str(k / "old_job.xlsx"), 12, ITEMS[2:])
        out = tmp_path / "report.md"
        assert accuracy_gate.main(["--loo", str(d), "--knowledge", str(tmp_path / "knowledge"), "--out", str(out),
                                   "--cleanup"]) == 0
        import json
        data = json.loads((tmp_path / "report.json").read_text())
        kn = {r["original_name"]: r for r in data["knowledge"]}
        assert kn["old_job.xlsx"]["tag"] == "reference_estimate" and kn["old_job.xlsx"]["status"] == "analysed"
        assert kn["old_job.xlsx"]["prices"] == 2
        # a.xlsx held out: its last two items come from the knowledge folder, the first two from b.xlsx
        a = next(r for r in data["results"] if r["name"] == "a.xlsx")["score"]
        assert a["how"]["direct"] == 4
    finally:
        object.__setattr__(llm.settings, "llm_enabled", True)
        db.execute("DELETE FROM jobs")
