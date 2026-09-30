import pytest
from openpyxl import Workbook

from app.ingest.structure import analyse_workbook, classify_header
from fixtures.make_fixtures import build_all


@pytest.fixture(scope="module")
def files(tmp_path_factory):
    return build_all(str(tmp_path_factory.getbasetemp() / "ingest_fixtures"))


@pytest.fixture(scope="module")
def lv(files):
    return analyse_workbook(files["lv"])


def _cols(sheet):
    return {c.letter: (c.meaning, c.source) for c in sheet.columns}


def test_classify_header():
    assert classify_header("Darba nosaukums") == ("item", None)
    assert classify_header("Daudz.") == ("qty", None)
    assert classify_header("Mērv.") == ("unit", None)
    assert classify_header("Materialer i alt") == ("material", "total")
    assert classify_header("Darba alga uz vienību") == ("labour", "unit")
    assert classify_header("Timepris")[0] == "hourly_rate"
    assert classify_header("Kilde")[0] == "source"
    assert classify_header("I alt") == ("total", None)


def test_lv_workbook_language_currency(lv):
    assert lv.language == "LV"
    assert lv.currency == "EUR"
    assert [s.kind for s in lv.sheets] == ["estimate", "estimate", "summary"]
    assert lv.model_calls == 0 and lv.model_cells == 0


def test_lv_header_and_unlabelled_unit_block(lv):
    s = lv.sheets[0]
    assert s.header_rows == (5, 6) and s.header_row == 6
    assert s.first_data_row == 8                      # numbering row 7 skipped
    cols = _cols(s)
    assert cols["A"][0] == "no" and cols["B"][0] == "item" and cols["C"][0] == "unit" and cols["D"][0] == "qty"
    # unlabelled columns under the merged "Vienības izmaksas": found by position + numeric relationships
    assert cols["E"] == ("norm_h", "position")
    assert cols["F"] == ("hourly_rate", "position")
    assert cols["G"] == ("unit_labour", "position")
    assert cols["H"] == ("unit_material", "position")
    assert cols["I"] == ("unit_mechanisms", "position")
    assert cols["J"] == ("unit_total", "position")
    assert s.unit_block == {"norm_h": "E", "hourly_rate": "F", "unit_labour": "G", "unit_material": "H",
                            "unit_mechanisms": "I", "unit_total": "J"}
    assert s.total_block == {"total_norm_h": "K", "total_labour": "L", "total_material": "M",
                             "total_mechanisms": "N", "total": "O"}
    assert all(c.source == "header" for c in s.columns if c.letter in "KLMNO")
    assert s.col("unit_labour").confidence >= 0.9    # verified: total_labour == qty * unit_labour


def test_lv_rows_sections_subtotals_markups(lv):
    s = lv.sheets[0]
    kinds = {r.row: r.kind for r in s.rows}
    assert kinds[8] == "section" and kinds[14] == "section"
    assert kinds[13] == "subtotal" and kinds[19] == "subtotal"
    assert kinds[20] == "total" and kinds[25] == "total"
    assert kinds[27] == "note"
    assert [r.row for r in s.items] == [9, 10, 11, 12, 15, 16, 17, 18]
    assert [(x.title, x.row_start, x.row_end, x.subtotal_row) for x in s.sections] == [
        ("Kabeļi un caurules", 8, 13, 13), ("Elektroierīces", 14, 19, 19)]
    r15 = next(r for r in s.rows if r.row == 15)
    assert r15.values["unit_material"] == pytest.approx(6.8)        # "6,80" stored as text
    assert r15.cells["unit_material"] == "H15" and r15.formulas["total_material"] is True
    assert r15.section_title == "Elektroierīces"
    names = [m["key"] for m in s.markups]
    assert names == ["virsizdevumi", "pelna", "darba_deveja_socialais_nodoklis", "transporta_izdevumi"]
    transport = s.markups[-1]
    assert transport["pct"] == pytest.approx(0.03) and transport["pct_cell"] == "N24"
    assert s.totals["total"] == pytest.approx(1894.87)


def test_hourly_rates_per_sheet_and_section(lv):
    s1, s2 = lv.sheets[0], lv.sheets[1]
    assert [r["rate"] for r in s1.hourly_rates] == [12.5]
    assert {x.hourly_rate for x in s1.sections} == {12.5}
    assert [r["rate"] for r in s2.hourly_rates] == [14.0, 15.0]
    assert [(x.title, x.hourly_rate) for x in s2.sections] == [("Datu tīkls un signalizācija", 14.0),
                                                               ("Zemējums", 15.0)]


def test_electrical_detection(lv):
    assert lv.sheets[0].is_electrical is True and lv.sheets[0].electrical_confidence >= 0.6
    assert lv.sheets[1].is_electrical is True


def test_non_electrical_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Būvdarbi"
    ws.append(["Nr.", "Darba nosaukums", "Mērv.", "Daudz.", "Cena", "Summa"])
    for i, t in enumerate(["Betona grīdas izbūve", "Sienu apmetums", "Krāsošana 2 kārtās", "Flīzēšana",
                           "Jumta seguma montāža"], start=1):
        ws.append([i, t, "m2", 10 * i, 5.0, 50.0 * i])
    p = tmp_path / "build.xlsx"
    wb.save(p)
    s = analyse_workbook(str(p)).sheets[0]
    assert s.kind == "estimate"
    assert s.is_electrical is False and s.electrical_confidence <= 0.25


def test_blank_uses_header_and_position_only(files):
    ws = analyse_workbook(files["lv_blank"])
    for s in ws.sheets[:2]:
        cols = _cols(s)
        assert [cols[x][0] for x in "EFGHIJ"] == ["norm_h", "hourly_rate", "unit_labour", "unit_material",
                                                  "unit_mechanisms", "unit_total"]
        assert all(cols[x][1] == "position" for x in "EFGHIJ")
        assert all(s.col(m).confidence <= 0.6 for m in ("norm_h", "unit_labour"))
        assert len(s.items) == (8 if s.idx == 0 else 6)
        assert all(r.qty for r in s.items)
        assert all(r.values.get("unit_labour") is None for r in s.items)


def test_danish_estimate_with_kilde(files):
    ws = analyse_workbook(files["da"])
    assert ws.language == "DA" and ws.currency == "DKK"
    s = ws.sheets[0]
    cols = _cols(s)
    assert cols["E"][0] == "norm_h"                  # "Timer" per unit
    assert cols["F"][0] == "unit_material"           # generic "Materialer" resolved by numbers
    assert cols["G"][0] == "hourly_rate"
    assert cols["H"][0] == "total_labour" and cols["I"][0] == "total_material" and cols["J"][0] == "total"
    assert cols["K"][0] == "source"
    assert s.items[0].source_ref == "Solar prisliste 2024"
    assert [x.title for x in s.sections] == ["Stærkstrøm", "Svagstrøm"]
    kinds = {r.row: r.kind for r in s.rows}
    assert kinds[14] == "total"                       # "I alt" below several sections is the sheet total
    assert [m["key"] for m in s.markups] == ["avance", "moms"]
    assert s.hourly_rates[0]["rate"] == 450


def test_norms_file(files):
    ws = analyse_workbook(files["norms"], tag="hourly_norms")
    assert [s.kind for s in ws.sheets] == ["norms", "norms"]
    s = ws.sheets[0]
    assert _cols(s)["A"][0] == "category" and _cols(s)["D"][0] == "norm_h"


def test_docx_and_pdf_tables(files):
    d = analyse_workbook(files["docx"])
    s = d.sheets[0]
    assert s.kind == "estimate" and len(s.items) == 3
    assert s.items[1].values["unit_total"] == pytest.approx(11.18)
    if files["pdf"]:
        p = analyse_workbook(files["pdf"])
        assert p.language == "EN" and len(p.sheets[0].items) == 3


def _ambiguous(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Nr.", "Darba nosaukums", "Mērv.", "Daudz.", "X1", "X2"])
    for i in range(1, 6):
        ws.append([i, f"Kabelis NYM 3x{1.5 if i % 2 else 2.5}", "m", 10 * i, 0.5 + i, 7 + i * 3])
    p = tmp_path / "amb.xlsx"
    wb.save(p)
    return str(p)


def test_llm_only_for_ambiguous_columns(tmp_path):
    path = _ambiguous(tmp_path)
    calls = []

    def fake_llm(prompt, schema):
        calls.append(prompt)
        return {"columns": [{"letter": "E", "meaning": "unit_material", "confidence": 0.8},
                            {"letter": "F", "meaning": "unit_labour", "confidence": 0.7},
                            {"letter": "B", "meaning": "qty"}]}          # B is decided by code: ignored
    ws = analyse_workbook(path, llm=fake_llm)
    s = ws.sheets[0]
    assert len(calls) == 1 and ws.model_calls == 1 and ws.model_cells > 0 and ws.model_rows > 0
    cols = _cols(s)
    assert cols["E"] == ("unit_material", "model") and cols["F"] == ("unit_labour", "model")
    assert cols["B"][0] == "item"
    # without a model: deterministic, low confidence
    ws2 = analyse_workbook(path)
    s2 = ws2.sheets[0]
    assert ws2.model_calls == 0
    assert all(c.confidence <= 0.5 for c in s2.columns if c.letter in "EF")


def test_llm_failure_is_harmless(tmp_path):
    def boom(prompt, schema):
        raise RuntimeError("model down")
    ws = analyse_workbook(_ambiguous(tmp_path), llm=boom)
    assert ws.model_calls == 1 and ws.sheets[0].items


def test_messy_workbook(tmp_path):
    """Header below a long title block, hidden rows, numbers stored as text, formatted empty trailing rows."""
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Tāme 2"
    for i in range(1, 12):
        ws.cell(row=i, column=1, value=f"Pasūtītājs / objekts / adrese, rinda {i}")
    hdr = ["Nr.p.k.", "Darbu un materiālu nosaukums", "Mērvienība", "Daudzums", "Vienības cena, EUR",
           "Kopā, EUR", "Piezīmes"]
    for c, h in enumerate(hdr, start=1):
        ws.cell(row=14, column=c, value=h)
    data = [
        ("1", "Kabelis NYM 3x2,5", "tek.m", "1 250", "2,10", "2 625,00", ""),
        ("2", "Kontaktligzda IP44", "gab", "12", "8,50", "102,00", "virsapmetuma"),
        ("3", "Vecais gaismeklis (neizmanto)", "gab", "3", "10,00", "30,00", ""),
        ("4", "Automātslēdzis C16 1P", "gab.", "6", "4,20", "25,20", ""),
    ]
    for r, row in enumerate(data, start=15):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v or None)
    ws.row_dimensions[17].hidden = True
    ws.cell(row=19, column=2, value="Kopā").font = Font(bold=True)
    ws.cell(row=19, column=6, value="2 782,20")
    for r in range(20, 400):                        # formatted but empty
        ws.cell(row=r, column=6).number_format = "0.00"
    p = tmp_path / "messy.xlsx"
    wb.save(p)
    s = analyse_workbook(str(p)).sheets[0]
    assert s.header_row == 14 and s.first_data_row == 15
    cols = {c.letter: c.meaning for c in s.columns}
    assert cols == {"A": "no", "B": "item", "C": "unit", "D": "qty", "E": "unit_total", "F": "total", "G": "notes"}
    items = s.items
    assert [r.row for r in items] == [15, 16, 17, 18]
    assert items[0].qty == 1250 and items[0].values["unit_total"] == 2.1 and items[0].values["total"] == 2625
    assert items[2].hidden is True
    assert s.last_data_row == 19 and s.rows[-1].kind == "total"
