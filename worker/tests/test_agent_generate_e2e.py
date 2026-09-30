"""Generate from a work list: structure card first, generation only after Generate; vague list → clarify
card; Use a different reference swaps the template. Deterministic path (no model)."""
import agent_fixtures as fx  # noqa: I001  (sets env first)

import os

import openpyxl
import pytest
from openpyxl.styles import Font

from app import db
from app.agent import generate
from app.agent import run as agent_run
from app.llm import client as llm


@pytest.fixture(autouse=True)
def _env():
    fx.migrate()
    object.__setattr__(llm.settings, "llm_enabled", False)
    created: list[str] = []
    yield created
    object.__setattr__(llm.settings, "llm_enabled", True)
    fx.cleanup_files(created)
    db.execute("DELETE FROM jobs")


def make_reference(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elektro"
    ws["A1"] = "Lokālā tāme"
    ws["A2"] = "Stundas likme: 12,00 EUR/h"
    ws.append([])
    ws.append(["Nr.", "Darba nosaukums", "Mērv.", "Daudz.", "Laika norma c/h", "Darba alga EUR", "Materiāli EUR",
               "Kopā EUR", "Darbietilpība c/h", "Darba alga EUR", "Materiāli EUR", "Kopā EUR"])
    for c in ws[4]:
        c.font = Font(bold=True)
    ws.append([None, "Kabeļu līnijas"])
    ws["B5"].font = Font(bold=True)
    ws.append([1, "Kabelis NYM-J 3x1,5 mm2 guldīšana", "m", 100, 0.1, 1.2, 0.85, "=F6+G6", "=D6*E6", "=D6*F6", "=D6*G6", "=J6+K6"])
    ws.append([2, "Kabelis NYM-J 5x2,5 mm2 guldīšana", "m", 50, 0.14, 1.68, 2.1, "=F7+G7", "=D7*E7", "=D7*F7", "=D7*G7", "=J7+K7"])
    ws.append([None, "Kopā", None, None, None, None, None, None, "=SUM(I6:I7)", "=SUM(J6:J7)", "=SUM(K6:K7)", "=SUM(L6:L7)"])
    ws.append([None, "Ierīces"])
    ws["B9"].font = Font(bold=True)
    ws.append([3, "Kontaktligzda 2-vietīga IP44 montāža", "gab.", 10, 0.5, 6.0, 9.5, "=F10+G10", "=D10*E10", "=D10*F10", "=D10*G10", "=J10+K10"])
    ws.append([None, "Kopā", None, None, None, None, None, None, "=SUM(I10:I10)", "=SUM(J10:J10)", "=SUM(K10:K10)", "=SUM(L10:L10)"])
    ws.append([None, "Kopā tāmē", None, None, None, None, None, None, None, "=J8+J11", "=K8+K11", "=L8+L11"])
    ws.append([None, "Sociālais nodoklis 23,59%", None, None, None, None, None, None, None, "=J12*0.2359"])
    ws.column_dimensions["B"].width = 50
    wb.save(path)


def _ref_file(name: str) -> str:
    d = os.path.join(os.environ["DATA_DIR"], "knowledge", fx.uid())
    os.makedirs(d)
    path = os.path.join(d, "original.xlsx")
    make_reference(path)
    fid = fx.make_file(name, items=[
        {"row": 6, "text": "Kabelis NYM-J 3x1,5 mm2 guldīšana", "unit": "m", "norm_h": 0.1, "labour": 1.2, "material": 0.85, "rate": 12},
        {"row": 7, "text": "Kabelis NYM-J 5x2,5 mm2 guldīšana", "unit": "m", "norm_h": 0.14, "labour": 1.68, "material": 2.1, "rate": 12},
        {"row": 10, "text": "Kontaktligzda 2-vietīga IP44 montāža", "unit": "gab.", "norm_h": 0.5, "labour": 6.0, "material": 9.5, "rate": 12},
    ])
    db.execute("UPDATE files SET stored_path=%s WHERE id=%s", (path, fid))
    return fid


def test_parse_work_list_lines():
    items, unparsed = generate.parse_work_list(
        "Kabeļi:\n- Kabelis NYM-J 3x1,5 — 120 m\n- 24 gab. kontaktligzda 2-vietīga IP44\nRozetes\n")
    assert items[0] == {"text": "Kabelis NYM-J 3x1,5", "qty": 120.0, "unit": "m", "section": "Kabeļi"}
    assert items[1]["qty"] == 24.0 and items[1]["unit"].startswith("gab")
    assert items[2]["qty"] is None and unparsed == 1


def _decide(card_id: str, action: str, data: dict | None = None) -> None:
    import json
    db.execute("UPDATE cards SET decision=%s::jsonb, decided_at=now() WHERE id=%s",
               (json.dumps({"action": action, "data": data or {}}), card_id))


def test_generate_waits_for_structure_then_builds_from_template(_env):
    ref = _ref_file("Tame_ref.xlsx")
    ref2 = _ref_file("Tame_other.xlsx")
    _env.extend([ref, ref2])
    u = fx.make_user()
    conv = fx.make_conversation(u)
    run = fx.make_run(conv, allowed=[ref, ref2], text=(
        "Sagatavo tāmi:\n1. Kabelis NYM-J 3x1,5 mm2 guldīšana — 200 m\n2. Kabelis NYM-J 5x2,5 mm2 guldīšana — 60 m\n"
        "3. Kontaktligzda 2-vietīga IP44 montāža — 18 gab.\n"))
    rid = str(run["id"])
    agent_run.run_agent({"run_id": rid})
    assert db.fetchone("SELECT status FROM runs WHERE id=%s", (rid,))["status"] == "waiting"
    assert db.fetchone("SELECT COUNT(*) AS n FROM documents WHERE run_id=%s", (rid,))["n"] == 0  # nothing yet
    card = db.fetchone("SELECT * FROM cards WHERE run_id=%s AND kind='structure'", (rid,))
    p = card["payload"]
    assert p["language"] == "LV" and p["sheets"][0]["name"] == "Elektro"
    assert {s["title"] for s in p["sheets"][0]["sections"]} >= {"Kabeļu līnijas", "Ierīces"}
    assert any(c["meaning"] == "qty" for c in p["columns"]) and p["alternatives"]

    # Use a different reference → old card replaced, new proposal on the other template
    _decide(str(card["id"]), "use_reference", {"file_id": p["alternatives"][0]["file_id"]})
    agent_run.resume_run({"run_id": rid, "card_id": str(card["id"])})
    assert db.fetchone("SELECT status FROM cards WHERE id=%s", (card["id"],))["status"] == "replaced"
    card2 = db.fetchone("""SELECT * FROM cards WHERE run_id=%s AND kind='structure' AND status='pending'""", (rid,))
    assert card2["payload"]["template_file_id"] == p["alternatives"][0]["file_id"]

    # Generate
    _decide(str(card2["id"]), "generate")
    agent_run.resume_run({"run_id": rid, "card_id": str(card2["id"])})
    assert db.fetchone("SELECT status FROM runs WHERE id=%s", (rid,))["status"] == "done"
    assert db.fetchone("SELECT status FROM cards WHERE id=%s", (card2["id"],))["status"] == "done"
    doc = db.fetchone("SELECT * FROM documents WHERE run_id=%s", (rid,))
    assert doc["mode"] == "generate" and doc["language"] == "LV"
    rows = db.fetchall("SELECT * FROM estimate_rows WHERE document_id=%s ORDER BY row_idx", (doc["id"],))
    assert len(rows) == 3 and all(r["price_source"] == "exact" for r in rows)
    ws = openpyxl.load_workbook(doc["stored_path"])["Elektro"]
    assert ws.column_dimensions["B"].width == 50                       # template formatting kept
    texts = [ws.cell(r, 2).value for r in range(5, ws.max_row + 1)]
    assert "Kabeļu līnijas" in texts and "Ierīces" in texts
    item_row = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == "Kabelis NYM-J 3x1,5 mm2 guldīšana")
    assert ws.cell(item_row, 4).value == 200 and ws.cell(item_row, 10).value == f"=D{item_row}*F{item_row}"
    assert ws.cell(item_row, 6).value == 1.2                            # value cell filled
    assert any(isinstance(ws.cell(r, 10).value, str) and "0.2359" in ws.cell(r, 10).value
               for r in range(5, ws.max_row + 1))                        # markup row carried over


def test_vague_list_asks_clarifying_questions(_env):
    ref = _ref_file("Tame_ref3.xlsx")
    _env.append(ref)
    u = fx.make_user()
    conv = fx.make_conversation(u)
    run = fx.make_run(conv, allowed=[ref], text="Uztaisi tāmi dzīvoklim:\n- kabeļi\n- rozetes\n- sadale\n")
    agent_run.run_agent({"run_id": str(run["id"])})
    card = db.fetchone("SELECT * FROM cards WHERE run_id=%s", (run["id"],))
    assert card["kind"] == "clarify" and card["payload"]["questions"]
    assert db.fetchone("SELECT COUNT(*) AS n FROM cards WHERE run_id=%s AND kind='structure'", (run["id"],))["n"] == 0
