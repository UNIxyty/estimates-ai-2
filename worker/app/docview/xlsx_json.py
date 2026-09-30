"""xlsx -> compact row/cell JSON for the virtualised DocViewer grid.

The whole workbook is converted once (formulas + cached values + styles + structure), stored
gzip-compressed under ``${DATA_DIR}/cache/view/<key>.json.gz`` and kept in a small in-process LRU,
so page requests are a list slice.

Model (cached)::

    {"v": 1, "locale": "en"|"eu", "sheets": [
      {"name", "idx", "state", "row_count", "col_count", "col_widths": [px...], "default_row_h": px,
       "frozen": {"rows", "cols"}, "merges": [{"r1","c1","r2","c2"}], "hidden_rows": [r...],
       "hidden_cols": [c...],
       "rows": [{"r": 1, "h": px|null, "kind": "item|section|subtotal|total|header|note|blank",
                 "cells": [[c, value, display, flags(, formula)] ...], "hidden"?: true}]}]}

``flags`` is a compact string: letters ``b`` bold, ``i`` italic, ``u`` underline, ``f`` formula,
``m`` merged-range origin, ``r``/``c`` right/centre aligned (numbers default to ``r`` like Excel),
``w`` wrap; optionally followed by ``;bg=RRGGBB`` (solid fill) and ``;fg=RRGGBB`` (font colour).
Example: ``"bfr;bg=FFF2CC"``. Formula cells carry the formula text as a 5th element.
"""
from __future__ import annotations

import bisect
import colorsys
import datetime as _dt
import importlib
import importlib.util
import logging
import os
import re
import zipfile
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import coordinate_to_tuple

from . import cache
from .formula import DEFERRED, Evaluator
from .numfmt import format_value, guess_locale

log = logging.getLogger(__name__)

MODEL_VERSION = 1
MAX_LIMIT = 500
MAX_MATCHES = 500
FILTERS = ("all", "web", "check", "no_price", "edited", "flagged")
KINDS = ("item", "section", "subtotal", "total", "header", "note", "blank")

# ------------------------------------------------------------------ colours

_THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5",
                "accent6", "hlink", "folHlink"]
_OFFICE_THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000",
                 "5B9BD5", "70AD47", "0563C1", "954F72"]
_THEME_RE = re.compile(
    r"<a:(dk1|lt1|dk2|lt2|accent[1-6]|hlink|folHlink)>\s*<a:(?:srgbClr val|sysClr[^>]*?lastClr)=\"([0-9A-Fa-f]{6})\"")


def _theme_colors(wb) -> list[str]:
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return list(_OFFICE_THEME)
    text = raw.decode("utf-8", "replace") if isinstance(raw, bytes) else str(raw)
    found = {name: val.upper() for name, val in _THEME_RE.findall(text)}
    return [found.get(n, d) for n, d in zip(_THEME_ORDER, _OFFICE_THEME)]


def _tint(hex6: str, tint: float) -> str:
    if not tint:
        return hex6
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


def _color_hex(color, theme: list[str]) -> str | None:
    if color is None:
        return None
    try:
        t = color.type
        if t == "rgb":
            v = color.rgb
            if not isinstance(v, str):
                return None
            v = v.upper()
            if len(v) == 8:
                if v[:2] == "00" and v != "00000000":
                    pass  # alpha is ignored by Excel for fills
                v = v[2:]
            return v if len(v) == 6 else None
        if t == "theme":
            idx = int(color.theme)
            if 0 <= idx < len(theme):
                return _tint(theme[idx], float(color.tint or 0))
            return None
        if t == "indexed":
            idx = int(color.indexed)
            if 0 <= idx < len(COLOR_INDEX) and idx < 64:
                return _tint(COLOR_INDEX[idx][2:].upper(), float(color.tint or 0))
    except (AttributeError, TypeError, ValueError):
        return None
    return None


# ------------------------------------------------------------------ styles


class _StyleInfo:
    __slots__ = ("letters", "tail", "bold", "fmt", "halign", "fill")

    def __init__(self, letters: str, tail: str, bold: bool, fmt: str, halign: str | None, fill: bool):
        self.letters, self.tail, self.bold, self.fmt, self.halign, self.fill = letters, tail, bold, fmt, halign, fill


def _style_info(cell, theme: list[str]) -> _StyleInfo:
    font = cell.font
    al = cell.alignment
    letters = ""
    bold = bool(font is not None and font.b)
    if bold:
        letters += "b"
    if font is not None and font.i:
        letters += "i"
    if font is not None and font.u:
        letters += "u"
    tail = ""
    fill = cell.fill
    has_fill = False
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        bg = _color_hex(fill.fgColor, theme)
        if bg and bg != "FFFFFF":
            tail += f";bg={bg}"
            has_fill = True
    if font is not None and font.color is not None:
        fg = _color_hex(font.color, theme)
        if fg and fg != "000000":
            tail += f";fg={fg}"
    h = al.horizontal if al is not None else None
    halign = "r" if h == "right" else "c" if h in ("center", "centerContinuous") else "l" if h in (
        "left", "justify", "distributed", "fill") else None
    if al is not None and al.wrap_text:
        tail = "w" + tail  # keep letters before ';'
    return _StyleInfo(letters, tail, bold, cell.number_format or "General", halign, has_fill)


def _json_value(v: Any) -> Any:
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return None
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat()
    if isinstance(v, _dt.timedelta):
        return v.total_seconds()
    return str(v)


def _formula_text(raw: Any) -> str | None:
    if isinstance(raw, str) and raw.startswith("=") and len(raw) > 1:
        return raw
    text = getattr(raw, "text", None)  # ArrayFormula / DataTableFormula
    if isinstance(text, str):
        return text if text.startswith("=") else "=" + text
    return None


# ------------------------------------------------------------------ row kinds

_TOTAL_RE = re.compile(r"(?i)\b(pavisam|grand\s+total|kopā\s+ar\s+pvn|total\s+incl|i\s+alt\s+inkl)")
_SUBTOTAL_RE = re.compile(
    r"(?i)(?<!\w)(kopā|kopa|starpsumma|sub-?total|total|i\s+alt|mellemsum|summa\s+kopā|kopsumma)(?!\w)")
_HEADER_WORDS = re.compile(
    r"(?i)(nosaukum|apraksts|daudz|mērv|merv|vien|cena|summa|\bnr\b|\bpos|name|description|qty|"
    r"quantity|unit|price|amount|betegnelse|antal|enhed|pris|darba|materiā|labou?r|material|norm)")


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def heuristic_kinds(rows: list[dict], bold_rows: set[int], fill_rows: set[int]) -> None:
    """Assign ``kind`` in place using light rules (used when ingest structure is unavailable)."""
    last_item_i = -1
    header_done = False
    for i, row in enumerate(rows):
        cells = row["cells"]
        vals = [c for c in cells if c[1] is not None and c[1] != ""]
        if not vals:
            row["kind"] = "blank"
            continue
        texts = [c for c in vals if isinstance(c[1], str) and c[1].strip()]
        nums = [c for c in vals if _is_num(c[1])]
        # A leading position number ("1", "2.1") next to a title does not make a row priced.
        if nums and texts and vals[0] is nums[0] and nums[0][0] < texts[0][0] and len(nums) == 1:
            nums = []
        text_join = " ".join(c[1] for c in texts)
        r = row["r"]
        kind = None
        if not header_done and r <= 40 and len(texts) >= 3 and not nums \
                and len(_HEADER_WORDS.findall(text_join)) >= 2:
            kind = "header"
            header_done = True
        elif texts and _SUBTOTAL_RE.search(text_join) and (nums or len(texts) <= 2 or any("f" in c[3].split(";")[0] for c in vals)):
            kind = "total" if _TOTAL_RE.search(text_join) else "subtotal"
        elif not nums and texts and len(texts) <= 3 and (r in bold_rows or r in fill_rows):
            kind = "section"
        elif nums:
            kind = "item"
        else:
            kind = "note"
        row["kind"] = kind
        if kind == "item":
            last_item_i = i
    # Title / meta lines above the column header are notes, not sections.
    header_i = next((i for i, row in enumerate(rows) if row["kind"] == "header"), -1)
    for row in rows[:max(header_i, 0)]:
        if row["kind"] != "blank":
            row["kind"] = "note"
    # With no explicit grand total, the last subtotal-like row after the final item is the total.
    if last_item_i >= 0:
        tail = rows[last_item_i + 1:]
        if not any(row["kind"] == "total" for row in tail):
            for row in tail:
                if row["kind"] == "subtotal":
                    row["kind"] = "total"
                    break


def _ingest_kinds(path: str) -> dict[str, dict[int, str]] | None:
    """Row kinds from app.ingest.structure.analyse_workbook, when that module is available."""
    if os.environ.get("DOCVIEW_USE_INGEST", "1") == "0":
        return None
    try:
        mod = importlib.import_module("app.ingest.structure")
        analyse = getattr(mod, "analyse_workbook")
    except Exception:  # noqa: BLE001 - optional dependency, written concurrently
        return None
    try:
        res = analyse(path)
    except Exception as e:  # noqa: BLE001
        log.warning("docview: analyse_workbook failed for %s: %s", path, e)
        return None

    def g(o: Any, *names: str) -> Any:
        for n in names:
            if isinstance(o, dict) and n in o:
                return o[n]
            if hasattr(o, n):
                return getattr(o, n)
        return None

    sheets = g(res, "sheets")
    if sheets is None and isinstance(res, (list, tuple)):
        sheets = res
    if not sheets:
        return None
    out: dict[str, dict[int, str]] = {}
    try:
        for s in sheets:
            name = g(s, "name", "sheet_name", "title")
            rows = g(s, "rows") or []
            if name is None:
                continue
            m: dict[int, str] = {}
            for ri in rows:
                r, k = g(ri, "row", "r"), g(ri, "kind")
                k = getattr(k, "value", k)  # enums
                if isinstance(r, int) and isinstance(k, str) and k in KINDS:
                    m[r] = k
            out[str(name)] = m
    except Exception as e:  # noqa: BLE001
        log.warning("docview: unexpected analyse_workbook shape: %s", e)
        return None
    return out or None


# ------------------------------------------------------------------ conversion

def _load(path: str, **kw):
    try:
        return openpyxl.load_workbook(path, **kw)
    except zipfile.BadZipFile:
        raise cache.Unreadable("not a valid xlsx file (bad zip)") from None
    except (KeyError, ValueError, TypeError, OSError, IndexError, AttributeError) as e:
        raise cache.Unreadable(f"cannot read workbook: {type(e).__name__}: {e}"[:300]) from None
    except Exception as e:  # noqa: BLE001 - openpyxl raises many types on corrupt files
        raise cache.Unreadable(f"cannot read workbook: {type(e).__name__}: {e}"[:300]) from None


def convert_workbook(path: str) -> dict:
    wb = _load(path, data_only=False, rich_text=False, keep_links=False)
    theme = _theme_colors(wb)
    worksheets = list(wb.worksheets)

    # Pass 1: raw cells of every sheet (needed for the formula fallback across sheets).
    raw_cells: dict[str, dict[tuple[int, int], Any]] = {}
    formula_pos: dict[str, set[tuple[int, int]]] = {}
    for ws in worksheets:
        cells = {}
        fpos = set()
        for key, cell in getattr(ws, "_cells", {}).items():
            if isinstance(cell, MergedCell):
                continue
            v = cell._value if hasattr(cell, "_value") else cell.value
            if v is None:
                continue
            f = _formula_text(v)
            if f is not None:
                fpos.add(key)
                v = f
            cells[key] = v
        raw_cells[ws.title] = cells
        formula_pos[ws.title] = fpos

    # Pass 2: cached values of formula cells (read-only is fine here: no merges/styles needed).
    cached: dict[str, dict[tuple[int, int], Any]] = {name: {} for name in raw_cells}
    if any(formula_pos.values()):
        vwb = _load(path, data_only=True, read_only=True, keep_links=False)
        try:
            for vws in vwb.worksheets:
                fpos = formula_pos.get(vws.title)
                if not fpos:
                    continue
                max_r = max(r for r, _ in fpos)
                out = cached[vws.title]
                try:
                    vws.reset_dimensions()
                except AttributeError:
                    pass
                for row in vws.iter_rows(max_row=max_r):
                    for c in row:
                        rr = getattr(c, "row", None)
                        if rr is None:
                            continue
                        k = (rr, c.column)
                        if k in fpos and c.value is not None:
                            out[k] = c.value
        finally:
            vwb.close()
    evaluator = Evaluator(raw_cells, cached)
    # Evaluate uncached formulas top-to-bottom so running totals (=H9+G10 ...) resolve with shallow
    # recursion: each one's predecessor is already memoised.
    pending = [(title, key) for title, fpos in formula_pos.items() for key in sorted(fpos)
               if key not in cached[title]]
    for _ in range(4):  # cells whose chain was too deep are retried once their predecessors are known
        pending = [(t, k) for t, k in pending if evaluator.try_value(t, k[0], k[1]) is DEFERRED]
        if not pending:
            break

    formats: set[str] = set()
    style_cache: dict[tuple, _StyleInfo] = {}
    sheets_out = []
    kinds_from_ingest = None
    ingest_tried = False

    # Collect formats first so the locale is known before rendering displays.
    try:
        formats.update(str(f) for f in wb._number_formats)
    except (AttributeError, TypeError):
        pass
    loc = guess_locale(formats)

    for idx, ws in enumerate(worksheets):
        title = ws.title
        merges = []
        merge_origins = set()
        max_r = max_c = 0
        for mr in ws.merged_cells.ranges:
            merges.append({"r1": mr.min_row, "c1": mr.min_col, "r2": mr.max_row, "c2": mr.max_col})
            merge_origins.add((mr.min_row, mr.min_col))
        rows_cells: dict[int, list] = {}
        bold_rows: set[int] = set()
        fill_rows: set[int] = set()
        fill_only: list[tuple[int, int, list]] = []
        fpos = formula_pos[title]
        for (r, c), cell in sorted(ws._cells.items()):
            if isinstance(cell, MergedCell):
                continue
            try:
                skey = tuple(cell._style)
            except (AttributeError, TypeError):
                skey = None
            si = style_cache.get(skey) if skey is not None else None
            if si is None:
                si = _style_info(cell, theme)
                if skey is not None:
                    style_cache[skey] = si
            raw = raw_cells[title].get((r, c))
            is_f = (r, c) in fpos
            value = evaluator.value(title, r, c) if is_f else raw
            is_origin = (r, c) in merge_origins
            if value is None and not si.fill and not is_origin:
                continue
            letters = si.letters
            if is_f:
                letters += "f"
            if is_origin:
                letters += "m"
            halign = si.halign
            if halign is None and (_is_num(value) or isinstance(value, (_dt.date, _dt.time))):
                halign = "r"
            if halign in ("r", "c"):
                letters += halign
            flags = letters + si.tail
            display = format_value(value, si.fmt, loc).strip() if value is not None else ""
            entry = [c, _json_value(value), display, flags]
            if is_f:
                entry.append(raw)
            if value is None:
                fill_only.append((r, c, entry))
                continue
            rows_cells.setdefault(r, []).append(entry)
            if r > max_r:
                max_r = r
            if c > max_c:
                max_c = c
            if si.bold and isinstance(value, str):
                bold_rows.add(r)
            if si.fill:
                fill_rows.add(r)
        for m in merges:
            max_r = max(max_r, m["r2"])
            max_c = max(max_c, m["c2"])
        for r, c, entry in fill_only:  # styled-but-empty cells only inside the used area
            if r <= max_r and c <= max_c:
                lst = rows_cells.setdefault(r, [])
                lst.append(entry)
                if len(lst) > 1 and lst[-2][0] > c:
                    lst.sort(key=lambda e: e[0])
                if "bg=" in entry[3]:
                    fill_rows.add(r)

        # geometry
        fmt = ws.sheet_format
        default_w = fmt.defaultColWidth or ((fmt.baseColWidth or 8) + 0.43)
        default_px = int(round(default_w * 7 + 5))
        col_widths = [default_px] * max_c
        hidden_cols: set[int] = set()
        for dim in ws.column_dimensions.values():
            lo, hi = dim.min or 0, dim.max or 0
            if not lo:
                try:
                    lo = hi = openpyxl.utils.column_index_from_string(dim.index)
                except (ValueError, AttributeError):
                    continue
            for c in range(lo, min(hi, max_c) + 1):
                if dim.width and dim.customWidth is not False:
                    col_widths[c - 1] = int(round(dim.width * 7 + 5))
                if dim.hidden:
                    hidden_cols.add(c)
        default_row_h = int(round((fmt.defaultRowHeight or 15) * 4 / 3))
        heights: dict[int, int] = {}
        hidden_rows: set[int] = set()
        for r, dim in ws.row_dimensions.items():
            if r > max_r:
                continue
            if dim.ht is not None and (dim.customHeight or dim.ht):
                heights[r] = int(round(dim.ht * 4 / 3))
            if dim.hidden:
                hidden_rows.add(r)
        frozen = {"rows": 0, "cols": 0}
        if ws.freeze_panes:
            try:
                fr, fc = coordinate_to_tuple(ws.freeze_panes)
                frozen = {"rows": fr - 1, "cols": fc - 1}
            except (ValueError, TypeError):
                pass

        rows = []
        for r in range(1, max_r + 1):
            row = {"r": r, "h": heights.get(r), "kind": "blank", "cells": rows_cells.get(r, [])}
            if r in hidden_rows:
                row["hidden"] = True
            rows.append(row)

        heuristic_kinds(rows, bold_rows, fill_rows)
        if not ingest_tried:
            ingest_tried = True
            kinds_from_ingest = _ingest_kinds(path)
        if kinds_from_ingest and title in kinds_from_ingest:
            km = kinds_from_ingest[title]
            for row in rows:
                k = km.get(row["r"])
                if k and not (k == "blank" and row["cells"] and any(c[1] not in (None, "") for c in row["cells"])):
                    row["kind"] = k

        sheets_out.append({
            "name": title, "idx": idx, "state": ws.sheet_state or "visible",
            "row_count": max_r, "col_count": max_c, "col_widths": col_widths,
            "default_row_h": default_row_h, "frozen": frozen, "merges": merges,
            "hidden_rows": sorted(hidden_rows), "hidden_cols": sorted(hidden_cols),
            "kind_source": "ingest" if kinds_from_ingest and title in kinds_from_ingest else "heuristic",
            "rows": rows,
        })
    try:
        wb.close()
    except Exception:  # noqa: BLE001
        pass
    return {"v": MODEL_VERSION, "locale": loc, "sheets": sheets_out}


# ------------------------------------------------------------------ cached access

class Workbook:
    """A loaded conversion plus per-sheet lowercase search text (built once per process)."""

    __slots__ = ("model", "search")

    def __init__(self, model: dict) -> None:
        self.model = model
        self.search = [
            ["\x1f".join(c[2].lower() for c in row["cells"] if c[2]) for row in s["rows"]]
            for s in model["sheets"]
        ]


def _ingest_available() -> bool:
    if os.environ.get("DOCVIEW_USE_INGEST", "1") == "0":
        return False
    try:
        return importlib.util.find_spec("app.ingest.structure") is not None
    except (ImportError, ValueError):
        return False


def load(path: str) -> Workbook:
    # Whether ingest structure was used is part of the key, so kinds refresh once it ships.
    tag = f"xlsx-v{MODEL_VERSION}-{'i' if _ingest_available() else 'h'}"
    key = cache.cache_key(tag, path)
    return cache.get_or_build(key, ".json.gz", lambda: convert_workbook(path),
                              cache.gzip_json_dump, cache.gzip_json_load, Workbook)


def sheet_meta(sheet: dict) -> dict:
    return {k: v for k, v in sheet.items() if k != "rows"}


def workbook_meta(wb: Workbook) -> dict:
    return {"locale": wb.model["locale"], "sheets": [sheet_meta(s) for s in wb.model["sheets"]]}


def find_sheet(wb: Workbook, sheet: str | int | None) -> int | None:
    sheets = wb.model["sheets"]
    if sheet is None or sheet == "":
        return 0 if sheets else None
    try:
        i = int(sheet)
        return i if 0 <= i < len(sheets) else None
    except (TypeError, ValueError):
        for s in sheets:
            if s["name"] == sheet:
                return s["idx"]
        low = str(sheet).lower()
        for s in sheets:
            if s["name"].lower() == low:
                return s["idx"]
    return None


def _marker_pred(filter_: str):
    if filter_ == "flagged":
        return lambda m: m is not None and m.get("marker") is not None
    want = {"web": "WEB", "check": "CHECK", "no_price": "NO PRICE", "edited": "EDITED"}[filter_]
    return lambda m: m is not None and want in m.get("_eff", ())


def query_rows(wb: Workbook, sheet_idx: int, *, offset: int = 0, limit: int = 200, filter_: str = "all",
               q: str | None = None, around: int | None = None,
               markers: dict[int, dict] | None = None) -> dict:
    sheet = wb.model["sheets"][sheet_idx]
    rows = sheet["rows"]
    limit = max(1, min(int(limit), MAX_LIMIT))
    offset = max(0, int(offset))
    q = (q or "").strip().lower()
    filter_ = filter_ or "all"

    idxs: list[int] | None = None  # None = all rows, unfiltered
    if filter_ != "all":
        pred = _marker_pred(filter_)
        mk = markers or {}
        idxs = [i for i, row in enumerate(rows) if pred(mk.get(row["r"]))]
    if q:
        text = wb.search[sheet_idx]
        base = idxs if idxs is not None else range(len(rows))
        idxs = [i for i in base if q in text[i]]

    total = len(rows) if idxs is None else len(idxs)
    if around is not None:
        if idxs is None:
            pos = bisect.bisect_left([row["r"] for row in rows], int(around)) if rows else 0
        else:
            pos = bisect.bisect_left([rows[i]["r"] for i in idxs], int(around))
        pos = min(pos, max(total - 1, 0))
        offset = (pos // limit) * limit
    page_idx = range(offset, min(offset + limit, total)) if idxs is None else idxs[offset:offset + limit]

    out_rows = []
    for i in page_idx:
        row = rows[i]
        if markers is not None:
            m = markers.get(row["r"])
            if m is not None:
                row = dict(row)
                row.update({k: v for k, v in m.items() if not k.startswith("_")})
        out_rows.append(row)
    res: dict[str, Any] = {"sheet": sheet_idx, "offset": offset, "limit": limit, "total": total, "rows": out_rows}
    if q:
        res["matches"] = [rows[i]["r"] for i in (idxs or [])[:MAX_MATCHES]]
    return res
