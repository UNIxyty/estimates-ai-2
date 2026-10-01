"""Generate from a work list:
  1. read the list (deterministic line parser first; model only for free prose)
  2. too vague → clarifying questions card instead of guessing
  3. pick the closest reference as template (deterministic ranking)
  4. post the structure proposal card (sheets, sections, columns, pricing logic, language)
  5. generate only after Generate; Change structure / Use a different reference revise the proposal
"""
from __future__ import annotations

import json
import os
import re

from .. import db
from ..ingest.language import detect_language
from ..ingest.structure import analyse_workbook
from ..llm import client as llm
from . import builder, cards, common
from .context import RunContext, RunWaiting
from .knowledge import lexical_similarity, tokens
from .pricing import RowSpec
from .writer import layout_from_structure

_QTY_UNIT = r"(\d+(?:[.,]\d+)?)\s*(m2|m²|m3|m|gab\.?|gb|pcs?|stk\.?|kompl\.?|kpl|set|sets|ea|each|obj\.?|lot|kg|l)\b\.?"
_TRAIL = re.compile(rf"^(?P<text>.+?)[\s,;:–—-]+{_QTY_UNIT}\s*$", re.I)
_LEAD = re.compile(rf"^{_QTY_UNIT}\s*[x×]?\s+(?P<text>.+)$", re.I)
_BULLET = re.compile(r"^\s*(?:\d+[.)]|[-•*])\s*")
# A line like "Sagatavo tāmi:" / "Please make an estimate for:" is the request, not a section heading.
_REQUEST = re.compile(r"\b(sagatavo|izveido|uztaisi|aprēķini|generate|create|make|build|prepare|price|lav|opret|"
                      r"beregn|tām|estimate|budget|tilbud|overslag|please|lūdzu)\w*", re.I)


def parse_work_list(text: str) -> tuple[list[dict], int]:
    """Returns (items, unparsed_line_count). Lines ending with ':' become section headings."""
    items, unparsed, section = [], 0, None
    for raw in text.splitlines():
        line = _BULLET.sub("", raw).strip()
        if not line:
            continue
        if line.endswith(":") and len(line) < 80:
            if not _REQUEST.search(line):
                section = line[:-1].strip()
            continue
        m = _TRAIL.match(line) or _LEAD.match(line)
        if m:
            qty = float(m.group(1 if m.re is _LEAD else 2).replace(",", "."))
            unit = m.group(2 if m.re is _LEAD else 3)
            items.append({"text": m.group("text").strip(" -–—,;:"), "qty": qty, "unit": unit, "section": section})
        else:
            items.append({"text": line, "qty": None, "unit": None, "section": section})
            unparsed += 1
    return items, unparsed


EXTRACT_SCHEMA = {
    "type": "object",
    "properties": {
        "items": {"type": "array", "items": {"type": "object", "properties": {
            "text": {"type": "string"}, "qty": {"type": ["number", "null"]}, "unit": {"type": ["string", "null"]},
            "section": {"type": ["string", "null"]}}, "required": ["text"]}},
        "vague": {"type": "boolean"},
        "questions": {"type": "array", "items": {"type": "object", "properties": {
            "text": {"type": "string"}, "options": {"type": "array", "items": {"type": "string"}}},
            "required": ["text"]}},
    },
    "required": ["items", "vague", "questions"],
}
EXTRACT_SYSTEM = (
    "You turn an electrician's work description into estimate line items (text, qty, unit, section). "
    "Keep the user's wording and language for item texts; never invent quantities. If the description is "
    "too vague to price (missing quantities, sizes, cable types, counts), set vague=true and ask at most 5 "
    "short, concrete questions (with options when useful)."
)


def _work_text(rc: RunContext, attachments: list[dict], message_text: str) -> str:
    parts = [message_text or ""]
    for a in attachments:
        p = a["stored_path"]
        ext = a["ext"].lower()
        try:
            if ext in ("txt", "csv"):
                parts.append(open(p, encoding="utf-8", errors="replace").read())
            elif ext == "docx":
                import docx
                d = docx.Document(p)
                parts.append("\n".join(par.text for par in d.paragraphs))
                for t in d.tables:
                    for row in t.rows:
                        parts.append(" — ".join(c.text.strip() for c in row.cells if c.text.strip()))
            elif ext == "pdf":
                import pdfplumber
                with pdfplumber.open(p) as pdf:
                    parts.append("\n".join((pg.extract_text() or "") for pg in pdf.pages[:30]))
        except Exception as e:  # noqa: BLE001
            rc.emitter.step_warn("read", f"Couldn't read {a['original_name']}: {e}")
    return "\n".join(x for x in parts if x).strip()


def _rank_templates(rc: RunContext, items: list[dict], lang: str | None) -> list[tuple[str, float]]:
    refs = [fid for fid in rc.allowed if rc.kb.files.get(fid, {}).get("tag") == "reference_estimate"]
    toks = [tokens(i["text"]) for i in items]
    scores = []
    for fid in refs:
        fitems = [it for it in rc.kb.items if it.file_id == fid]
        if not fitems:
            continue
        cover = sum(max((lexical_similarity(t, it.toks) for it in fitems), default=0) for t in toks) / max(len(toks), 1)
        bonus = 0.1 if lang and rc.kb.language_of(fid) == lang else 0
        scores.append((fid, round(cover + bonus, 4)))
    return sorted(scores, key=lambda t: -t[1])


def _template_sheet(path: str):
    st = analyse_workbook(path)
    sheets = [s for s in st.sheets if any(r.kind == "item" for r in s.rows)]
    sheets.sort(key=lambda s: (not s.is_electrical, -sum(1 for r in s.rows if r.kind == "item")))
    return st, (sheets[0] if sheets else None)


def _group_sections(items: list[dict], sheet) -> list[dict]:
    """Keep the user's own sections if given; else map each item to the template section whose rows are
    most similar; unmatched items go to a final 'Other works' section."""
    if any(i.get("section") for i in items):
        out: dict[str, list] = {}
        for i in items:
            out.setdefault(i.get("section") or "Other works", []).append(i)
        return [{"title": k, "items": v} for k, v in out.items()]
    sec_rows: dict[str, list] = {}
    for r in sheet.rows:
        if r.kind == "item" and getattr(r, "section_title", None):
            sec_rows.setdefault(r.section_title, []).append(tokens(r.text or ""))
    out: dict[str, list] = {t: [] for t in sec_rows}
    other = []
    for i in items:
        t = tokens(i["text"])
        best = max(((title, max((lexical_similarity(t, x) for x in rows), default=0)) for title, rows in sec_rows.items()),
                   key=lambda p: p[1], default=(None, 0))
        (out[best[0]] if best[0] and best[1] >= 0.35 else other).append(i)
    secs = [{"title": k, "items": v} for k, v in out.items() if v]
    if other:
        secs.append({"title": "Other works", "items": other})
    return secs


def _proposal(rc: RunContext, template_id: str, items: list[dict], ranked: list[tuple[str, float]],
              language: str | None = None) -> dict:
    f = db.fetchone("SELECT * FROM files WHERE id=%s", (template_id,))
    path = f["work_path"] or f["stored_path"]
    _, sheet = _template_sheet(path)
    lang = language or f["language"]
    logic = db.fetchall("""SELECT COALESCE(override_sentence, sentence) AS s FROM file_logic
                           WHERE file_id=%s ORDER BY ord LIMIT 8""", (template_id,))
    return {
        "template_file_id": template_id, "template_name": f["original_name"], "language": lang,
        "sheets": [{"name": sheet.name if sheet else "Estimate", "sections": _group_sections(items, sheet)}],
        "columns": [{"letter": c.letter, "header": c.header, "meaning": c.meaning} for c in (sheet.columns if sheet else [])
                    if c.meaning != "unknown"],
        "pricing_logic": [r["s"] for r in logic],
        "alternatives": [{"file_id": fid, "name": rc.kb.file_name(fid), "score": s} for fid, s in ranked
                         if fid != template_id][:5],
        "item_count": len(items),
    }


REVISE_SCHEMA = {"type": "object", "properties": {"sections": {"type": "array", "items": {"type": "object", "properties": {
    "title": {"type": "string"}, "items": {"type": "array", "items": {"type": "object", "properties": {
        "text": {"type": "string"}, "qty": {"type": ["number", "null"]}, "unit": {"type": ["string", "null"]}},
        "required": ["text"]}}}, "required": ["title", "items"]}},
    "language": {"type": ["string", "null"]}}, "required": ["sections"]}


def run(rc: RunContext, attachments: list[dict], message_text: str, card: dict | None = None) -> None:
    st = rc.state.setdefault("generate", {})

    # --- resume paths --------------------------------------------------------------
    if card is not None and card["kind"] == "clarify":
        answers = cards.decision_data(card).get("answers") or {}
        st["answers"] = answers
        st.pop("items", None)
    if card is not None and card["kind"] == "structure":
        action = (card.get("decision") or {}).get("action")
        data = cards.decision_data(card)
        prop = card["payload"]
        if action == "use_reference":
            fid = data.get("file_id")
            if fid not in rc.allowed:
                raise ValueError("reference not allowed for this chat")
            items = [i for s in prop["sheets"][0]["sections"] for i in s["items"]]
            new = _proposal(rc, fid, items, _rank_templates(rc, items, None))
            cards.update(str(card["id"]), status="replaced", expect_status=("pending", "replaced"))
            _post_proposal(rc, new, f"Switched the template to {new['template_name']}.")
            raise RunWaiting()
        if action == "change":
            instructions = data.get("instructions", "")
            rc.emitter.step_started("revise", "Revising the structure")
            out = llm.complete_json("generate", "You revise the section structure of an electrical estimate. "
                                    "Keep every item; move, rename, split or merge sections as asked.",
                                    f"Current structure:\n{json.dumps(prop['sheets'][0]['sections'], ensure_ascii=False)}"
                                    f"\n\nUser's change request:\n{instructions}", REVISE_SCHEMA, ctx=rc.ctx,
                                    forced_tier=rc.forced_tier)
            new = {**prop, "sheets": [{**prop["sheets"][0], "sections": out["sections"]}],
                   "language": out.get("language") or prop["language"]}
            rc.emitter.step_done("revise")
            cards.update(str(card["id"]), status="changed", expect_status=("pending", "changed"))
            _post_proposal(rc, new, "Here's the revised structure.")
            raise RunWaiting()
        if action == "generate":
            cards.update(str(card["id"]), status="generating", expect_status=("pending", "generating"))
            lang = data.get("language")
            if isinstance(lang, str) and re.fullmatch(r"[A-Za-z]{2}", lang):  # Language selector on the card
                prop = {**prop, "language": lang.upper()}
            _generate(rc, prop, card)
            cards.update(str(card["id"]), status="done")
            if rc.state.get("pending"):
                raise RunWaiting()
            return

    # --- 1. read the work list --------------------------------------------------------
    if "items" not in st:
        rc.emitter.step_started("read", "Reading the work list")
        text = _work_text(rc, attachments, message_text)
        if st.get("answers"):
            text += "\n\nAnswers to your questions:\n" + "\n".join(f"- {k}: {v}" for k, v in st["answers"].items())
        items, unparsed = parse_work_list(text)
        need_model = not items or unparsed > max(2, len(items) // 3)
        questions: list[dict] = []
        vague = False
        if need_model and llm.available():
            out = llm.complete_json("generate", EXTRACT_SYSTEM, text[:30000], EXTRACT_SCHEMA, ctx=rc.ctx,
                                    forced_tier=rc.forced_tier)
            items, vague, questions = out.get("items") or [], bool(out.get("vague")), out.get("questions") or []
        missing = [i for i in items if i.get("qty") is None]
        if not items or len(missing) > max(1, len(items) // 3):
            vague = True
            if not questions:
                questions = [{"text": f"What quantity (and unit) for “{i['text'][:70]}”?"} for i in missing[:5]] or \
                            [{"text": "Could you list the works with quantities, e.g. “NYM 3×1.5 cable — 120 m”?"}]
        rc.emitter.step_done("read", f"{len(items)} item(s)")
        if vague and not st.get("answers"):
            msg = rc.new_assistant_message()
            common.tools.ask_clarifying_questions(rc, questions, message_id=str(msg["id"]))
            rc.complete_message(str(msg["id"]), content="The list is too vague to price reliably, so I have a few "
                                                        "questions first.")
            rc.save_state(generate=st)
            raise RunWaiting()
        st["items"] = items
        st["lang"] = detect_language([i["text"] for i in items]) if items else None
        rc.save_state(generate=st)

    # --- 2. template + 3. proposal ----------------------------------------------------
    rc.emitter.step_started("template", "Choosing the closest reference as template")
    ranked = _rank_templates(rc, st["items"], st.get("lang"))
    if not ranked:
        msg = rc.new_assistant_message()
        rc.complete_message(str(msg["id"]), content="None of the references available to this chat is an analysed "
                                                    "reference estimate I can use as a template.")
        return
    template_id = ranked[0][0]
    rc.emitter.step_done("template", rc.kb.file_name(template_id))
    _post_proposal(rc, _proposal(rc, template_id, st["items"], ranked),
                   f"I'd build this on {rc.kb.file_name(template_id)}. Check the structure, then press Generate.")
    raise RunWaiting()


def _post_proposal(rc: RunContext, proposal: dict, text: str) -> None:
    msg = rc.new_assistant_message()
    common.tools.propose_structure(rc, proposal, message_id=str(msg["id"]))
    rc.complete_message(str(msg["id"]), content=text)


TRANSLATE_SCHEMA = {"type": "object", "properties": {"texts": {"type": "array", "items": {"type": "string"}}},
                    "required": ["texts"]}


def _generate(rc: RunContext, prop: dict, card: dict) -> None:
    f = db.fetchone("SELECT * FROM files WHERE id=%s AND deleted_at IS NULL", (prop["template_file_id"],))
    if not f or str(f["id"]) not in rc.allowed:
        raise ValueError("template file is no longer available")
    path = f["work_path"] or f["stored_path"]
    _, sheet = _template_sheet(path)
    sections = prop["sheets"][0]["sections"]
    lang = prop.get("language") or f["language"]
    items = [i for s in sections for i in s["items"]]
    src_lang = detect_language([i["text"] for i in items]) if items else lang
    if lang and src_lang and lang != src_lang and llm.available():
        rc.emitter.step_started("translate", f"Writing item texts in {lang}")
        out = llm.complete_json("generate", f"Translate electrical estimate item texts into {lang}. Keep sizes, "
                                "types and brand names exactly. Same order, same count.",
                                json.dumps([i["text"] for i in items], ensure_ascii=False), TRANSLATE_SCHEMA,
                                ctx=rc.ctx, forced_tier=rc.forced_tier)
        texts = out.get("texts") or []
        if len(texts) == len(items):
            for i, t in zip(items, texts):
                i["source_text"], i["text"] = i["text"], t
        rc.emitter.step_done("translate")
    rc.emitter.step_started("build", f"Building the workbook from {f['original_name']}")
    layout = layout_from_structure(sheet)
    out_path = os.path.join(common.settings.data_dir, "documents", "build", f"{rc.run_id}.xlsx")
    built = builder.build(path, out_path, sheet=sheet, sections=sections, lang=lang, cols=layout.cols)
    rc.emitter.step_done("build", f"{len(built)} rows in {len(sections)} section(s)")
    specs = [RowSpec(sheet=sheet.name, row=b.row, text=b.text, unit=b.unit, qty=b.qty, section=b.section)
             for b in built]
    rate = next((float(h["rate"]) for h in (sheet.hourly_rates or []) if isinstance(h, dict) and h.get("rate")), None)
    rc.save_state(lang=lang, target_rate=rate)
    priced, stats = common.price_with_events(rc, specs, task="generate", target_rate=rate)
    rc.check()
    base = os.path.splitext(f["original_name"])[0]
    common.finish_priced(rc, priced=priced, stats=stats, layouts={sheet.name: layout}, src_path=out_path,
                         name=f"New estimate (from {base}).xlsx", mode="generate", lang=lang,
                         currency=sheet.currency or "EUR", sheets=[sheet.name], template_file_id=str(f["id"]))
