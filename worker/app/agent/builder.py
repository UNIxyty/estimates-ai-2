"""Build a new estimate workbook from a reference estimate's own structure and formatting.

Keeps the template sheet's header rows (styles, merges, column widths, print setup) untouched, clears
the data area, then writes sections / items / subtotals using the template's own row styles, and
re-creates per-row formulas by translating the template item row's formulas to each new row. Rows below
the template's data area (grand total, markups) are carried over with their formulas shifted."""
from __future__ import annotations

import copy
import os
import shutil
from dataclasses import dataclass

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from .writer import TOTAL_MEANINGS

SUBTOTAL_WORD = {"LV": "Kopā", "EN": "Subtotal", "DA": "I alt", "NO": "Sum", "SV": "Summa", "LT": "Iš viso",
                 "ET": "Kokku", "DE": "Zwischensumme", "FI": "Yhteensä"}
TOTAL_WORD = {"LV": "Kopā tāmē", "EN": "Total", "DA": "Total", "NO": "Totalt", "SV": "Totalt", "LT": "Iš viso sąmatoje",
              "ET": "Kokku", "DE": "Gesamt", "FI": "Yhteensä"}


@dataclass
class BuiltRow:
    row: int
    text: str
    unit: str | None
    qty: float | None
    section: str


def _row_style(ws, r: int, max_col: int) -> list[dict]:
    out = []
    for c in range(1, max_col + 1):
        cell = ws.cell(r, c)
        out.append({"font": copy.copy(cell.font), "border": copy.copy(cell.border), "fill": copy.copy(cell.fill),
                    "number_format": cell.number_format, "alignment": copy.copy(cell.alignment),
                    "protection": copy.copy(cell.protection), "value": cell.value})
    return out


def _apply_style(ws, r: int, style: list[dict]) -> None:
    for c, s in enumerate(style, start=1):
        cell = ws.cell(r, c)
        cell.font, cell.border, cell.fill = s["font"], s["border"], s["fill"]
        cell.number_format, cell.alignment, cell.protection = s["number_format"], s["alignment"], s["protection"]


def build(template_path: str, out_path: str, *, sheet, sections: list[dict], lang: str | None,
          cols: dict[str, str]) -> list[BuiltRow]:
    """sheet: ingest SheetStructure of the template sheet; sections: [{title, items:[{text,unit,qty}]}];
    cols: meaning → letter. Returns the item rows written (for pricing)."""
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    shutil.copyfile(template_path, out_path)
    wb = openpyxl.load_workbook(out_path)
    for name in list(wb.sheetnames):
        if name != sheet.name:
            del wb[name]
    ws = wb[sheet.name]
    max_col = max(ws.max_column, 1)
    first = sheet.first_data_row or ((sheet.header_row or 1) + 1)
    last = sheet.last_data_row or ws.max_row

    kinds = {r.row: r.kind for r in sheet.rows}
    sample = {k: next((r for r in range(first, last + 1) if kinds.get(r) == k), None)
              for k in ("section", "item", "subtotal")}
    item_style = _row_style(ws, sample["item"] or first, max_col)
    section_style = _row_style(ws, sample["section"], max_col) if sample["section"] else item_style
    subtotal_style = _row_style(ws, sample["subtotal"], max_col) if sample["subtotal"] else section_style
    item_formulas = {c: s["value"] for c, s in enumerate(item_style, start=1)
                     if isinstance(s["value"], str) and s["value"].startswith("=")}
    section_merges = [m for m in ws.merged_cells.ranges if sample["section"] and m.min_row == m.max_row == sample["section"]]
    tail_first = last + 1
    tail = [(_row_style(ws, r, max_col), ws.row_dimensions[r].height) for r in range(tail_first, ws.max_row + 1)]
    tail_merges = [m for m in ws.merged_cells.ranges if m.min_row >= tail_first]

    # clear data area + tail (unmerge first: openpyxl doesn't shift merged ranges on delete)
    for m in list(ws.merged_cells.ranges):
        if m.max_row >= first:
            ws.unmerge_cells(str(m))
    ws.delete_rows(first, ws.max_row - first + 1)

    item_col = cols.get("item", "B")
    no_col, unit_col, qty_col = cols.get("no"), cols.get("unit"), cols.get("qty")
    total_cols = [cols[m] for m in TOTAL_MEANINGS if m in cols]
    r = first
    built: list[BuiltRow] = []
    subtotal_rows: list[int] = []
    n = 0
    for sec in sections:
        _apply_style(ws, r, section_style)
        ws[f"{item_col}{r}"] = sec["title"]
        for m in section_merges:
            ws.merge_cells(start_row=r, start_column=m.min_col, end_row=r, end_column=m.max_col)
        r += 1
        start = r
        for it in sec.get("items", []):
            n += 1
            _apply_style(ws, r, item_style)
            if no_col:
                ws[f"{no_col}{r}"] = n
            ws[f"{item_col}{r}"] = it["text"]
            if unit_col:
                ws[f"{unit_col}{r}"] = it.get("unit")
            if qty_col and it.get("qty") is not None:
                ws[f"{qty_col}{r}"] = it["qty"]
            for c, f in item_formulas.items():
                ws.cell(r, c).value = Translator(f, origin=f"{get_column_letter(c)}{sample['item'] or first}") \
                    .translate_formula(f"{get_column_letter(c)}{r}")
            built.append(BuiltRow(r, it["text"], it.get("unit"), it.get("qty"), sec["title"]))
            r += 1
        _apply_style(ws, r, subtotal_style)
        ws[f"{item_col}{r}"] = f"{SUBTOTAL_WORD.get((lang or 'EN').upper(), 'Subtotal')}: {sec['title']}"
        for tc in total_cols:
            ws[f"{tc}{r}"] = f"=SUM({tc}{start}:{tc}{r - 1})" if r > start else 0
        subtotal_rows.append(r)
        r += 1

    # grand total + carried-over tail (markups etc.), formulas shifted by the row offset
    _apply_style(ws, r, subtotal_style)
    ws[f"{item_col}{r}"] = TOTAL_WORD.get((lang or "EN").upper(), "Total")
    for tc in total_cols:
        ws[f"{tc}{r}"] = "=" + "+".join(f"{tc}{s}" for s in subtotal_rows) if subtotal_rows else 0
    grand = r
    offset = (grand + 1) - tail_first
    for i, (style, height) in enumerate(tail):
        rr = grand + 1 + i
        _apply_style(ws, rr, style)
        for c, s in enumerate(style, start=1):
            v = s["value"]
            if isinstance(v, str) and v.startswith("="):
                v = Translator(v, origin=f"{get_column_letter(c)}{tail_first + i}").translate_formula(
                    f"{get_column_letter(c)}{rr}")
            ws.cell(rr, c).value = v
        if height:
            ws.row_dimensions[rr].height = height
    for m in tail_merges:
        ws.merge_cells(start_row=m.min_row + offset, start_column=m.min_col, end_row=m.max_row + offset,
                       end_column=m.max_col)
    try:
        wb.calculation.fullCalcOnLoad = True
    except AttributeError:
        pass
    wb.save(out_path)
    return built
