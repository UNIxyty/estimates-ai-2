"""Unit-rate accuracy gate: blank each hand-priced EU BOQ, fill it from the others, compare.

Blanking clears what M.G.S. typed in and the client would not have: item install / supply rates (and totals that are
typed in rather than formulas), contractor-item rates and months, lump sums, delivery amounts, prelim rates and
weeks, and the Summary's subcontractor / offer date / validity cells. Quantities stay.

The setup card is answered the way the estimator would (never with prices): programme weeks per phase from the
original's contractor months, prelims weeks from its Site Manager line, delivery % as typed in, material supply in
scope when the original priced supply, today's offer date + 2 months validity.

Metrics: install €/unit exact and within ±10 %, supply within ±15 %, sheet and Summary totals within ±5 %, the share
of lines priced without a model call, cost per run, and how many lines were CHECK / NO PRICE / interpolated.
"""
from __future__ import annotations

import os
import re
import shutil
import statistics
from dataclasses import dataclass, field
from datetime import date, timedelta

import openpyxl

from .boq import BoqWorkbook, read_boq
from .fill import WEEKS_PER_MONTH, Setup, _phase_no

INSTALL_TOL, SUPPLY_TOL, TOTAL_TOL = 0.10, 0.15, 0.05


def _clear(ws, coord: str | None) -> bool:
    if not coord:
        return False
    cell = ws[coord]
    if isinstance(cell.value, str) and cell.value.startswith("="):
        return False
    try:
        cell.value = None
    except AttributeError:  # merged non-anchor
        return False
    return True


def make_blank(src: str, dst: str, wb: BoqWorkbook | None = None) -> BoqWorkbook:
    """Copy src to dst with every M.G.S. input cleared. Returns the parsed original (the truth)."""
    wb = wb or read_boq(src, file_name=os.path.basename(src))
    shutil.copyfile(src, dst)
    book = openpyxl.load_workbook(dst)
    for sh in wb.pricing_sheets:
        ws = book[sh.name]
        for r in sh.rows:
            for role in ("rate", "supply", "total", "supply_total"):
                _clear(ws, r.cells.get(role))
            if r.kind == "contractor" and r.rate_basis == "monthly":
                _clear(ws, r.cells.get("qty"))
            if r.kind == "delivery" and r.uom_norm == "%":
                _clear(ws, r.cells.get("qty"))
            if r.kind in ("lump_sum", "delivery"):
                for ph in sh.phases:
                    if ph.get("money_col"):
                        _clear(ws, f"{ph['money_col']}{r.row}")
            if r.kind == "item":
                for ph in sh.phases:
                    if ph.get("money_col"):
                        _clear(ws, f"{ph['money_col']}{r.row}")
    for p in wb.prelims:
        ws = book[p["sheet"]]
        _clear(ws, p["cells"].get("rate"))
        _clear(ws, p["cells"].get("total"))
        if p["rate_basis"] == "weekly":
            _clear(ws, p["cells"].get("qty"))
    if "Summary" in book.sheetnames:
        for v in (wb.subcontractor or {}).values():
            _clear(book["Summary"], v.get("cell"))
    book.save(dst)
    return wb


def setup_for(truth: BoqWorkbook, today: date | None = None) -> Setup:
    """What the estimator would answer on the setup card for this job (no prices)."""
    months: dict[int | None, float] = {}
    per_sheet: dict[str, float] = {}
    for sh in truth.pricing_sheets:
        for r in sh.rows:
            if r.kind == "contractor" and r.rate_basis == "monthly" and r.qty:
                ph = _phase_no(r.description) or _phase_no(r.phase)
                months[ph] = max(months.get(ph, 0), r.qty)
                key = f"{sh.name} · Phase {ph}" if ph else sh.name
                per_sheet[key] = max(per_sheet.get(key, 0), r.qty)
    phase_weeks = {(f"Phase {k}" if k else "Programme"): round(v * WEEKS_PER_MONTH, 2) for k, v in sorted(
        months.items(), key=lambda kv: kv[0] or 0)}
    # a package that runs shorter or longer than the programme gets its own duration (Goodman: containment 18
    # months, lighting 12; FRA3H: external 1 month, internal 6)
    sheet_weeks = {}
    for key, m in per_sheet.items():
        ph = _phase_no(key.split(" · ")[-1]) if " · " in key else None
        overall = months.get(ph)
        if overall is not None and abs(overall - m) > 1e-9:
            sheet_weeks[key] = round(m * WEEKS_PER_MONTH, 2)
    sm = [p["qty"] for p in truth.prelims if p["rate_basis"] == "weekly" and p["qty"]
          and re.search(r"site\s*manager", p["description"], re.I)]
    from .analysis import delivery_info
    pct = [d["typed_percent"] or d["computed_percent"] for d in delivery_info(truth)
           if d.get("typed_percent") or d.get("computed_percent")]
    has_supply = any(r.supply_rate for sh in truth.pricing_sheets for r in sh.rows if r.kind == "item")
    today = today or date.today()
    return Setup(phase_weeks=phase_weeks, sheet_weeks=sheet_weeks, prelims_weeks=max(sm) if sm else None,
                 participating={ln["label"]: not ln.get("not_participating") for ln in truth.summary_lines},
                 material_supply=has_supply, delivery_pct=round(statistics.median(pct), 2) if pct else 12.0,
                 offer_date=today.isoformat(), validity_date=(today + timedelta(days=61)).isoformat())


@dataclass
class Score:
    name: str
    install_n: int = 0
    install_exact: int = 0
    install_tol: int = 0
    supply_n: int = 0
    supply_tol: int = 0
    lines: int = 0
    check: int = 0
    no_price: int = 0
    interpolated: int = 0
    model_free: int = 0
    sheets: list[dict] = field(default_factory=list)       # [{sheet, hand, agent, ok}]
    summary: list[dict] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    cost_usd: float = 0.0
    seconds: float = 0.0
    misses: list[dict] = field(default_factory=list)


def score_plan(name: str, truth: BoqWorkbook, plan, verified: dict | None = None) -> Score:
    s = Score(name=name)
    by = {(ln.sheet, ln.row): ln for ln in plan.lines}
    for sh in truth.pricing_sheets:
        for r in sh.rows:
            if r.kind != "item":
                continue
            ln = by.get((sh.name, r.row))
            if r.install_rate and r.rate_basis != "supply_only":
                s.install_n += 1
                got = ln.install_rate if ln else None
                if got is not None and abs(got - r.install_rate) < 1e-9:
                    s.install_exact += 1
                if got is not None and abs(got - r.install_rate) <= INSTALL_TOL * r.install_rate:
                    s.install_tol += 1
                elif len(s.misses) < 400:
                    s.misses.append({"sheet": sh.name, "row": r.row, "what": "install", "text": r.description[:70],
                                     "hand": r.install_rate, "agent": got,
                                     "method": ln.install.method if ln and ln.install else None})
            if r.supply_rate:
                s.supply_n += 1
                got = ln.supply_rate if ln else None
                if got is not None and abs(got - r.supply_rate) <= SUPPLY_TOL * r.supply_rate:
                    s.supply_tol += 1
    for ln in plan.lines:
        s.lines += 1
        s.model_free += 1           # unit-rate filling makes no model call
        s.check += "CHECK" in ln.flags
        s.no_price += "NO PRICE" in ln.flags
        s.interpolated += bool(ln.install and ln.install.method == "band")
    # sheet totals: the agent's recalculated totals (or the plan's) against the original's cached totals
    vt = (verified or {}).get("totals") or {}
    for sh in truth.pricing_sheets:
        hand = sh.totals.get("total")
        if not hand:
            continue
        agent = vt.get(sh.name, {}).get("total")
        if agent is None:
            pt = plan.sheet_totals.get(sh.name) or {}
            agent = pt.get("total")
        ok = agent is not None and abs(agent - hand) <= TOTAL_TOL * abs(hand)
        s.sheets.append({"sheet": sh.name, "hand": round(hand, 2), "agent": agent, "ok": ok})
    hand_summary = {ln["label"]: ln for ln in truth.summary_lines}
    for label, agent in ((verified or {}).get("summary") or {}).items():
        h = hand_summary.get(label, {}).get("value")
        if h:
            s.summary.append({"label": label, "hand": h, "agent": agent, "ok": abs(agent - h) <= TOTAL_TOL * abs(h)})
    s.problems = list((verified or {}).get("problems") or [])
    return s


def summary_values(path: str, truth: BoqWorkbook) -> None:
    """Attach the original's cached Summary values to truth.summary_lines (for the ±5 % Summary check)."""
    try:
        v = openpyxl.load_workbook(path, data_only=True)
    except Exception:  # noqa: BLE001
        return
    if "Summary" not in v.sheetnames:
        return
    for ln in truth.summary_lines:
        if ln.get("cell"):
            x = v["Summary"][ln["cell"]].value
            if isinstance(x, (int, float)):
                ln["value"] = float(x)


def render(scores: list[Score], *, title: str = "Unit-rate accuracy gate (leave-one-out)") -> str:
    def pct(a: int, b: int) -> str:
        return f"{a}/{b} ({a / b:.1%})" if b else "n/a"

    ie = sum(s.install_exact for s in scores)
    it = sum(s.install_tol for s in scores)
    n = sum(s.install_n for s in scores)
    st = sum(s.supply_tol for s in scores)
    sn = sum(s.supply_n for s in scores)
    sheets = [x for s in scores for x in s.sheets]
    summ = [x for s in scores for x in s.summary]
    lines = [f"# {title}", "",
             f"**Install €/unit exact: {pct(ie, n)}** · within ±10 %: {pct(it, n)} · "
             f"**Supply within ±15 %: {pct(st, sn)}**",
             f"Sheet totals within ±5 %: {pct(sum(x['ok'] for x in sheets), len(sheets))} · Summary lines within ±5 %: "
             f"{pct(sum(x['ok'] for x in summ), len(summ))}",
             f"Lines priced without a model call: {pct(sum(s.model_free for s in scores), sum(s.lines for s in scores))}"
             f" · cost per run: ${statistics.mean([s.cost_usd for s in scores]) if scores else 0:.4f}",
             f"CHECK {sum(s.check for s in scores)} · NO PRICE {sum(s.no_price for s in scores)} · "
             f"interpolated from a band {sum(s.interpolated for s in scores)}", "",
             "| Held-out | Install exact | Install ±10 % | Supply ±15 % | Sheets ±5 % | Summary ±5 % | CHECK | NO PRICE "
             "| Band | Recalc problems |", "|---|---|---|---|---|---|---|---|---|---|"]
    for s in scores:
        lines.append(f"| {s.name} | {pct(s.install_exact, s.install_n)} | {pct(s.install_tol, s.install_n)} | "
                     f"{pct(s.supply_tol, s.supply_n)} | {pct(sum(x['ok'] for x in s.sheets), len(s.sheets))} | "
                     f"{pct(sum(x['ok'] for x in s.summary), len(s.summary))} | {s.check} | {s.no_price} | "
                     f"{s.interpolated} | {len(s.problems)} |")
    lines += ["", "## Sheet totals", "", "| Held-out | Sheet | Hand | Agent | Δ |", "|---|---|---|---|---|"]
    for s in scores:
        for x in s.sheets:
            d = "–" if x["agent"] is None else f"{(x['agent'] - x['hand']) / x['hand']:+.1%}"
            ag = "–" if x["agent"] is None else f"{x['agent']:,.0f}"
            lines.append(f"| {s.name} | {x['sheet']} | {x['hand']:,.0f} | {ag} | {d} |")
    return "\n".join(lines) + "\n"


def offline_loo(paths: dict[str, str], work_dir: str, *, recalc: bool = True) -> list[Score]:
    """Leave-one-out over parsed workbooks (no database): every file is blanked, then filled from the others."""
    import time

    from .fill import Planner, extras_from_workbook, verify, write_plan
    from .match import RefMeta, UnitRateKB, Want, rows_from_workbook
    books: dict[str, tuple[BoqWorkbook, RefMeta]] = {}
    for name, p in paths.items():
        wb = read_boq(p, file_name=name)
        summary_values(p, wb)
        books[name] = (wb, RefMeta(file_id=name, file_name=name[:40], market=wb.market, client=wb.client,
                                   end_client=wb.end_client, package=wb.package, doc_date=wb.doc_date))
    os.makedirs(work_dir, exist_ok=True)
    scores = []
    for held, (truth, _meta) in books.items():
        t0 = time.time()
        others = [(w, m) for n, (w, m) in books.items() if n != held]
        kb = UnitRateKB([r for w, m in others for r in rows_from_workbook(w, m)],
                        Want(market=truth.market, client=truth.client, package=truth.package,
                             end_client=truth.end_client))
        extras = [e for w, m in others for e in extras_from_workbook(w, m)]
        ratios = [(m, rr) for w, m in others for rr in w.ratios]
        safe = re.sub(r"[^\w.-]+", "_", held)[:60]
        blank = os.path.join(work_dir, f"blank_{safe}.xlsx")
        make_blank(paths[held], blank, read_boq(paths[held], file_name=held))
        bwb = read_boq(blank, file_name=held)
        setup = setup_for(truth)
        plan = Planner(bwb, kb, setup, extras=extras, ref_ratios=ratios).plan()
        out = os.path.join(work_dir, f"filled_{safe}.xlsx")
        write_plan(blank, out, plan, bwb)
        ver = verify(out, plan, bwb) if recalc else None
        sc = score_plan(held, truth, plan, ver)
        sc.seconds = round(time.time() - t0, 1)
        scores.append(sc)
    return scores


# ------------------------------------------------------------------ through the chat run path (database)


def _answers_for(card: dict, setup: Setup) -> dict:
    """The setup card answered like the estimator (programme, packages, scope, delivery %, dates; never prices)."""
    out: dict = {}
    for q in card["payload"]["questions"]:
        qid = q["id"]
        if qid.startswith("weeks:"):
            ph = qid.split(":", 1)[1]
            out[qid] = setup.phase_weeks.get(ph) or setup.programme_weeks or q.get("default")
        elif qid.startswith("weeks_sheet:"):
            out[qid] = setup.sheet_weeks.get(qid.split(":", 1)[1]) or setup.programme_weeks
        elif qid == "packages":
            out[qid] = [o for o in q["options"] if setup.participating.get(o, True)]
        elif qid == "prelims_weeks":
            out[qid] = setup.prelims_weeks
        elif qid == "material_supply":
            out[qid] = "Yes" if setup.material_supply else "No"
        elif qid == "delivery_pct":
            out[qid] = setup.delivery_pct
        elif qid in ("offer_date", "validity_date"):
            out[qid] = getattr(setup, qid)
        elif qid == "subcontractor":
            out[qid] = setup.subcontractor
    return out


def score_document(name: str, truth: BoqWorkbook, run: dict) -> Score:
    from .. import db
    s = Score(name=name)
    doc = db.fetchone("SELECT * FROM documents WHERE run_id=%s", (run["id"],))
    rows = {(r["sheet_name"], r["row_idx"]): r for r in
            db.fetchall("SELECT * FROM estimate_rows WHERE document_id=%s", (doc["id"],))} if doc else {}
    for sh in truth.pricing_sheets:
        for r in sh.rows:
            if r.kind != "item":
                continue
            got = rows.get((sh.name, r.row))
            if r.install_rate and r.rate_basis != "supply_only":
                s.install_n += 1
                v = float(got["unit_labour"]) if got and got["unit_labour"] is not None else None
                s.install_exact += v is not None and abs(v - r.install_rate) < 1e-9
                if v is not None and abs(v - r.install_rate) <= INSTALL_TOL * r.install_rate:
                    s.install_tol += 1
                elif len(s.misses) < 400:
                    s.misses.append({"sheet": sh.name, "row": r.row, "what": "install", "text": r.description[:70],
                                     "hand": r.install_rate, "agent": v,
                                     "method": got["install_method"] if got else None})
            if r.supply_rate:
                s.supply_n += 1
                v = float(got["unit_material"]) if got and got["unit_material"] is not None else None
                s.supply_tol += v is not None and abs(v - r.supply_rate) <= SUPPLY_TOL * r.supply_rate
    for r in rows.values():
        s.lines += 1
        s.model_free += not r["model_used"]
        s.check += "CHECK" in (r["flags"] or [])
        s.no_price += "NO PRICE" in (r["flags"] or [])
        s.interpolated += r.get("install_method") == "band"
    checks = (doc or {}).get("checks") or {}
    for sh in truth.pricing_sheets:
        hand = sh.totals.get("total")
        if hand:
            agent = (checks.get("totals") or {}).get(sh.name, {}).get("total")
            s.sheets.append({"sheet": sh.name, "hand": round(hand, 2), "agent": agent,
                             "ok": agent is not None and abs(agent - hand) <= TOTAL_TOL * abs(hand)})
    hand_summary = {ln["label"]: ln for ln in truth.summary_lines}
    for label, agent in (checks.get("summary") or {}).items():
        h = hand_summary.get(label, {}).get("value")
        if h:
            s.summary.append({"label": label, "hand": h, "agent": agent, "ok": abs(agent - h) <= TOTAL_TOL * abs(h)})
    s.problems = list(checks.get("problems") or [])
    u = db.fetchone("SELECT COALESCE(SUM(cost_usd),0) c FROM usage WHERE run_id=%s", (run["id"],))
    s.cost_usd = float(u["c"])
    return s


def db_loo(files: list[str], work_dir: str) -> tuple[list[Score], list[dict]]:
    """Leave-one-out through ingestion and the chat run path: every file is ingested as uploaded (.xls too), then each
    one is blanked and filled by a run whose allowed references are the other files."""
    import time

    from .. import db
    from ..accuracy_gate import fill as gate_fill
    from ..accuracy_gate import gate_user, ingest
    from ..agent import run as agent_run
    from ..agent.common import ensure_xlsx
    user = gate_user()
    ids: dict[str, str] = {}
    paths: dict[str, str] = {}
    os.makedirs(work_dir, exist_ok=True)
    for f in files:
        name = os.path.basename(f)
        ids[name] = ingest(f, "reference_estimate", str(user["id"]))
        if f.lower().endswith(".xls"):
            tmp = os.path.join(work_dir, "xls_" + re.sub(r"[^\w.-]+", "_", name))
            os.makedirs(tmp, exist_ok=True)
            shutil.copyfile(f, os.path.join(tmp, name))
            paths[name] = ensure_xlsx(os.path.join(tmp, name))
        else:
            paths[name] = f
    ingested = db.fetchall("""SELECT id, original_name, status, pricing_model, market, client, end_client, package,
                                     (SELECT count(*) FROM price_items p WHERE p.file_id=f.id) n
                              FROM files f WHERE id = ANY(%s::uuid[])""", (list(ids.values()),))
    scores = []
    for held, p in paths.items():
        truth = read_boq(p, file_name=held)
        summary_values(p, truth)
        safe = re.sub(r"[^\w.-]+", "_", held)[:60]
        blank = os.path.join(work_dir, f"blank_{safe}.xlsx")
        make_blank(p, blank, read_boq(p, file_name=held))
        t0 = time.time()
        allowed = [v for k, v in ids.items() if k != held]
        res = gate_fill(blank, held, user, allowed)
        run = res["run"]
        card = db.fetchone("""SELECT * FROM cards WHERE run_id=%s AND kind='clarify' AND status='pending'
                              AND payload->>'purpose'='unit_rate_setup'""", (run["id"],))
        if card:
            db.execute("UPDATE cards SET status='answered', decision=%s WHERE id=%s",
                       (db.jsonb({"action": "answer", "data": {"answers": _answers_for(card, setup_for(truth))}}),
                        card["id"]))
            agent_run.resume_run({"run_id": str(run["id"]), "card_id": str(card["id"])})
            run = db.fetchone("SELECT * FROM runs WHERE id=%s", (run["id"],))
        sc = score_document(held, truth, run)
        sc.seconds = round(time.time() - t0, 1)
        scores.append(sc)
        print(f"  {held[:60]}: {run['status']} install exact {sc.install_exact}/{sc.install_n}", flush=True)
    return scores, [dict(r, id=str(r["id"])) for r in ingested]
