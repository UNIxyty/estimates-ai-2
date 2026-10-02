"""Unit-rate BOQ reader (EU data-centre subcontract BOQs: Total = Quantity × Rate, no hours, no hourly rate).

Reads a workbook into a BoqWorkbook: Summary (project / date / package / subcontractor fields / cost lines incl.
"Not participating"), Site Attendance, Preliminaries (weekly staff rates × weeks, items), and every pricing sheet:

  layout A  Description | [Building A | Building B] | Quantity | UoM | Rate | Total
  layout B  … | Quantity | UoM | Rate | Total | Unit price (UNIT COST) | Total          (install + material supply)
  layout C  per-phase (qty, €) pairs or qty columns [× INTERNAL/EXTERNAL] | Quantity | UoM | Rate | Total [| …]

Columns are found by header labels; phase € columns by their formulas (=qty × Rate), never by the header text
(FRA5 Lighting labels one of them "Labor"). Section headers carry the pricing basis ("(Free Issued for Install)",
"(Supply & Install)", Optional supply-of-brackets sheets); notes ("Note: - DB/Equip Header is equal to 5 meters …")
are kept with the rows they govern. Contractor Items (months × monthly rate, tools / fixation lump sums, materials
delivery %) and accessory ratios (qty of each row per anchor row in a Supply & Install section) are extracted too.

Input cells (light-blue theme-4 fill in client BOQs, or no fill in MGS's own CostX sheets) are recorded as a hint;
"not a formula" is what makes a cell writable."""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import asdict, dataclass, field
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .attributes import unit_rate_attributes

# ------------------------------------------------------------------ small helpers


def _s(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def _num(v: Any) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = _s(v).replace("€", "").replace("\xa0", "").replace(" ", "")
    if not t:
        return None
    if "," in t and "." in t:
        t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _is_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


UOM_MAP = {"no": "no", "nr": "no", "no.": "no", "nr.": "no", "pcs": "no", "pc": "no", "ea": "no", "each": "no",
           "gab": "no", "gab.": "no", "stk": "no", "m": "m", "lm": "m", "mtr": "m", "m.": "m", "item": "item",
           "items": "item", "ls": "item", "sum": "item", "wks": "week", "wk": "week", "week": "week", "weeks": "week",
           "month": "month", "months": "month", "mth": "month", "mths": "month", "%": "%", "m2": "m2", "m²": "m2",
           "set": "set", "kg": "kg"}


def norm_uom(u: Any) -> str | None:
    t = _s(u).lower()
    return UOM_MAP.get(t, t or None)


def _fill(cell) -> str | None:
    f = cell.fill
    if not f or f.fill_type is None or f.fgColor is None:
        return None
    c = f.fgColor
    if c.type == "theme":
        return f"theme{c.theme}{'+' if (c.tint or 0) > 0 else ''}{round(c.tint or 0, 2)}"
    if c.type == "rgb" and c.rgb and c.rgb not in ("00000000",):
        return f"#{c.rgb[-6:].upper()}"
    if c.type == "indexed":
        return f"idx{c.indexed}"
    return None


def is_input_fill(code: str | None) -> bool:
    """Light-blue input fill of client BOQs: theme colour 4 with a positive tint, or the .xls palette blues."""
    if not code:
        return False
    return bool(re.match(r"theme4\+", code)) or code in ("#99CCFF", "#CCFFFF", "#C5D9F1", "#DCE6F1", "#BDD7EE",
                                                         "#DDEBF7", "#9BC2E6") or code in ("idx44", "idx41")


class Values:
    """Cell values for one sheet: the cached value when the file has one (saved by Excel / LibreOffice), else a safe
    evaluation of simple same-sheet arithmetic (=SUM(C15:D15), =C46*12, =C56/1.2, =C15+D15). Anything else is None."""

    _SAFE = re.compile(r"^[0-9.+\-*/() eE]+$")

    def __init__(self, wsf, wsv):
        self.f, self.v = wsf, wsv
        self._cache: dict[str, float | None] = {}

    def cell(self, r: int, c: int):
        return self.get(f"{get_column_letter(c)}{r}")

    def get(self, coord: str, _depth: int = 0):
        coord = coord.replace("$", "")
        if coord in self._cache:
            return self._cache[coord]
        cached = self.v[coord].value
        if cached is not None:
            return cached
        raw = self.f[coord].value
        if not _is_formula(raw):
            return raw
        if _depth > 40:
            return None
        self._cache[coord] = None  # cycle guard
        expr = raw[1:].replace("$", "")
        if "!" in expr or re.search(r"[A-Z]{2,}\(", expr.replace("SUM(", "")):
            return None

        def rng(m):
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            tot = 0.0
            for rr in range(r1, r2 + 1):
                for cc in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                    v = _num(self.get(f"{get_column_letter(cc)}{rr}", _depth + 1))
                    tot += v or 0.0
            return repr(tot)
        expr = re.sub(r"SUM\(\s*([A-Z]{1,3})(\d+)\s*:\s*([A-Z]{1,3})(\d+)\s*\)", rng, expr)
        expr = re.sub(r"SUM\(([^()]*)\)", r"(\1)", expr)
        expr = re.sub(r"\b([A-Z]{1,3})(\d+)\b",
                      lambda m: repr(_num(self.get(f"{m.group(1)}{m.group(2)}", _depth + 1)) or 0.0), expr)
        expr = expr.replace(",", "+")
        if not self._SAFE.match(expr):
            return None
        try:
            val = float(eval(expr, {"__builtins__": {}}, {}))  # noqa: S307 - digits and + - * / ( ) only
        except Exception:  # noqa: BLE001
            val = None
        self._cache[coord] = val
        return val


def _refs(formula: str) -> list[str]:
    return re.findall(r"\$?([A-Z]{1,3})\$?\d+", formula or "")


# ------------------------------------------------------------------ data model

BASIS_FREE_ISSUE = "install_only"


@dataclass
class BoqRow:
    sheet: str
    row: int
    kind: str                         # item | contractor | lump_sum | delivery
    description: str
    tag: str | None = None            # column A text: finish (PG/HDG) or category (Ladder/Bracketry)
    section: str | None = None
    section_basis: str | None = None  # install_only | supply_only | supply_and_install
    notes: list[str] = field(default_factory=list)
    qty: float | None = None
    uom: str | None = None
    uom_norm: str | None = None
    install_rate: float | None = None
    supply_rate: float | None = None
    total: float | None = None
    supply_total: float | None = None
    rate_basis: str | None = None     # install_only | supply_only | supply_and_install | lump_sum | weekly | monthly | percent
    package: str | None = None
    phase: str | None = None          # contractor items: the phase they belong to
    phase_qty: dict[str, float] = field(default_factory=dict)
    qty_formula: str | None = None
    cells: dict[str, str] = field(default_factory=dict)      # {"qty":"E15","rate":"G15","total":"H15","supply":…}
    input_hint: dict[str, bool] = field(default_factory=dict)  # rate/supply/total cell: blue fill & not a formula
    attrs: dict[str, Any] = field(default_factory=dict)
    flags: list[str] = field(default_factory=list)


@dataclass
class BoqSheet:
    name: str
    kind: str                         # pricing | prelims | summary | attendance | takeoff | calc | other
    layout: str | None = None         # A | B | C
    header_row: int | None = None
    cols: dict[str, str] = field(default_factory=dict)        # desc qty uom rate total supply supply_total tag
    phases: list[dict] = field(default_factory=list)          # [{label, area, qty_col, money_col}]
    buildings: list[dict] = field(default_factory=list)       # [{label, qty_col}]
    areas: list[str] = field(default_factory=list)
    has_supply: bool = False
    package: str | None = None
    notes: list[str] = field(default_factory=list)            # sheet-level notes (before the first section)
    sections: list[dict] = field(default_factory=list)        # [{title, row, basis, notes}]
    rows: list[BoqRow] = field(default_factory=list)
    total_cells: dict[str, str] = field(default_factory=dict)
    totals: dict[str, float] = field(default_factory=dict)
    input_fill: bool = False


@dataclass
class BoqWorkbook:
    project: str | None = None
    end_client: str | None = None
    client: str | None = None
    market: str | None = None
    doc_date: str | None = None       # ISO date
    package: str | None = None        # containment | lighting | gs_sp | cable | electrical (several)
    currency: str = "EUR"
    language: str = "EN"
    subcontractor: dict = field(default_factory=dict)   # {name:{cell,value}, offer_date:{…}, validity:{…}}
    summary_lines: list[dict] = field(default_factory=list)
    sheets: list[BoqSheet] = field(default_factory=list)
    prelims: list[dict] = field(default_factory=list)
    attendance: dict = field(default_factory=dict)
    delivery: list[dict] = field(default_factory=list)
    ratios: list[dict] = field(default_factory=list)
    is_takeoff: bool = False          # CostX export (Folder | Dimension Group | qty | UOM) present

    @property
    def pricing_sheets(self) -> list[BoqSheet]:
        return [s for s in self.sheets if s.kind == "pricing"]

    def as_dict(self) -> dict:
        d = asdict(self)
        return d


# ------------------------------------------------------------------ detection

_NORM_HEADER = re.compile(r"laika\s*norm|norma|darbietilp|c\s*/\s*h|stundas|hourly|samaksas\s*likme|norm\s*h|\btimer\b",
                          re.I)


def _row_texts(ws, r: int, max_col: int) -> dict[int, str]:
    out = {}
    for c in range(1, max_col + 1):
        v = ws.cell(r, c).value
        if isinstance(v, str) and v.strip():
            out[c] = _s(v)
    return out


def _find_pricing_header(ws, max_scan: int = 40) -> tuple[int, dict[str, int]] | None:
    for r in range(1, min(ws.max_row, max_scan) + 1):
        texts = {c: t.lower() for c, t in _row_texts(ws, r, min(ws.max_column, 40)).items()}
        qty = [c for c, t in texts.items() if t in ("quantity", "qty", "quantities")]
        rate = [c for c, t in texts.items() if t == "rate"]
        if qty and rate:
            cols = {"qty": qty[0], "rate": next((c for c in rate if c > qty[0]), rate[0])}
            tot = sorted(c for c, t in texts.items() if t.startswith("total"))
            after_rate = [c for c in tot if c > cols["rate"]]
            if after_rate:
                cols["total"] = after_rate[0]
            sup = [c for c, t in texts.items() if t in ("unit price", "unit cost", "unit costs", "material unit price")]
            if sup:
                cols["supply"] = sup[0]
                st = [c for c in tot if c > sup[0]]
                if st:
                    cols["supply_total"] = st[0]
            u = [c for c, t in texts.items() if t in ("uom", "unit", "uoм", "units")]
            if u:
                cols["uom"] = u[0]
            ref = [c for c, t in texts.items() if t == "ref"]
            if ref:
                cols["ref"] = ref[0]
            desc = [c for c, t in texts.items() if t == "description"]
            if desc:
                cols["desc"] = desc[0]
            notes = [c for c, t in texts.items() if t == "notes"]
            if notes:
                cols["notes"] = notes[0]
            fin = [c for c, t in texts.items() if t == "finish"]
            if fin:
                cols["tag"] = fin[0]
            return r, cols
    return None


def is_unit_rate_workbook(path: str) -> bool:
    """A BOQ priced per unit: a Quantity/Rate/Total pricing header and no norm-hours or hourly-rate column, or a CostX
    takeoff export. Cheap check used by ingestion to choose the reader."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        found = False
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=40, max_col=40, values_only=True))
            texts = [[_s(v).lower() for v in r if isinstance(v, str)] for r in rows]
            flat = [t for r in texts for t in r]
            if any(_NORM_HEADER.search(t) for t in flat):
                return False
            for r in texts:
                if ("quantity" in r or "qty" in r) and "rate" in r:
                    found = True
                if "folder" in r and "dimension group" in r:
                    found = True
        return found
    finally:
        wb.close()


# ------------------------------------------------------------------ reader


_MONTHS = {m: i for i, m in enumerate(["january", "february", "march", "april", "may", "june", "july", "august",
                                         "september", "october", "november", "december"], 1)}


def parse_date(v: Any) -> str | None:
    if isinstance(v, (dt.datetime, dt.date)):
        return (v.date() if isinstance(v, dt.datetime) else v).isoformat()
    t = _s(v).lower()
    m = re.search(r"(\d{1,2})(?:st|nd|rd|th)?\s*(?:of\s+)?([a-z]+)\s*,?\s*(\d{4})", t)
    if m and m.group(2) in _MONTHS:
        try:
            return dt.date(int(m.group(3)), _MONTHS[m.group(2)], int(m.group(1))).isoformat()
        except ValueError:
            return None
    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", t)
    if m:
        try:
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
        except ValueError:
            return None
    return None


_MARKET_HINTS = [(re.compile(r"\bams\d*\b|amsterdam|nl\b|netherlands|holland", re.I), "NL"),
                 (re.compile(r"\bfra\d*\w*\b|frankfurt|germany|deutschland|\bde\b", re.I), "DE"),
                 (re.compile(r"\briga\b|rīga|latvia|latvija|\blv\b", re.I), "LV"),
                 (re.compile(r"\bcph\d*\b|copenhagen|denmark|danmark", re.I), "DK"),
                 (re.compile(r"\bdub\d*\b|dublin|ireland", re.I), "IE"),
                 (re.compile(r"\blon\d*\b|london|\buk\b", re.I), "UK")]


def guess_market(*texts: str | None) -> str | None:
    for t in texts:
        for rx, m in _MARKET_HINTS:
            if t and rx.search(t):
                return m
    return None


def guess_package(*texts: str | None) -> str | None:
    t = " ".join(x for x in texts if x).lower()
    if re.search(r"containment|ladder|tray|bracket|optional", t):
        return "containment"
    if re.search(r"lighting", t):
        return "lighting"
    if re.search(r"general\s*services|gs[\s-]*sp|small\s*power|\bgs\b|socket", t):
        return "gs_sp"
    if re.search(r"\bcable\b", t):
        return "cable"
    if re.search(r"prelim", t):
        return "prelims"
    return None


def read_boq(path: str, *, file_name: str | None = None) -> BoqWorkbook:
    wbf = openpyxl.load_workbook(path)                 # formulas + styles
    wbv = openpyxl.load_workbook(path, data_only=True)  # cached values
    out = BoqWorkbook()
    all_text: list[str] = []
    for ws in wbf.worksheets:
        name = ws.title.strip()
        low = name.lower()
        wv = Values(ws, wbv[ws.title])
        first = {t.lower() for t in _row_texts(ws, 1, min(ws.max_column, 12)).values()}
        if "folder" in first and "dimension group" in first:
            out.is_takeoff = True
            out.sheets.append(BoqSheet(name=ws.title, kind="takeoff"))
            continue
        if low == "summary":
            _read_summary(ws, wv, out)
            out.sheets.append(BoqSheet(name=ws.title, kind="summary"))
            continue
        if "attendance" in low:
            out.attendance = _read_attendance(ws)
            out.sheets.append(BoqSheet(name=ws.title, kind="attendance"))
            continue
        hdr = _find_pricing_header(ws)
        if hdr and "ref" in hdr[1] and ("prelim" in low or "desc" in hdr[1]) and "prelim" in low:
            out.prelims.extend(_read_prelims(ws, wv, hdr))
            out.sheets.append(BoqSheet(name=ws.title, kind="prelims", header_row=hdr[0],
                                       package=guess_package(name)))
            continue
        if hdr:
            sh = _read_pricing(ws, wv, hdr)
            sh.package = guess_package(name) or None
            out.sheets.append(sh)
            continue
        kind = "calc" if re.search(r"calc|total|sheet\d", low) else "other"
        out.sheets.append(BoqSheet(name=ws.title, kind=kind))
    for ws in wbf.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), max_col=min(ws.max_column, 12), values_only=True):
            all_text.extend(_s(v) for v in row if isinstance(v, str))
    joined = " ".join(all_text)
    if re.search(r"winthrop", joined, re.I):
        out.client = "Winthrop"
    proj = out.project or ""
    m = re.search(r"\b(goodman|ntt|equinix|digital\s*realty|vantage|cyrusone|microsoft|google|amazon|aws|meta)\b",
                  f"{proj} {file_name or ''}", re.I)
    if m:
        out.end_client = m.group(1).upper() if m.group(1).lower() in ("ntt", "aws") else m.group(1).title()
    out.market = guess_market(proj, file_name)
    pkgs = {s.package for s in out.pricing_sheets if s.package}
    sum_pkg = guess_package(out.package or "", file_name or "")
    if len(pkgs) > 1:
        out.package = "electrical"
    else:
        out.package = next(iter(pkgs), None) or sum_pkg
    for s in out.pricing_sheets:
        s.package = s.package or out.package
        for r in s.rows:
            r.package = r.package or (s.package if r.kind == "item" else "contractor_items")
    out.ratios = _ratios(out)
    return out


# ------------------------------------------------------------------ Summary / attendance / prelims


def _read_summary(ws, wv, out: BoqWorkbook) -> None:
    for r in range(1, min(ws.max_row, 60) + 1):
        texts = _row_texts(ws, r, min(ws.max_column, 12))
        for c, t in texts.items():
            tl = t.lower().rstrip(": ")
            right = next((ws.cell(r, cc) for cc in range(c + 1, c + 6) if ws.cell(r, cc).value not in (None, "")), None)
            val = wv.cell(right.row, right.column) if right is not None else None
            if tl == "project" and out.project is None:
                out.project = _s(val) or None
            elif tl == "date" and out.doc_date is None:
                out.doc_date = parse_date(val)
            elif tl == "package" and out.package is None:
                out.package = _s(val) or None
            elif tl in ("name of subcontractor", "date of offer", "offer valid to date", "offer validity"):
                key = {"name of subcontractor": "name", "date of offer": "offer_date"}.get(tl, "validity")
                # the input cell is the yellow / blank cell in the value column (E), whatever it holds now
                cell = right if right is not None else ws.cell(r, c + 2)
                if right is None:
                    for cc in range(c + 1, c + 6):
                        f = _fill(ws.cell(r, cc))
                        if f and f.upper().endswith("FFFF00"):
                            cell = ws.cell(r, cc)
                            break
                out.subcontractor[key] = {"cell": cell.coordinate, "value": _s(val) or None}
    # cost summary lines
    start = None
    for r in range(1, min(ws.max_row, 80) + 1):
        if any(t.lower() == "cost summary" for t in _row_texts(ws, r, min(ws.max_column, 12)).values()):
            start = r
            break
    if start is None:
        return
    for r in range(start + 1, min(ws.max_row, start + 40) + 1):
        texts = _row_texts(ws, r, min(ws.max_column, 14))
        # a label cell may be a formula (='Containment Install'!B4): use its cached text
        for c, txt in list(texts.items()):
            if txt.startswith("="):
                cv = wv.cell(r, c)
                if isinstance(cv, str) and cv.strip():
                    texts[c] = _s(cv)
        label = next((t for c, t in sorted(texts.items())
                      if len(t) > 3 and not t.startswith("=") and t.lower() not in ("prelims", "e", "o")), "")
        if not label:
            continue
        if label.lower().startswith("total"):
            break
        link = None
        link_cell = None
        for c in range(1, min(ws.max_column, 14) + 1):
            v = ws.cell(r, c).value
            if _is_formula(v) and "!" in v:
                m = re.search(r"=\s*(?:\[\d+\])?'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)", v)
                if m:
                    link = {"sheet": m.group(1).strip(), "cell": f"{m.group(2)}{m.group(3)}"}
                    link_cell = ws.cell(r, c).coordinate
                    break
        not_part = any(re.search(r"not\s*participat", t, re.I) for t in texts.values())
        code = next((t for t in texts.values() if t in ("Prelims", "E", "O")), None)
        out.summary_lines.append({"row": r, "label": label, "code": code, "link": link, "cell": link_cell,
                                  "not_participating": not_part,
                                  "package": guess_package(label),
                                  "supply": bool(re.search(r"supply", label, re.I)),
                                  "prelims": bool(re.search(r"prelim", label, re.I)),
                                  "optional": code == "O" or bool(re.search(r"optional", label, re.I))})


def _read_attendance(ws) -> dict:
    """Tick matrix: which attendance items the main contractor provides vs the subcontractor (kept as issued)."""
    hdr = None
    for r in range(1, min(ws.max_row, 40) + 1):
        texts = {c: t.lower() for c, t in _row_texts(ws, r, min(ws.max_column, 10)).items()}
        if any("sub-contractor" in t or "subcontractor" in t for t in texts.values()):
            hdr = (r, texts)
            break
    items = 0
    by_sub = 0
    if hdr:
        sub_col = next(c for c, t in hdr[1].items() if "sub-contractor" in t or "subcontractor" in t)
        for r in range(hdr[0] + 1, ws.max_row + 1):
            d = ws.cell(r, 2).value
            if isinstance(d, str) and d.strip():
                items += 1
                if _s(ws.cell(r, sub_col).value):
                    by_sub += 1
    return {"items": items, "by_subcontractor": by_sub}


def _read_prelims(ws, wv, hdr) -> list[dict]:
    r0, cols = hdr
    desc_c = cols.get("desc", 2)
    unit_c = cols.get("uom")
    out = []
    for r in range(r0 + 1, ws.max_row + 1):
        d = _s(ws.cell(r, desc_c).value)
        if not d:
            continue
        if re.search(r"total\s*-\s*c/f|^total", d, re.I):
            break
        qty = _num(wv.cell(r, cols["qty"]))
        rate = _num(wv.cell(r, cols["rate"]))
        unit = norm_uom(ws.cell(r, unit_c).value) if unit_c else None
        total = _num(wv.cell(r, cols["total"])) if cols.get("total") else None
        if qty is None and rate is None and unit is None:
            continue
        basis = "weekly" if unit == "week" else "monthly" if unit == "month" else "lump_sum"
        out.append({"sheet": ws.title, "row": r, "ref": _s(ws.cell(r, cols.get("ref", 1)).value) or None,
                    "description": d, "qty": qty, "uom": unit, "rate": rate, "total": total, "rate_basis": basis,
                    "cells": {"qty": ws.cell(r, cols["qty"]).coordinate, "rate": ws.cell(r, cols["rate"]).coordinate,
                              **({"total": ws.cell(r, cols["total"]).coordinate} if cols.get("total") else {})},
                    "rate_is_formula": _is_formula(ws.cell(r, cols["rate"]).value)})
    return out


# ------------------------------------------------------------------ pricing sheets

_PHASE = re.compile(r"^phase\s*(\d+)", re.I)
_AREA = re.compile(r"^(internal|external)$", re.I)
_BUILDING = re.compile(r"^building\s+[a-z0-9]+$", re.I)
_CONTRACTOR = re.compile(r"^contractor\s*items", re.I)
_TOTAL_ROW = re.compile(r"^total\b|c/f\s*to\s*summary|^sub[\s-]*total", re.I)
_NOTE = re.compile(r"^note\b|^notes?:", re.I)
_LUMP = re.compile(r"fixation|power\s*tools|tools\b", re.I)
_MONTHLY = re.compile(r"accom\w*dation|transport|scissor|lift|mewp|cherry\s*picker", re.I)
_DELIVERY = re.compile(r"delivery", re.I)
_HEADERISH = re.compile(r"^(description|all areas|dimension group|site|item|items|building [a-z0-9]+)$", re.I)


def _basis_from_section(title: str, sheet_name: str) -> str | None:
    t = title.lower()
    if "free issue" in t or "free issued" in t:
        return "install_only"
    if re.search(r"supply\s*(?:&|and)\s*install", t):
        return "supply_and_install"
    if re.search(r"supply\s*only|supply of (?:the )?brackets|optional price", t) or \
            re.search(r"optional", sheet_name, re.I):
        return "supply_only"
    if re.search(r"install(?:ation)?\s*only", t):
        return "install_only"
    return None


def _read_pricing(ws, wv, hdr) -> BoqSheet:
    r0, cols = hdr
    sh = BoqSheet(name=ws.title, kind="pricing", header_row=r0)
    qty_c, rate_c = cols["qty"], cols["rate"]
    if "uom" not in cols:  # UoM label may sit on a different row (CostX "UOM" two rows below)
        for r in range(max(1, r0 - 3), min(ws.max_row, r0 + 4) + 1):
            for c, t in _row_texts(ws, r, min(ws.max_column, 40)).items():
                if t.lower() in ("uom", "unit"):
                    cols["uom"] = c
                    break
            if "uom" in cols:
                break
    # description column: most text cells in the data rows, left of the quantity column
    counts: dict[int, int] = {}
    for r in range(r0 + 1, min(ws.max_row, r0 + 120) + 1):
        for c in range(1, qty_c):
            v = ws.cell(r, c).value
            if isinstance(v, str) and len(v.strip()) > 6 and not v.startswith("="):
                counts[c] = counts.get(c, 0) + 1
    desc_c = max(counts, key=lambda c: (counts[c], -c)) if counts else cols.get("desc", 2)
    cols["desc"] = desc_c
    if "tag" not in cols and desc_c > 1:
        cols["tag"] = 1
    sh.cols = {k: get_column_letter(v) for k, v in cols.items()}
    # notes above the header (yellow "Luminaires … will be free issued" lines, "Note: …")
    for r in range(max(1, r0 - 6), r0):
        for c in range(1, qty_c):
            cell = ws.cell(r, c)
            if isinstance(cell.value, str) and len(cell.value.strip()) > 15 and not cell.value.startswith("=") and \
                    (_NOTE.match(_s(cell.value)) or (_fill(cell) or "").endswith("FFFF00")):
                sh.notes.append(_s(cell.value))
    sh.has_supply = "supply" in cols

    # phase / building / area labels near the header
    labels: dict[int, list[tuple[int, str]]] = {}
    for r in range(max(1, r0 - 8), min(ws.max_row, r0 + 10) + 1):
        for c, t in _row_texts(ws, r, qty_c).items():
            labels.setdefault(c, []).append((r, t))
    phase_rows = {r for lst in labels.values() for r, t in lst if _PHASE.match(t)}
    phase_at: dict[int, str] = {}
    if phase_rows:
        prow = min(phase_rows)
        for c, lst in labels.items():
            for r, t in lst:
                if r == prow and not _HEADERISH.match(t) and len(t) <= 30:
                    phase_at[c] = t
    area_at = {c: t.upper() for c, lst in labels.items() for _, t in lst if _AREA.match(t)}
    bld_at = {c: t for c, lst in labels.items() for _, t in lst if _BUILDING.match(t)}
    sh.buildings = [{"label": t, "qty_col": get_column_letter(c)} for c, t in sorted(bld_at.items())]

    # money columns = formula "=<qty col>*<rate col>" in data rows; their qty column is the other reference
    rate_l = get_column_letter(rate_c)
    money_of: dict[int, int] = {}
    for r in range(r0 + 1, min(ws.max_row, r0 + 80) + 1):
        for c in range(1, qty_c):
            v = ws.cell(r, c).value
            if _is_formula(v) and re.fullmatch(r"=\s*\+?\$?[A-Z]{1,3}\$?\d+\s*\*\s*\$?[A-Z]{1,3}\$?\d+\s*", v):
                refs = _refs(v)
                if rate_l in refs:
                    other = next((x for x in refs if x != rate_l), None)
                    if other:
                        money_of[c] = column_index_from_string(other)
    if phase_at:
        group_cols = sorted(phase_at)
        def group_for(c: int) -> str:
            left = [pc for pc in group_cols if pc <= c]
            return phase_at[left[-1]] if left else phase_at[group_cols[0]]
        # qty columns: those feeding a money column, every INTERNAL/EXTERNAL column, and (FRA5 GS: qty-only
        # phases) the labelled group columns themselves when nothing else marks them
        money_cols = set(money_of)
        qcols = set(money_of.values()) | {c for c in area_at if c < qty_c}
        if not qcols:
            qcols = {c for c in group_cols if c < qty_c}
        qcols -= money_cols
        for qc in sorted(qcols):
            money = next((m for m, q in money_of.items() if q == qc), None)
            if money is None and qc + 1 < qty_c and qc + 1 not in qcols and qc + 1 not in phase_at \
                    and (qc + 1) not in area_at:
                money = qc + 1  # the € column next to it whose formula points elsewhere (FR12X Phase 3 EXT)
            label = _s(group_for(qc))
            m = _PHASE.match(label)
            sh.phases.append({"label": f"Phase {m.group(1)}" if m else label, "area": area_at.get(qc),
                              "qty_col": get_column_letter(qc),
                              "money_col": get_column_letter(money) if money else None})
        sh.areas = sorted({p["area"] for p in sh.phases if p["area"]})
    sh.layout = "C" if sh.phases else ("B" if sh.has_supply else "A")

    # rows
    section: str | None = None
    basis: str | None = _basis_from_section("", ws.title)
    sec_notes: list[str] = []
    contractor = False
    first_section_seen = False
    for r in range(r0 + 1, ws.max_row + 1):
        dcell = ws.cell(r, desc_c)
        d = _s(dcell.value) if isinstance(dcell.value, str) and not str(dcell.value).startswith("=") else ""
        row_texts = _row_texts(ws, r, min(ws.max_column, 30))
        qty_raw = ws.cell(r, qty_c).value
        qty = _num(wv.cell(r, qty_c))
        uom = ws.cell(r, cols["uom"]).value if cols.get("uom") else None
        rate = _num(wv.cell(r, rate_c))
        supply = _num(wv.cell(r, cols["supply"])) if cols.get("supply") else None
        total = _num(wv.cell(r, cols["total"])) if cols.get("total") else None
        stotal = _num(wv.cell(r, cols["supply_total"])) if cols.get("supply_total") else None
        any_total_word = any(_TOTAL_ROW.search(t) for t in row_texts.values())
        if any_total_word and not d:
            for k, cl in (("total", cols.get("total")), ("supply_total", cols.get("supply_total"))):
                if cl and total is not None and k == "total":
                    sh.total_cells.setdefault("total", ws.cell(r, cl).coordinate)
                    sh.totals.setdefault("total", total)
                if cl and stotal is not None and k == "supply_total":
                    sh.total_cells.setdefault("supply_total", ws.cell(r, cl).coordinate)
                    sh.totals.setdefault("supply_total", stotal)
            continue
        if not d or _HEADERISH.match(d):
            continue
        if _NOTE.match(d) or (_fill(dcell) or "").endswith("FFFF00") and qty is None and uom in (None, ""):
            (sec_notes if first_section_seen else sh.notes).append(d)
            continue
        if _CONTRACTOR.match(d):
            contractor = True
            section, sec_notes = "Contractor Items", []
            sh.sections.append({"title": "Contractor Items", "row": r, "basis": None, "notes": []})
            continue
        has_uom = bool(_s(uom))
        is_item = has_uom or qty is not None or _is_formula(qty_raw) or rate is not None or supply is not None
        if not is_item and (_DELIVERY.search(d) or _LUMP.search(d)) and (not d.isupper() or any(
                _num(wv.cell(r, c)) for c in range(desc_c + 1, min(ws.max_column, 40) + 1))):
            # a delivery / lump-sum line with only an amount (FRA3H GS: H113 = 42 534), or with nothing yet (a blank)
            is_item = True
        if not is_item and contractor and _MONTHLY.search(d) and not d.isupper():
            is_item = True  # a contractor item of a blank: no months, no rate yet
        if not is_item:
            # a section header (ALL CAPS group, a "(Free Issued for Install)" heading, a bracket-group title, …)
            section = d
            first_section_seen = True
            basis = _basis_from_section(d, ws.title) or _basis_from_section("", ws.title)
            sec_notes = []
            sh.sections.append({"title": d, "row": r, "basis": basis, "notes": sec_notes})
            continue
        row = BoqRow(sheet=ws.title, row=r, kind="item", description=d, section=section, section_basis=basis,
                     notes=list(sh.notes) + list(sec_notes), qty=qty, uom=_s(uom) or None, uom_norm=norm_uom(uom),
                     install_rate=rate, supply_rate=supply, total=total, supply_total=stotal)
        tag_c = cols.get("tag")
        if tag_c and tag_c != desc_c:
            tv = ws.cell(r, tag_c).value
            row.tag = _s(tv) if isinstance(tv, str) else None
        row.cells = {"qty": ws.cell(r, qty_c).coordinate, "rate": ws.cell(r, rate_c).coordinate}
        if cols.get("total"):
            row.cells["total"] = ws.cell(r, cols["total"]).coordinate
        if cols.get("supply"):
            row.cells["supply"] = ws.cell(r, cols["supply"]).coordinate
        if cols.get("supply_total"):
            row.cells["supply_total"] = ws.cell(r, cols["supply_total"]).coordinate
        if cols.get("uom"):
            row.cells["uom"] = ws.cell(r, cols["uom"]).coordinate
        for k in ("rate", "supply", "total", "supply_total"):
            if k in row.cells:
                cell = ws[row.cells[k]]
                row.input_hint[k] = is_input_fill(_fill(cell)) and not _is_formula(cell.value)
                if is_input_fill(_fill(cell)):
                    sh.input_fill = True
        if _is_formula(qty_raw):
            row.qty_formula = qty_raw
        for p in sh.phases:
            pv = _num(wv.get(f"{p['qty_col']}{r}"))
            if pv:
                key = p["label"] + (f" {p['area']}" if p["area"] else "")
                row.phase_qty[key] = round(row.phase_qty.get(key, 0) + pv, 6)
        for b in sh.buildings:
            bv = _num(wv.get(f"{b['qty_col']}{r}"))
            if bv:
                row.phase_qty[b["label"]] = bv
        # contractor items / lump sums / delivery
        # accommodation / transport / scissor lifts …: a monthly rate in the thousands, or (a blank) no rate yet
        monthly_word = bool(_MONTHLY.search(d)) and (row.uom_norm in ("month", None)) and \
            (rate is None or rate >= 1000) and not re.search(r"\d+\s*(?:mm|m)\b|cable|tray|ladder", d, re.I)
        if contractor or monthly_word or _DELIVERY.search(d) or (_LUMP.search(d) and not has_uom and qty is None):
            row.kind = "contractor"
            m = re.search(r"phase\s*(\d+)", d, re.I)
            row.phase = f"Phase {m.group(1)}" if m else None
            if _DELIVERY.search(d):
                row.kind = "delivery"
                row.rate_basis = "percent"
                pct = qty if row.uom_norm == "%" else None
                row.attrs = {"percent": pct}
                amounts = [x for x in (total, stotal) if x]
                if not amounts:
                    amounts = [v for c, v in ((c, _num(wv.cell(r, c))) for c in range(qty_c + 1, ws.max_column + 1))
                               if v and (pct is None or v != pct)]
                row.attrs["amount"] = max(amounts) if amounts else None
            elif _LUMP.search(d) and not _MONTHLY.search(d):
                row.kind = "lump_sum"
                row.rate_basis = "lump_sum"
                amounts = [x for x in (total, stotal) if x]
                if not amounts:
                    nums = [_num(wv.cell(r, c)) for c in range(desc_c + 1, ws.max_column + 1)]
                    nums = [n for n in nums if n]
                    amounts = [max(nums)] if nums else []
                row.attrs = {"amount": max(amounts) if amounts else None}
                if row.install_rate is None and qty and qty > 100 and not rate:
                    row.attrs["amount"] = row.attrs["amount"] or qty
            else:
                row.rate_basis = "monthly" if (row.uom_norm in ("month", None) and _MONTHLY.search(d)) else \
                    ("lump_sum" if re.search(r"\btools\b", d, re.I) else "monthly")
            sh.rows.append(row)
            continue
        # pricing basis of an item row: section basis, else what was priced
        if basis:
            row.rate_basis = basis
        elif rate is not None and supply is not None:
            row.rate_basis = "supply_and_install"
        elif supply is not None and rate is None:
            row.rate_basis = "supply_only"
        elif rate is not None:
            row.rate_basis = "install_only"
        # nothing priced and no section heading saying how (a blank): the basis stays unknown
        if row.rate_basis == "supply_only" and row.install_rate is not None and not sh.has_supply:
            # supply-only sheets put the supply €/unit in the Rate column
            row.supply_rate, row.install_rate = row.install_rate, None
        row.attrs = unit_rate_attributes(d, finish=row.tag, section=section)
        if row.attrs.get("conflict"):
            row.flags.append("CONFLICTING_DESCRIPTION")
        sh.rows.append(row)
    return sh


# ------------------------------------------------------------------ accessory ratios

_ANCHOR = re.compile(r"luminaire\s*bracket|^junction\s*box|fire\s*rated\s*junction|end\s*feed", re.I)
_ACCESSORY_SECTION = re.compile(r"cabl|wiring|secondary|accessor|spur", re.I)


def _ratios(wb: BoqWorkbook) -> list[dict]:
    """In every Supply & Install section: quantity of each row per anchor row (luminaire bracket / JB / socket …).
    Live formulas (Goodman) and hard values (FRA3H, FRA5) give the same ratios."""
    out = []
    for sh in wb.pricing_sheets:
        by_sec: dict[str, list[BoqRow]] = {}
        for r in sh.rows:
            if r.kind == "item" and r.section and r.qty and (
                    r.rate_basis == "supply_and_install" or (r.rate_basis is None and sh.has_supply)):
                by_sec.setdefault(r.section, []).append(r)
        for sec, rows in by_sec.items():
            if not _ACCESSORY_SECTION.search(sec):
                continue
            anchor = next((r for r in rows if _ANCHOR.search(r.description)), None)
            if not anchor or not anchor.qty:
                continue
            ratios = []
            for r in rows:
                if r is anchor or not r.qty:
                    continue
                ratios.append({"description": r.description, "row": r.row, "uom": r.uom_norm,
                               "ratio": round(r.qty / anchor.qty, 4), "formula": r.qty_formula})
            out.append({"sheet": sh.name, "section": sec, "anchor": anchor.description, "anchor_row": anchor.row,
                        "anchor_qty": anchor.qty, "rows": ratios})
    return out
