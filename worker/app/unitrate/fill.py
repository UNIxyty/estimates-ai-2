"""Fill a unit-rate BOQ blank (2.4): decide every value first (a pure plan), then write only input cells.

Plan, per pricing sheet that takes part:
  item rows        install €/unit and supply €/unit from unit-rate references (match.UnitRateKB); install never from
                   the web, supply from the web only when no reference has it (hook, allowlisted suppliers, EUR).
                   Free-issued sections and "material supply not in scope" leave supply empty; supply-only sheets
                   (Optional supply of brackets) take the supply rate in their Rate column. A line no reference
                   prices stays empty and is listed "not priced, check".
  contractor items months × monthly rate (accommodation / transport / scissor lifts …): the rate from the
                   preferred reference's same item and phase, months from the programme length you gave.
  lump sums        fixation materials / power tools: the preferred reference's lump sum scaled by this sheet's
                   install total against that reference sheet's (marked CHECK).
  delivery         delivery % × this sheet's material-supply total.
  prelims          weekly staff rates × weeks (roles the preferred reference priced, its weeks scaled to your
                   programme) and item rates (mobilisation, insurances …).
  summary          subcontractor name, offer date and validity in the Summary's input cells.
Not touched: Site Attendance, sheets the Summary marks "Not participating" (or that you untick), quantities, every
formula, every non-input cell (when the sheet marks input cells with a fill colour), styles.
Accessory ratios (cable ties / glands / labels … per luminaire bracket or junction box) are compared with the
references: a quantity more than 5 % off is marked CHECK with a message; the client's quantity is never changed.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Callable

import openpyxl

from ..matching.text import normalise_text
from .boq import BoqRow, BoqSheet, BoqWorkbook, guess_package
from .match import Quote, RefMeta, UnitRateKB

RATIO_TOL = 0.05
WEEKS_PER_MONTH = 52 / 12


# ------------------------------------------------------------------ inputs


@dataclass
class Setup:
    """The setup card's answers (unit-rate runs ask no hourly rate)."""
    phase_weeks: dict[str, float] = field(default_factory=dict)   # {"Phase 1": 43.3, …} or {"Programme": 52}
    sheet_weeks: dict[str, float] = field(default_factory=dict)   # a package with its own duration:
                                                                  # {"Lighting Install": 52, "Lighting External · Phase 1": 4}
    prelims_weeks: float | None = None                            # default: the whole programme
    participating: dict[str, bool] = field(default_factory=dict)  # Summary cost line → takes part
    material_supply: bool = True
    delivery_pct: float = 12.0
    offer_date: str | None = None                                 # ISO
    validity_date: str | None = None                              # ISO
    subcontractor: str = "M.G.S. IT LLC"

    @property
    def programme_weeks(self) -> float | None:
        return round(sum(self.phase_weeks.values()), 2) if self.phase_weeks else None


@dataclass
class ExtraRef:
    """A non-item reference line: prelim, contractor item, lump sum."""
    meta: RefMeta
    kind: str                 # prelim | contractor | lump_sum
    description: str
    rate: float | None = None
    qty: float | None = None
    uom: str | None = None
    basis: str | None = None
    phase: str | None = None
    sheet: str | None = None
    row: int | None = None
    sheet_install_total: float | None = None   # lump sums: the reference sheet's install total
    anchor_weeks: float | None = None          # prelims: that reference's Site Manager weeks


@dataclass
class Line:
    kind: str                 # item | contractor | lump_sum | delivery | prelim
    sheet: str
    row: int
    description: str
    section: str | None = None
    qty: float | None = None
    uom: str | None = None
    basis: str | None = None
    install_rate: float | None = None
    supply_rate: float | None = None
    amount: float | None = None              # lump sum / delivery money
    install: Quote | None = None
    supply: Quote | None = None
    note: str | None = None
    flags: list[str] = field(default_factory=list)
    writes: list[dict] = field(default_factory=list)   # [{cell, value, role, input}]
    source: dict | None = None               # non-item lines: the reference used

    @property
    def priced(self) -> bool:
        return self.install_rate is not None or self.supply_rate is not None or self.amount is not None

    @property
    def install_total(self) -> float:
        return round((self.qty or 0) * (self.install_rate or 0), 2) if self.kind in ("item", "contractor", "prelim") \
            else 0.0

    @property
    def supply_total(self) -> float:
        return round((self.qty or 0) * (self.supply_rate or 0), 2) if self.kind == "item" else 0.0


@dataclass
class Plan:
    lines: list[Line] = field(default_factory=list)
    summary_writes: list[dict] = field(default_factory=list)
    skipped_sheets: list[dict] = field(default_factory=list)   # [{sheet, reason}]
    checks: list[dict] = field(default_factory=list)           # accessory-ratio messages
    sheet_totals: dict[str, dict] = field(default_factory=dict)

    @property
    def not_priced(self) -> list[Line]:
        return [ln for ln in self.lines if ln.kind == "item" and "NO PRICE" in ln.flags]


# ------------------------------------------------------------------ helpers


def _phase_no(text: str | None) -> int | None:
    m = re.search(r"phase\s*(\d+)", text or "", re.I)
    return int(m.group(1)) if m else None


def _contractor_class(text: str) -> str:
    t = text.lower()
    for k, rx in (("accommodation", r"accom\w*dation"), ("transport", r"transport"),
                  ("lift", r"scissor|lift|mewp|cherry"), ("tools", r"tools")):
        if re.search(rx, t):
            return k
    return normalise_text(text)[:40]


def _lump_class(text: str) -> str:
    t = text.lower()
    fix, tools = bool(re.search(r"fixation", t)), bool(re.search(r"tools", t))
    return "fixation+tools" if fix and tools else "fixation" if fix else "tools" if tools else normalise_text(t)[:40]


def _w(cell: str | None, value: Any, role: str, hint: bool | None) -> list[dict]:
    return [{"cell": cell, "value": value, "role": role, "input": hint}] if cell else []


def _row_qty(r: BoqRow) -> float | None:
    if r.qty is not None:
        return r.qty
    if r.phase_qty:
        return round(sum(r.phase_qty.values()), 6)
    return None


def _excluded_sheets(wb: BoqWorkbook, setup: Setup) -> dict[str, str]:
    out: dict[str, str] = {}
    for ln in wb.summary_lines:
        target = (ln.get("link") or {}).get("sheet")
        if not target:
            continue
        if ln.get("not_participating"):
            out[target] = f"Summary says “Not participating” ({ln['label']})"
        elif setup.participating and setup.participating.get(ln["label"]) is False:
            out[target] = f"unticked in setup ({ln['label']})"
    return out


def _optional_supply_sheets(wb: BoqWorkbook) -> set[str]:
    return {(ln.get("link") or {}).get("sheet") for ln in wb.summary_lines
            if ln.get("optional") and ln.get("supply") and (ln.get("link") or {}).get("sheet")}


# ------------------------------------------------------------------ planning


class Planner:
    def __init__(self, wb: BoqWorkbook, kb: UnitRateKB, setup: Setup, *, extras: list[ExtraRef] | None = None,
                 ref_ratios: list[tuple[RefMeta, dict]] | None = None,
                 web_supply: Callable[[BoqRow], dict | None] | None = None):
        self.wb, self.kb, self.setup = wb, kb, setup
        self.extras = extras or []
        self.ref_ratios = ref_ratios or []
        self.web_supply = web_supply

    # -- preference among non-item references (same rule as the matcher)
    def _best(self, refs: list[ExtraRef]) -> ExtraRef | None:
        return min(refs, key=lambda e: self.kb.rank(e.meta)) if refs else None

    def plan(self) -> Plan:
        p = Plan()
        excluded = _excluded_sheets(self.wb, self.setup)
        optional_supply = _optional_supply_sheets(self.wb)
        for sh in self.wb.sheets:
            if sh.kind == "attendance":
                p.skipped_sheets.append({"sheet": sh.name, "reason": "Site Attendance is kept as issued"})
            elif sh.kind == "takeoff":
                p.skipped_sheets.append({"sheet": sh.name, "reason": "CostX take-off (quantities only)"})
        for sh in self.wb.pricing_sheets:
            if sh.name in excluded:
                p.skipped_sheets.append({"sheet": sh.name, "reason": excluded[sh.name]})
                continue
            self._plan_sheet(sh, p, supply_only_sheet=sh.name in optional_supply)
        self._plan_prelims(p, excluded)
        self._plan_summary(p)
        self._ratio_checks(p, excluded)
        return p

    # -- one pricing sheet
    def _plan_sheet(self, sh: BoqSheet, p: Plan, *, supply_only_sheet: bool) -> None:
        lines: list[Line] = []
        for r in sh.rows:
            if r.kind == "item":
                lines.append(self._item(sh, r, supply_only_sheet=supply_only_sheet))
        install_total = sum(ln.install_total for ln in lines)
        supply_total = sum(ln.supply_total for ln in lines)
        for r in sh.rows:
            if r.kind == "contractor":
                lines.append(self._contractor(sh, r))
            elif r.kind == "lump_sum":
                lines.append(self._lump(sh, r, install_total))
        for r in sh.rows:
            if r.kind == "delivery":
                lines.append(self._delivery(sh, r, supply_total))
        p.lines.extend(lines)
        p.sheet_totals[sh.name] = {
            "install": round(sum(ln.install_total for ln in lines), 2),
            "supply": round(sum(ln.supply_total for ln in lines), 2),
            "other": round(sum(ln.amount or 0 for ln in lines if ln.kind in ("lump_sum", "delivery")), 2)}
        p.sheet_totals[sh.name]["total"] = round(sum(p.sheet_totals[sh.name].values()), 2)

    def _item(self, sh: BoqSheet, r: BoqRow, *, supply_only_sheet: bool) -> Line:
        qty = _row_qty(r)
        basis = r.rate_basis
        if supply_only_sheet or (basis == "supply_only" and not sh.has_supply):
            basis = "supply_only"
        ln = Line(kind="item", sheet=sh.name, row=r.row, description=r.description, section=r.section, qty=qty,
                  uom=r.uom_norm or r.uom, basis=basis)
        free_issue = bool((r.attrs or {}).get("free_issue")) or basis == "install_only"
        want_install = basis != "supply_only"
        want_supply = self.setup.material_supply and not free_issue and (sh.has_supply or basis == "supply_only")
        if want_install:
            ln.install = self.kb.quote(r.description, uom=r.uom_norm, section=r.section, which="install_rate",
                                       attrs=r.attrs, sheet=sh.name)
            ln.install_rate = ln.install.rate
        if want_supply:
            q = self.kb.quote(r.description, uom=r.uom_norm, section=r.section, which="supply_rate", attrs=r.attrs,
                              sheet=sh.name)
            if q.rate is None and ln.install and ln.install.source is not None and \
                    ln.install.source.supply_rate is None and r.rate_basis is None:
                q = Quote(None, None, "high", None, [], None)
                want_supply = False  # the same item is install-only in the references (terminations, labels …)
            elif q.rate is None and self.web_supply:
                web = self.web_supply(r)
                if web and web.get("unit_price"):
                    q = Quote(float(web["unit_price"]), "web", "low", None, [], f"supplier price: {web.get('url', '')}",
                              check=True)
                    ln.source = {"web": web}
                    ln.flags.append("WEB")
            ln.supply = q if want_supply else None
            ln.supply_rate = q.rate
        notes = [q.note for q in (ln.install, ln.supply) if q and q.note]
        ln.note = "; ".join(dict.fromkeys(notes)) or None
        if any(q and q.check for q in (ln.install, ln.supply)):
            ln.flags.append("CHECK")
        if (want_install and ln.install_rate is None) or (want_supply and ln.supply_rate is None):
            if not ln.priced:
                ln.flags.append("NO PRICE")
                ln.note = "not priced, check: no unit-rate reference has this item"
            else:
                ln.flags.append("CHECK")
                missing = "install" if want_install and ln.install_rate is None else "supply"
                ln.note = "; ".join(x for x in (ln.note, f"no {missing} rate in the references") if x)
        # cells: rate (install, or supply on supply-only sheets), supply, and totals that are inputs (not formulas)
        if basis == "supply_only" and not sh.has_supply:
            ln.writes += _w(r.cells.get("rate"), ln.supply_rate, "rate", r.input_hint.get("rate"))
            ln.writes += _w(r.cells.get("total"), None if ln.supply_rate is None or qty is None
                            else round(qty * ln.supply_rate, 2), "total", r.input_hint.get("total"))
        else:
            ln.writes += _w(r.cells.get("rate"), ln.install_rate, "rate", r.input_hint.get("rate"))
            ln.writes += _w(r.cells.get("total"), None if ln.install_rate is None or qty is None
                            else round(qty * ln.install_rate, 2), "total", r.input_hint.get("total"))
            if sh.has_supply:
                ln.writes += _w(r.cells.get("supply"), ln.supply_rate, "supply", r.input_hint.get("supply"))
                ln.writes += _w(r.cells.get("supply_total"), None if ln.supply_rate is None or qty is None
                                else round(qty * ln.supply_rate, 2), "supply_total", r.input_hint.get("supply_total"))
        # layout C: a phase € cell that is an input (not a formula) gets phase qty × rate
        if sh.phases and ln.install_rate is not None:
            for ph in sh.phases:
                key = ph["label"] + (f" {ph['area']}" if ph.get("area") else "")
                if ph.get("money_col") and r.phase_qty.get(key):
                    ln.writes += _w(f"{ph['money_col']}{r.row}", round(r.phase_qty[key] * ln.install_rate, 2),
                                    "phase_money", None)
        ln.writes = [w for w in ln.writes if w["value"] is not None]
        return ln

    def _contractor(self, sh: BoqSheet, r: BoqRow) -> Line:
        cls = _contractor_class(r.description)
        phase = _phase_no(r.description) or _phase_no(r.phase)
        cands = [e for e in self.extras if e.kind == "contractor" and _contractor_class(e.description) == cls]
        best = self._best(cands)
        same_file = [e for e in cands if best and e.meta.file_id == best.meta.file_id]
        pick = next((e for e in same_file if phase and _phase_no(e.phase or e.description) == phase), None) or \
            (same_file[0] if same_file else None)
        qty = r.qty
        ln = Line(kind="contractor", sheet=sh.name, row=r.row, description=r.description, section=r.section,
                  uom=r.uom_norm or "month", basis=r.rate_basis or "monthly")
        if qty is None and (r.rate_basis or "monthly") == "monthly":
            weeks = self._weeks_for_phase(phase, sh.name)
            if weeks:
                qty = max(1, round(weeks / WEEKS_PER_MONTH))
                ln.writes += _w(r.cells.get("qty"), qty, "qty", None)
                ln.note = f"{qty} month(s) from the programme ({weeks:g} weeks)"
        ln.qty = qty
        if pick is None or pick.rate is None:
            ln.flags.append("NO PRICE")
            ln.note = "not priced, check: no reference prices this contractor item"
            return ln
        ln.install_rate = pick.rate
        ln.source = {"file_name": pick.meta.file_name, "sheet": pick.sheet, "row": pick.row, "rate": pick.rate,
                     "description": pick.description}
        alts = sorted({(e.meta.file_name, e.rate) for e in cands if e.rate})
        ln.note = "; ".join(x for x in (ln.note, f"€{pick.rate:,.0f}/{'month' if ln.basis == 'monthly' else 'item'} "
                                                 f"from {pick.meta.file_name}" +
                                        (f" (others: {', '.join(f'{n} €{v:,.0f}' for n, v in alts[:4])})"
                                         if len({v for _, v in alts}) > 1 else "")) if x)
        ln.writes += _w(r.cells.get("rate"), pick.rate, "rate", r.input_hint.get("rate"))
        if qty is not None:
            ln.writes += _w(r.cells.get("total"), round(qty * pick.rate, 2), "total", r.input_hint.get("total"))
        ln.writes = [w for w in ln.writes if w["value"] is not None]
        return ln

    def _weeks_for_phase(self, phase: int | None, sheet: str | None = None) -> float | None:
        sw = self.setup.sheet_weeks
        if sheet and phase is not None and f"{sheet} · Phase {phase}" in sw:
            return sw[f"{sheet} · Phase {phase}"]
        if sheet and sheet in sw:
            return sw[sheet]
        pw = self.setup.phase_weeks
        if not pw:
            return None
        if phase is not None:
            for k, v in pw.items():
                if _phase_no(k) == phase:
                    return v
        return self.setup.programme_weeks if phase is None else None

    def _lump(self, sh: BoqSheet, r: BoqRow, install_total: float) -> Line:
        cls = _lump_class(r.description)
        ln = Line(kind="lump_sum", sheet=sh.name, row=r.row, description=r.description, section=r.section,
                  qty=r.qty, basis="lump_sum")
        # "Fixation materials, power tools" on one line = the two lump sums the references keep apart
        parts = ["fixation", "tools"] if cls == "fixation+tools" else [cls]
        picks = []
        for part in parts:
            cands = [e for e in self.extras if e.kind == "lump_sum" and e.rate and _lump_class(e.description) == part]
            if cls == "fixation+tools":
                cands = [e for e in self.extras if e.kind == "lump_sum" and e.rate and
                         _lump_class(e.description) == cls] or cands
            best = self._best(cands)
            if best is not None and best not in picks:
                picks.append(best)
        if not picks:
            ln.flags.append("NO PRICE")
            ln.note = "not priced, check: no reference has this lump sum"
            return ln
        amount, notes = 0.0, []
        for best in picks:
            if best.sheet_install_total and install_total:
                a = round(best.rate * install_total / best.sheet_install_total, -2)
                notes.append(f"{best.meta.file_name}: {best.description[:30]} €{best.rate:,.0f} on "
                             f"€{best.sheet_install_total:,.0f} of install work, scaled to this sheet's "
                             f"€{install_total:,.0f}")
            else:
                a = best.rate
                notes.append(f"€{best.rate:,.0f} as in {best.meta.file_name}")
            amount += a
        ln.amount = amount
        ln.note = "; ".join(notes)
        ln.flags.append("CHECK")
        ln.source = {"file_name": picks[0].meta.file_name, "sheet": picks[0].sheet, "row": picks[0].row,
                     "amount": picks[0].rate}
        ln.writes = _lump_writes(sh, r, amount)
        return ln

    def _delivery(self, sh: BoqSheet, r: BoqRow, supply_total: float) -> Line:
        pct = self.setup.delivery_pct
        ln = Line(kind="delivery", sheet=sh.name, row=r.row, description=r.description, section=r.section,
                  basis="percent")
        if not self.setup.material_supply or not supply_total:
            ln.note = "no material supply priced on this sheet, so no delivery"
            return ln
        amount = round(supply_total * pct / 100, 2)
        ln.amount = amount
        ln.note = f"{pct:g} % of the material-supply total €{supply_total:,.2f}"
        writes: list[dict] = []
        if r.uom_norm == "%":
            writes += _w(r.cells.get("qty"), pct, "qty", None)
        for role in ("supply_total", "total"):
            if r.cells.get(role):
                writes += _w(r.cells[role], amount, role, r.input_hint.get(role))
                break
        else:
            writes += _w(r.cells.get("rate"), amount, "rate", r.input_hint.get("rate"))
        ln.writes = writes
        return ln

    # -- prelims
    def _plan_prelims(self, p: Plan, excluded: dict[str, str]) -> None:
        if not self.wb.prelims:
            return
        weeks = self.setup.prelims_weeks or self.setup.programme_weeks
        # one reference prelims sheet per blank prelims sheet: which roles are priced at all is the estimator's
        # choice per project, so its role set and weeks are copied, not the union of every reference
        groups: dict[tuple[str, str | None], list[ExtraRef]] = {}
        for e in self.extras:
            if e.kind == "prelim" and e.rate:
                groups.setdefault((e.meta.file_id, e.sheet), []).append(e)
        chosen: dict[str, list[ExtraRef]] = {}
        for sheet in dict.fromkeys(pr["sheet"] for pr in self.wb.prelims):
            pkg = guess_package(sheet) or self.wb.package
            if groups:
                key = min(groups, key=lambda k: (0 if (guess_package(k[1] or "") or groups[k][0].meta.package) == pkg
                                                 else 1, self.kb.rank(groups[k][0].meta)))
                chosen[sheet] = groups[key]
        for pr in self.wb.prelims:
            if pr["sheet"] in excluded:
                continue
            refs = chosen.get(pr["sheet"], [])
            ln = Line(kind="prelim", sheet=pr["sheet"], row=pr["row"], description=pr["description"],
                      section="Preliminaries", qty=pr["qty"], uom=pr["uom"], basis=pr["rate_basis"])
            if pr.get("rate_is_formula"):
                continue
            key = normalise_text(pr["description"])
            cands = [e for e in refs if normalise_text(e.description) == key and (e.basis == pr["rate_basis"])]
            best = self._best(cands)
            if best is None:
                continue  # roles / items the references never price stay empty (not every prelim line is priced)
            ln.install_rate = best.rate
            ln.source = {"file_name": best.meta.file_name, "sheet": best.sheet, "row": best.row, "rate": best.rate,
                         "qty": best.qty}
            if pr["rate_basis"] == "weekly":
                q = pr["qty"]
                if q is None and weeks and best.qty and best.anchor_weeks:
                    q = round(weeks * best.qty / best.anchor_weeks)
                    ln.writes += _w(pr["cells"].get("qty"), q, "qty", None)
                    ln.note = (f"€{best.rate:,.0f}/week × {q:g} weeks ({best.meta.file_name} ran this role "
                               f"{best.qty:g} of {best.anchor_weeks:g} weeks)")
                ln.qty = q
                ln.flags.append("CHECK")
            else:
                if ln.qty is None and best.qty:
                    ln.qty = best.qty
                    ln.writes += _w(pr["cells"].get("qty"), best.qty, "qty", None)
                ln.note = f"€{best.rate:,.0f} as in {best.meta.file_name}"
            ln.writes += _w(pr["cells"].get("rate"), best.rate, "rate", None)
            if ln.qty is not None:
                ln.writes += _w(pr["cells"].get("total"), round(ln.qty * best.rate, 2), "total", None)
            p.lines.append(ln)
        tot: dict[str, float] = {}
        for ln in p.lines:
            if ln.kind == "prelim":
                tot[ln.sheet] = tot.get(ln.sheet, 0) + ln.install_total
        for s, v in tot.items():
            p.sheet_totals[s] = {"install": round(v, 2), "supply": 0.0, "other": 0.0, "total": round(v, 2)}

    # -- summary fields
    def _plan_summary(self, p: Plan) -> None:
        sub = self.wb.subcontractor or {}
        vals = {"name": self.setup.subcontractor, "offer_date": self.setup.offer_date,
                "validity": self.setup.validity_date}
        for key, v in vals.items():
            if v and sub.get(key, {}).get("cell"):
                p.summary_writes.append({"sheet": "Summary", "cell": sub[key]["cell"], "value": v, "role": key})

    # -- accessory ratios
    def _ratio_checks(self, p: Plan, excluded: dict[str, str]) -> None:
        if not self.ref_ratios:
            return
        ref_by: dict[tuple[str, str], list[tuple[RefMeta, float]]] = {}
        for meta, rr in self.ref_ratios:
            a = _anchor_key(rr["anchor"])
            for x in rr["rows"]:
                ref_by.setdefault((a, normalise_text(x["description"])), []).append((meta, x["ratio"]))
        by_row = {(ln.sheet, ln.row): ln for ln in p.lines}
        for rr in self.wb.ratios:
            if rr["sheet"] in excluded:
                continue
            a = _anchor_key(rr["anchor"])
            for x in rr["rows"]:
                refs = ref_by.get((a, normalise_text(x["description"])))
                if not refs:
                    continue
                meta, want = min(refs, key=lambda t: self.kb.rank(t[0]))
                if want and abs(x["ratio"] - want) / want > RATIO_TOL:
                    msg = (f"{x['description'][:50]}: {x['ratio']:g} per {rr['anchor'][:30]} here, "
                           f"{want:g} in {meta.file_name} — quantity kept as issued, please check")
                    p.checks.append({"sheet": rr["sheet"], "row": x["row"], "section": rr["section"], "message": msg,
                                     "ratio": x["ratio"], "reference_ratio": want, "reference": meta.file_name})
                    ln = by_row.get((rr["sheet"], x["row"]))
                    if ln is not None:
                        if "CHECK" not in ln.flags:
                            ln.flags.append("CHECK")
                        ln.note = "; ".join(t for t in (ln.note, msg) if t)


def _anchor_key(text: str) -> str:
    t = text.lower()
    for k in ("luminaire bracket", "fire rated junction", "junction box", "end feed"):
        if k in t:
            return k
    return normalise_text(t)[:30]


def _lump_writes(sh: BoqSheet, r: BoqRow, amount: float) -> list[dict]:
    """Where the lump sum goes: Rate (qty 1 × rate), else the first input money cell (phase €, total, supply total)."""
    if r.qty and r.cells.get("rate"):
        return _w(r.cells["rate"], round(amount / r.qty, 2), "rate", r.input_hint.get("rate"))
    out: list[dict] = []
    for ph in sh.phases:
        if ph.get("money_col"):
            out += _w(f"{ph['money_col']}{r.row}", amount, "phase_money", None)
            return out
    for role in ("supply_total", "total"):
        if r.cells.get(role):
            return _w(r.cells[role], amount, role, r.input_hint.get(role))
    return out


def plan_fill(wb: BoqWorkbook, kb: UnitRateKB, setup: Setup, **kw) -> Plan:
    return Planner(wb, kb, setup, **kw).plan()


# ------------------------------------------------------------------ references → extras / ratios


def extras_from_workbook(wb: BoqWorkbook, meta: RefMeta) -> list[ExtraRef]:
    out: list[ExtraRef] = []
    anchor = {}
    for p in wb.prelims:
        if p["rate_basis"] == "weekly" and p["qty"] and re.search(r"site\s*manager", p["description"], re.I):
            anchor[p["sheet"]] = p["qty"]
    for p in wb.prelims:
        if p["rate"]:
            out.append(ExtraRef(meta=meta, kind="prelim", description=p["description"], rate=p["rate"], qty=p["qty"],
                                uom=p["uom"], basis=p["rate_basis"], sheet=p["sheet"], row=p["row"],
                                anchor_weeks=anchor.get(p["sheet"])))
    for sh in wb.pricing_sheets:
        install_total = sum((r.install_rate or 0) * (_row_qty(r) or 0) for r in sh.rows if r.kind == "item")
        for r in sh.rows:
            if r.kind == "contractor" and r.install_rate:
                out.append(ExtraRef(meta=meta, kind="contractor", description=r.description, rate=r.install_rate,
                                    qty=r.qty, basis=r.rate_basis, phase=r.phase, sheet=sh.name, row=r.row))
            elif r.kind == "lump_sum" and (r.attrs or {}).get("amount"):
                out.append(ExtraRef(meta=meta, kind="lump_sum", description=r.description, rate=r.attrs["amount"],
                                    sheet=sh.name, row=r.row, sheet_install_total=install_total or None))
    return out


# ------------------------------------------------------------------ writing


@dataclass
class WriteReport:
    written: int = 0
    skipped_formula: list[str] = field(default_factory=list)
    skipped_not_input: list[str] = field(default_factory=list)
    skipped_filled: list[str] = field(default_factory=list)
    skipped_merged: list[str] = field(default_factory=list)


def _merged_non_anchor(ws) -> set[str]:
    out: set[str] = set()
    for rng in ws.merged_cells.ranges:
        for i, row in enumerate(ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col,
                                             max_col=rng.max_col)):
            for j, cell in enumerate(row):
                if i or j:
                    out.add(cell.coordinate)
    return out


def _date_value(cell, iso: str):
    try:
        d = date.fromisoformat(iso[:10])
    except ValueError:
        return iso
    fmt = (cell.number_format or "").lower()
    if any(k in fmt for k in ("d", "y")) and fmt != "general" and "@" not in fmt:
        return datetime(d.year, d.month, d.day)
    return f"{d.day} {d.strftime('%B')} {d.year}"


def write_plan(src_path: str, dst_path: str, plan: Plan, wb: BoqWorkbook) -> WriteReport:
    """Write the plan's values into a copy of the blank. A cell is written only when it is not a formula, not a merged
    non-anchor, empty or 0, and (on sheets that mark input cells with a fill colour) an input cell."""
    os.makedirs(os.path.dirname(dst_path) or ".", exist_ok=True)
    if os.path.abspath(src_path) != os.path.abspath(dst_path):
        shutil.copyfile(src_path, dst_path)
    book = openpyxl.load_workbook(dst_path)
    rep = WriteReport()
    # a column that marks its input cells with the light-blue fill is written only in those cells; a Rate / Total
    # column that never uses the fill (FR12X, Small power) is the input column itself
    # (item rows and contractor / lump-sum rows separately: FR12X colours only its contractor-item rates)
    blue_cols = {(s.name, role, r.kind == "item") for s in wb.pricing_sheets for r in s.rows
                 for role, hint in r.input_hint.items() if hint}
    kind_at = {(s.name, r.row): r.kind == "item" for s in wb.pricing_sheets for r in s.rows}
    merged: dict[str, set[str]] = {}
    items = [(ln.sheet, w) for ln in plan.lines for w in ln.writes] + [(w["sheet"], w) for w in plan.summary_writes]
    for sheet, w in items:
        if sheet not in book.sheetnames:
            continue
        ws = book[sheet]
        coord = w["cell"]
        if coord in merged.setdefault(sheet, _merged_non_anchor(ws)):
            rep.skipped_merged.append(f"{sheet}!{coord}")
            continue
        cell = ws[coord]
        v = cell.value
        if isinstance(v, str) and v.startswith("="):
            rep.skipped_formula.append(f"{sheet}!{coord}")
            continue
        row_no = int(re.sub(r"^[A-Z]+", "", coord)) if re.match(r"^[A-Z]+\d+$", coord) else None
        if w["role"] in ("rate", "supply", "total", "supply_total") and w.get("input") is False and \
                (sheet, w["role"], kind_at.get((sheet, row_no), True)) in blue_cols:
            rep.skipped_not_input.append(f"{sheet}!{coord}")
            continue
        summary_field = w["role"] in ("name", "offer_date", "validity")
        if not summary_field and v not in (None, "", 0, 0.0):
            rep.skipped_filled.append(f"{sheet}!{coord}")
            continue
        if summary_field and v not in (None, "") and w["role"] == "name" and str(v).strip():
            rep.skipped_filled.append(f"{sheet}!{coord}")  # a name already typed in stays
            continue
        cell.value = _date_value(cell, w["value"]) if w["role"] in ("offer_date", "validity") else w["value"]
        rep.written += 1
    try:
        book.calculation.fullCalcOnLoad = True
    except AttributeError:
        pass
    book.save(dst_path)
    return rep


# ------------------------------------------------------------------ recalculation check (LibreOffice headless)

_ERR = re.compile(r"^#(REF!|VALUE!|DIV/0!|NAME\?|N/A|NUM!|NULL!)")


def recalculate(path: str, timeout: int = 240) -> str | None:
    """Recalculate every formula with LibreOffice headless; returns the recalculated copy (None without soffice)."""
    soffice = os.environ.get("SOFFICE_BIN") or shutil.which("soffice")
    if not soffice:
        return None
    out_dir = tempfile.mkdtemp(prefix="recalc_")
    profile = os.path.join(tempfile.gettempdir(), f"lo_{uuid.uuid4().hex}")
    src = os.path.join(out_dir, "in_" + os.path.basename(path))
    shutil.copyfile(path, src)
    subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--calc", "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(out_dir, "out"), src],
                   check=True, timeout=timeout, capture_output=True)
    res = os.path.join(out_dir, "out", os.path.basename(src))
    return res if os.path.exists(res) else None


def verify(filled_path: str, plan: Plan, wb: BoqWorkbook) -> dict:
    """Recalculate the filled workbook and check it: no error values in totals, every written row's formula total
    equals qty × rate, and the sheet totals the Summary links to. Returns {ok, recalculated, totals, problems}."""
    rc = recalculate(filled_path)
    if rc is None:
        return {"ok": None, "recalculated": False, "problems": ["LibreOffice is not available: totals not verified"],
                "totals": {}}
    vals = openpyxl.load_workbook(rc, data_only=True)
    problems: list[str] = []
    totals: dict[str, dict] = {}
    for sh in wb.pricing_sheets:
        if sh.name not in vals.sheetnames:
            continue
        ws = vals[sh.name]
        t = {}
        for k, coord in sh.total_cells.items():
            v = ws[coord].value
            if isinstance(v, str) and _ERR.match(v):
                problems.append(f"{sh.name}!{coord} shows {v}")
            elif isinstance(v, (int, float)):
                t[k] = round(float(v), 2)
        totals[sh.name] = t
    rows = {(ln.sheet, ln.row): ln for ln in plan.lines if ln.kind == "item"}
    by_sheet = {s.name: s for s in wb.pricing_sheets}
    for (sheet, r), ln in rows.items():
        sh = by_sheet.get(sheet)
        if not sh or sheet not in vals.sheetnames or ln.install_rate is None or ln.qty is None:
            continue
        brow = next((x for x in sh.rows if x.row == r), None)
        coord = (brow.cells.get("total") if brow else None)
        if not coord or ln.basis == "supply_only":
            continue
        v = vals[sheet][coord].value
        want = round(ln.qty * ln.install_rate, 2)
        if isinstance(v, (int, float)) and abs(float(v) - want) > max(0.02, 0.001 * abs(want)) and not sh.phases:
            problems.append(f"{sheet}!{coord} = {v:g}, expected {want:g} (qty {ln.qty:g} × {ln.install_rate:g})")
    summary: dict[str, float] = {}
    if "Summary" in vals.sheetnames:
        ws = vals["Summary"]
        for line in wb.summary_lines:
            if line.get("cell"):
                v = ws[line["cell"]].value
                if isinstance(v, str) and _ERR.match(v):
                    problems.append(f"Summary!{line['cell']} shows {v}")
                elif isinstance(v, (int, float)):
                    summary[line["label"]] = round(float(v), 2)
    return {"ok": not problems, "recalculated": True, "totals": totals, "summary": summary, "problems": problems[:50]}

