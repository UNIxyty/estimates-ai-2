"""Deterministic calculation-logic sentences per sheet / section.

Each record has a stable `logic_key` (e.g. `sheet:<name>:labour`) so a user's override of the sentence or
numbers survives re-analysis. Formulas (`=G12*D12`, `=SUM(J5:J40)`, `=J41*0.2`) confirm the logic when present;
otherwise the numbers are checked across rows.
"""
from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from typing import Any

from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

from ..matching.text import normalise_text
from . import formulas as fx
from .readers import SheetData
from .structure import RowInfo, SheetStructure, WorkbookStructure
from .util import close

LOGIC_KINDS = ("labour", "material", "subtotal", "markup", "rate", "other")

_BASE_LABEL = {"labour": "labour", "material": "materials", "total": "the direct-cost total",
               "mechanisms": "mechanisms", "total_with_markups": "the total including the markups above it"}
_KEYWORD_BASE = [
    (re.compile(r"(socialais nodoklis|soc nodoklis|darba devej|social tax|social security|arbejdsgiver|"
                r"sotsiaalmaks|sodra|sozialabgab)"), "labour"),
    (re.compile(r"(transport|korsel|materialu|materiali|material)"), "material"),
]
_MARKUP_LABEL = [
    (re.compile(r"virsizdevum|overhead|yleiskulu|uldkulud|накладн|pridetin|administration|narzut"), "Overheads"),
    (re.compile(r"pelna|profit|avance|fortjeneste|daekningsbidrag|gewinn|kasum|voitto|zysk|прибыль|margin"), "Profit"),
    (re.compile(r"socialais nodoklis|soc nodoklis|social tax|social security"), "Employer social tax"),
    (re.compile(r"transport|korsel"), "Transport costs"),
    (re.compile(r"\bpvn\b|\bvat\b|\bmoms\b|\bmva\b|mwst|\bpvm\b|\balv\b|\bkm\b|ндс"), "VAT"),
    (re.compile(r"neparedzet|contingenc|risiko|wagnis|risk"), "Contingency"),
    (re.compile(r"darba aizsardz|insurance|apdrosinas"), "Insurance / safety"),
]


@dataclass
class LogicRecord:
    sheet_name: str | None
    section_title: str | None
    ord: int
    kind: str
    logic_key: str
    sentence: str
    numbers: dict[str, Any] = field(default_factory=dict)
    source: str = "code"

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _fmt(x: float | None, nd: int = 2) -> str:
    if x is None:
        return "?"
    return f"{x:,.{nd}f}".replace(",", " ")


def _pct(x: float) -> str:
    s = f"{x * 100:.2f}".rstrip("0").rstrip(".")
    return f"{s}%"


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", normalise_text(s)).strip("_")[:48] or "x"


def _col_meaning(ss: SheetStructure) -> dict[int, str]:
    return {c.idx: c.meaning for c in ss.columns}


def _formula_example(sd: SheetData | None, rows: list[RowInfo], meaning: str) -> tuple[str | None, dict | None]:
    if sd is None:
        return None, None
    for r in rows:
        cell = r.cells.get(meaning)
        if not cell or not r.formulas.get(meaning):
            continue
        col, row = coordinate_from_string(cell)
        f = sd.formula(row, column_index_from_string(col))
        if f:
            return f"{cell} {f}", fx.shape(f)
    return None, None


def _shape_meanings(shape: dict | None, colmap: dict[int, str], row: int) -> list[str]:
    if not shape:
        return []
    refs = shape.get("refs") or ([shape["ref"]] if shape.get("ref") else [])
    out = []
    for (sheet, c, r) in refs:
        if sheet is None and r == row:
            out.append(colmap.get(c, "?"))
        else:
            out.append("?")
    return out


def derive_logic(ws: WorkbookStructure) -> list[LogicRecord]:
    out: list[LogicRecord] = []
    doc = ws.document
    for ss in ws.sheets:
        sd = doc.sheet(ss.name) if doc is not None else None
        if ss.kind == "norms":
            out.append(LogicRecord(ss.name, None, len(out), "other", f"sheet:{ss.name}:norms",
                                   f"Sheet '{ss.name}' lists labour norms in hours per unit "
                                   f"({sum(1 for r in ss.rows if r.kind == 'item')} rows).",
                                   {"rows": sum(1 for r in ss.rows if r.kind == "item")}))
            continue
        if ss.kind not in ("estimate", "prices"):
            continue
        out.extend(_sheet_logic(ss, sd, ws.currency, start=len(out)))
    for i, rec in enumerate(out):
        rec.ord = i
    return out


def _sheet_logic(ss: SheetStructure, sd: SheetData | None, currency: str | None, start: int) -> list[LogicRecord]:
    recs: list[LogicRecord] = []
    cur = ss.currency or currency or ""
    items = [r for r in ss.rows if r.kind == "item"]
    colmap = _col_meaning(ss)
    name = ss.name
    has = {c.meaning for c in ss.columns}

    # ---- hourly rate(s)
    rates = ss.hourly_rates
    distinct = sorted({round(x["rate"], 4) for x in rates})
    if len(distinct) == 1:
        x = rates[0]
        recs.append(LogicRecord(name, None, 0, "rate", f"sheet:{name}:rate",
                                f"Hourly rate {_fmt(x['rate'])} {cur}/h (from {x['source']}).",
                                {"hourly_rate": x["rate"], "currency": cur, "source": x["source"]}))
    elif len(distinct) > 1:
        recs.append(LogicRecord(name, None, 0, "rate", f"sheet:{name}:rate",
                                f"Hourly rate varies on this sheet: {', '.join(_fmt(r) for r in distinct)} {cur}/h "
                                f"(set per section).", {"hourly_rates": distinct, "currency": cur}))
        for s in ss.sections:
            if s.hourly_rate is not None:
                recs.append(LogicRecord(name, s.title, 0, "rate", f"sheet:{name}:section:{_slug(s.title)}:rate",
                                        f"Section '{s.title}': hourly rate {_fmt(s.hourly_rate)} {cur}/h.",
                                        {"hourly_rate": s.hourly_rate, "currency": cur,
                                         "rows": [s.row_start, s.row_end]}))

    # ---- labour
    ex_ul, sh_ul = _formula_example(sd, items, "unit_labour")
    ex_tl, sh_tl = _formula_example(sd, items, "total_labour")
    checked = matching = 0
    for r in items:
        ul, nh = r.values.get("unit_labour"), r.values.get("norm_h")
        rate = r.values.get("hourly_rate") or _section_rate(ss, r)
        if ul and nh and rate:
            checked += 1
            matching += close(ul, nh * rate, rel=0.01, abs_tol=0.011)
    ul_parts = _shape_meanings(sh_ul, colmap, _row_of(ex_ul))
    tl_parts = _shape_meanings(sh_tl, colmap, _row_of(ex_tl))
    rate_txt = (f"hourly rate {_fmt(distinct[0])} {cur}/h" if len(distinct) == 1 else
                ("the section's hourly rate" if distinct else "the hourly rate"))
    numbers: dict[str, Any] = {"currency": cur, "rows_checked": checked, "rows_matching": matching}
    if distinct:
        numbers["hourly_rate"] = distinct[0] if len(distinct) == 1 else distinct
    sentence = None
    confirmed = None
    if sh_ul and sh_ul.get("kind") == "product" and set(ul_parts) >= {"norm_h", "hourly_rate"}:
        confirmed = "formulas"
    elif checked and matching / checked >= 0.8:
        confirmed = "values"
    if "norm_h" in has and ("unit_labour" in has or "total_labour" in has):
        if "unit_labour" in has:
            sentence = f"Labour per unit = norm (h) × {rate_txt}; row labour = qty × unit labour."
        else:
            sentence = f"Row labour = qty × norm (h) × {rate_txt}."
            if sh_tl and sh_tl.get("kind") == "product" and {"qty", "norm_h"} <= set(tl_parts):
                confirmed = "formulas"
    elif "unit_labour" in has or "total_labour" in has:
        sentence = "Labour is priced per unit (no norm hours given); row labour = qty × unit labour."
    if sentence:
        if sh_ul and sh_ul.get("kind") == "product" and "ROUND" in (ex_ul or ""):
            m = re.search(r"ROUND\w*\(.*,\s*(\d+)\)", ex_ul or "")
            if m:
                sentence += f" Unit labour is rounded to {m.group(1)} decimals."
                numbers["rounding"] = int(m.group(1))
        if confirmed == "formulas":
            sentence += f" Confirmed by formulas (e.g. {ex_ul or ex_tl})."
        elif confirmed == "values":
            sentence += f" Checked on {matching} of {checked} rows."
        numbers["confirmed_by"] = confirmed
        if ex_ul or ex_tl:
            numbers["formula_example"] = ex_ul or ex_tl
        recs.append(LogicRecord(name, None, 0, "labour", f"sheet:{name}:labour", sentence, numbers))

    # ---- material
    ex_tm, sh_tm = _formula_example(sd, items, "total_material")
    tm_parts = _shape_meanings(sh_tm, colmap, _row_of(ex_tm))
    if "unit_material" in has or "total_material" in has:
        n_src = sum(1 for r in items if r.source_ref)
        sentence = "Material per unit is the supplier price of the item; row material = qty × unit material."
        m_numbers: dict[str, Any] = {"currency": cur}
        if sh_tm and sh_tm.get("kind") == "product" and {"qty", "unit_material"} <= set(tm_parts):
            sentence += f" Confirmed by formulas (e.g. {ex_tm})."
            m_numbers["confirmed_by"] = "formulas"
            m_numbers["formula_example"] = ex_tm
        if n_src:
            srcs = sorted({r.source_ref for r in items if r.source_ref})
            sentence += f" Price sources are given per row ({', '.join(srcs[:3])}{'…' if len(srcs) > 3 else ''})."
            m_numbers["sources"] = srcs[:20]
        recs.append(LogicRecord(name, None, 0, "material", f"sheet:{name}:material", sentence, m_numbers))

    # ---- unit total / row total composition
    ex_ut, sh_ut = _formula_example(sd, items, "unit_total")
    ex_t, sh_t = _formula_example(sd, items, "total")
    parts = _shape_meanings(sh_ut, colmap, _row_of(ex_ut))
    if sh_ut and sh_ut.get("kind") == "sum_refs" and parts:
        words = [p.replace("unit_", "") for p in parts if p != "?"]
        recs.append(LogicRecord(name, None, 0, "other", f"sheet:{name}:unit_total",
                                f"Unit total = {' + '.join(words)} per unit (e.g. {ex_ut}).",
                                {"parts": words, "formula_example": ex_ut}))
    tparts = _shape_meanings(sh_t, colmap, _row_of(ex_t))
    if sh_t and sh_t.get("kind") == "sum_refs" and tparts:
        words = [p.replace("total_", "") for p in tparts if p != "?"]
        recs.append(LogicRecord(name, None, 0, "other", f"sheet:{name}:row_total",
                                f"Row total = {' + '.join(words)} (e.g. {ex_t}).",
                                {"parts": words, "formula_example": ex_t}))

    # ---- subtotals
    subs = [r for r in ss.rows if r.kind == "subtotal"]
    totals = [r for r in ss.rows if r.kind == "total" and not r.markup]
    if subs or totals:
        ex = None
        for r in subs + totals:
            for m, cell in r.cells.items():
                if r.formulas.get(m) and sd is not None:
                    col, row = coordinate_from_string(cell)
                    f = sd.formula(row, column_index_from_string(col))
                    if f:
                        ex = f"{cell} {f}"
                        break
            if ex:
                break
        if subs:
            sentence = (f"Each of the {len(ss.sections)} sections ends with a subtotal row summing its item rows; "
                        f"the sheet total adds the section subtotals.")
        else:
            sentence = "The sheet total sums all item rows."
        if ex:
            sentence += f" (e.g. {ex})"
        recs.append(LogicRecord(name, None, 0, "subtotal", f"sheet:{name}:subtotal", sentence,
                                {"sections": len(ss.sections), "subtotal_rows": [r.row for r in subs],
                                 "total_rows": [r.row for r in totals], "formula_example": ex,
                                 "totals": ss.totals}))

    # ---- markups
    for mk in ss.markups:
        rec = _markup_logic(ss, sd, mk, cur)
        if rec:
            recs.append(rec)
    return recs


def _row_of(example: str | None) -> int:
    if not example:
        return -1
    m = re.match(r"[A-Z]+(\d+)\b", example)
    return int(m.group(1)) if m else -1


def _section_rate(ss: SheetStructure, r: RowInfo) -> float | None:
    for s in ss.sections:
        if s.row_start <= r.row <= s.row_end:
            return s.hourly_rate
    return ss.hourly_rates[0]["rate"] if len(ss.hourly_rates) == 1 else None


def _markup_logic(ss: SheetStructure, sd: SheetData | None, mk: dict[str, Any], cur: str) -> LogicRecord | None:
    name_raw = mk.get("name") or "Markup"
    n = normalise_text(name_raw)
    label = next((lbl for rx, lbl in _MARKUP_LABEL if rx.search(n)), None)
    pct = mk.get("pct")
    value = mk.get("value")
    base = None
    base_cell = None
    base_value = None
    confirmed = None
    colmap = _col_meaning(ss)
    rows_by = {r.row: r for r in ss.rows}
    pct_cell = mk.get("pct_cell")
    if mk.get("formula"):
        parsed = fx.parse_formula(mk["formula"])
        shape = fx.shape(parsed) if parsed else None
        refs = [ref for ref in (parsed.refs if parsed else []) if ref[0] is None]
        ref_cells = [f"{get_column_letter(c)}{r}" for (_, c, r) in refs]
        base_refs = [(c, r) for (_, c, r), cell in zip(refs, ref_cells) if cell != pct_cell]
        if shape and shape.get("kind") == "scale" and pct is None:
            pct = shape["factor"]
        if shape and shape.get("kind") == "product" and pct_cell and pct is not None:
            pass
        if base_refs:
            kinds = set()
            for c, r in base_refs:
                m = colmap.get(c, "")
                row = rows_by.get(r)
                if row is not None and row.markup:
                    kinds.add("markup")
                elif m in ("total_labour", "unit_labour"):
                    kinds.add("labour")
                elif m in ("total_material", "unit_material"):
                    kinds.add("material")
                elif m in ("total_mechanisms",):
                    kinds.add("mechanisms")
                else:
                    kinds.add("total")
            if "markup" in kinds:
                base = "total_with_markups"
            elif len(kinds) == 1:
                base = kinds.pop()
            else:
                base = "total"
            base_cell = "+".join(f"{get_column_letter(c)}{r}" for c, r in base_refs)
            if sd is not None:
                base_value = sum((_num(sd.get(r, c)) or 0.0) for c, r in base_refs)
            if pct is None and value is not None and base_value:
                pct = value / base_value
            if pct is not None and value is not None and base_value is not None and close(value, pct * base_value,
                                                                                             rel=0.005):
                confirmed = "formula"
    if base is None and pct is not None and value:
        for key, tot_key in (("labour", "total_labour"), ("material", "total_material"), ("total", "total")):
            tv = ss.totals.get(tot_key)
            if tv and close(value, pct * tv, rel=0.005):
                base, base_value, confirmed = key, tv, "values"
                break
    if base is None:
        base = next((b for rx, b in _KEYWORD_BASE if rx.search(n)), "total")
    if pct is None and value is None:
        return None
    pct_txt = _pct(pct) if pct is not None else "a fixed amount"
    what = label or name_raw
    shown = f"{what} ('{name_raw}')" if label and normalise_text(label) not in n else what
    sentence = f"{shown} = {pct_txt} of {_BASE_LABEL.get(base, base)}"
    if base_cell:
        sentence += f" ({base_cell})"
    sentence += "."
    if confirmed == "formula":
        sentence += f" Confirmed by formula {mk.get('value_cell')} {mk.get('formula')}."
    elif confirmed == "values":
        sentence += " Checked against the sheet totals."
    return LogicRecord(ss.name, None, 0, "markup", f"sheet:{ss.name}:markup:{mk.get('key') or _slug(name_raw)}",
                       sentence, {"name": name_raw, "label": label, "pct": pct, "base": base, "base_cell": base_cell,
                                  "base_value": base_value, "value": value, "value_cell": mk.get("value_cell"),
                                  "row": mk.get("row"), "confirmed_by": confirmed, "currency": cur})


def _num(v: Any) -> float | None:
    from .util import parse_number
    return parse_number(v)
