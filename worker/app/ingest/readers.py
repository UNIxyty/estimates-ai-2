"""File readers producing one in-memory model for xlsx / xls / docx / pdf.

xlsx is read with openpyxl in read-only (streaming) mode twice: `data_only=True` for values and cached
formula results (+ bold / number formats), `data_only=False` for formula text. Merged ranges and hidden
rows/columns are not exposed by read-only mode, so they are picked up by a streaming scan of the sheet XML.
Formula cells without a cached value (files written by tools that never calculate) are computed with the
small evaluator in `formulas.py`.
"""
from __future__ import annotations

import datetime as _dt
import logging
import os
import re
import shutil
import subprocess
import tempfile
import uuid
from dataclasses import dataclass, field
from typing import Any, Iterator

from openpyxl.utils.cell import column_index_from_string, get_column_letter

from . import formulas as fx

log = logging.getLogger(__name__)

SOFFICE = os.environ.get("SOFFICE_BIN") or shutil.which("soffice") or "/usr/bin/soffice"
MAX_ROWS = 200_000
MAX_COLS = 200


class IngestError(Exception):
    """Raised for files we can't read. `reason` is shown to the user as-is."""

    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


@dataclass
class Cell:
    value: Any            # formula text ("=D5*G5") for formula cells, else the constant
    cached: Any           # cached/computed value (what Excel shows)
    is_formula: bool
    number_format: str
    bold: bool
    merged: tuple[int, int, int, int] | None   # (min_row, min_col, max_row, max_col) when in a merged range


@dataclass
class SheetData:
    name: str
    idx: int
    rows: list[list[Any]] = field(default_factory=list)          # rows[r-1][c-1] -> cached value
    formulas: dict[tuple[int, int], str] = field(default_factory=dict)   # (row, col) 1-based -> "=..."
    bold: set[tuple[int, int]] = field(default_factory=set)
    number_formats: dict[tuple[int, int], str] = field(default_factory=dict)  # only non-General
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)
    hidden_rows: set[int] = field(default_factory=set)
    hidden_cols: set[int] = field(default_factory=set)
    state: str = "visible"
    source: str = "xlsx"          # xlsx | docx-table | pdf-table
    computed: set[tuple[int, int]] = field(default_factory=set)   # formula cells evaluated by us
    _merge_map: dict[tuple[int, int], tuple[int, int, int, int]] | None = None

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def get(self, row: int, col: int) -> Any:
        """Cached/displayed value at 1-based (row, col)."""
        if row < 1 or col < 1 or row > len(self.rows):
            return None
        r = self.rows[row - 1]
        return r[col - 1] if col <= len(r) else None

    def formula(self, row: int, col: int) -> str | None:
        return self.formulas.get((row, col))

    def is_bold(self, row: int, col: int) -> bool:
        return (row, col) in self.bold

    def number_format(self, row: int, col: int) -> str:
        return self.number_formats.get((row, col), "General")

    def merge_map(self) -> dict[tuple[int, int], tuple[int, int, int, int]]:
        if self._merge_map is None:
            mm: dict[tuple[int, int], tuple[int, int, int, int]] = {}
            for rng in self.merged:
                r1, c1, r2, c2 = rng
                if (r2 - r1 + 1) * (c2 - c1 + 1) > 5000:
                    continue
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        mm[(r, c)] = rng
            self._merge_map = mm
        return self._merge_map

    def merged_range(self, row: int, col: int) -> tuple[int, int, int, int] | None:
        return self.merge_map().get((row, col))

    def cell(self, row: int, col: int) -> Cell:
        f = self.formula(row, col)
        cached = self.get(row, col)
        return Cell(value=f if f is not None else cached, cached=cached, is_formula=f is not None,
                    number_format=self.number_format(row, col), bold=self.is_bold(row, col),
                    merged=self.merged_range(row, col))

    def iter_rows(self) -> Iterator[tuple[int, list[Any]]]:
        for i, r in enumerate(self.rows, start=1):
            yield i, r


@dataclass
class DocumentData:
    path: str                     # path that was read (work path for xls)
    kind: str                     # xlsx | docx | pdf
    sheets: list[SheetData] = field(default_factory=list)
    paragraphs: list[str] = field(default_factory=list)
    original_path: str | None = None
    work_path: str | None = None  # xls -> xlsx conversion result
    pages: int = 0

    def sheet(self, name: str) -> SheetData | None:
        for s in self.sheets:
            if s.name == name:
                return s
        return None


# ------------------------------------------------------------------ dispatch

def read_any(path: str, *, ext: str | None = None, work_dir: str | None = None) -> DocumentData:
    ext = (ext or os.path.splitext(path)[1].lstrip(".")).lower()
    if not os.path.exists(path):
        raise IngestError("The file is missing from storage.")
    if os.path.getsize(path) == 0:
        raise IngestError("The file is empty.")
    if ext in ("xlsx", "xlsm"):
        return read_xlsx(path)
    if ext == "xls":
        work = convert_to_xlsx(path, os.path.join(work_dir or os.path.dirname(path), "work.xlsx"))
        doc = read_xlsx(work)
        doc.original_path = path
        doc.work_path = work
        return doc
    if ext == "docx":
        return read_docx(path)
    if ext == "pdf":
        return read_pdf(path)
    raise IngestError(f"Unsupported file type: .{ext}")


# ------------------------------------------------------------------ xlsx

def _scan_sheet_xml(archive, member: str) -> tuple[list[tuple[int, int, int, int]], set[int], set[int]]:
    """Stream the sheet XML for <mergeCell>, hidden <row> and hidden <col> elements."""
    merged: list[tuple[int, int, int, int]] = []
    hidden_rows: set[int] = set()
    hidden_cols: set[int] = set()
    re_row = re.compile(rb"<(?:\w+:)?row\b([^>]*)>")
    re_col = re.compile(rb"<(?:\w+:)?col\b([^>]*)/?>")
    re_merge = re.compile(rb"<(?:\w+:)?mergeCell\b[^>]*\bref=\"([A-Z]+)(\d+):([A-Z]+)(\d+)\"")
    re_attr_r = re.compile(rb"\br=\"(\d+)\"")
    re_hidden = re.compile(rb"\bhidden=\"(?:1|true)\"")
    re_min = re.compile(rb"\bmin=\"(\d+)\"")
    re_max = re.compile(rb"\bmax=\"(\d+)\"")
    with archive.open(member) as fh:
        tail = b""
        while True:
            chunk = fh.read(1 << 20)
            if not chunk:
                break
            buf = tail + chunk
            cut = buf.rfind(b"<")
            if cut <= 0:
                cut = len(buf)
            data, tail = buf[:cut], buf[cut:]
            if b"hidden" in data:
                for m in re_row.finditer(data):
                    attrs = m.group(1)
                    if re_hidden.search(attrs):
                        rm = re_attr_r.search(attrs)
                        if rm:
                            hidden_rows.add(int(rm.group(1)))
                for m in re_col.finditer(data):
                    attrs = m.group(1)
                    if re_hidden.search(attrs):
                        mn, mx = re_min.search(attrs), re_max.search(attrs)
                        if mn and mx:
                            lo, hi = int(mn.group(1)), min(int(mx.group(1)), MAX_COLS)
                            hidden_cols.update(range(lo, hi + 1))
            if b"mergeCell" in data:
                for m in re_merge.finditer(data):
                    c1 = column_index_from_string(m.group(1).decode())
                    c2 = column_index_from_string(m.group(3).decode())
                    merged.append((int(m.group(2)), c1, int(m.group(4)), c2))
        if tail:
            for m in re_merge.finditer(tail):
                c1 = column_index_from_string(m.group(1).decode())
                c2 = column_index_from_string(m.group(3).decode())
                merged.append((int(m.group(2)), c1, int(m.group(4)), c2))
    return merged, hidden_rows, hidden_cols


def _clean_value(v: Any) -> Any:
    if isinstance(v, str):
        v = v.replace(" ", " ").strip()
        return v or None
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat()
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return v if v == v else None
    if v is None:
        return None
    return str(v)


def read_xlsx(path: str) -> DocumentData:
    import zipfile

    from openpyxl import load_workbook
    try:
        wb_v = load_workbook(path, read_only=True, data_only=True)
    except zipfile.BadZipFile:
        raise IngestError("The spreadsheet is damaged or password-protected (not a valid .xlsx file).") from None
    except Exception as e:  # noqa: BLE001
        msg = str(e).lower()
        if "encrypt" in msg or "password" in msg:
            raise IngestError("The spreadsheet is password-protected. Remove the password and upload again.") from e
        raise IngestError(f"The spreadsheet could not be opened ({type(e).__name__}).") from e
    doc = DocumentData(path=path, kind="xlsx")
    try:
        for idx, ws in enumerate(wb_v.worksheets):
            sd = SheetData(name=ws.title, idx=idx, state=getattr(ws, "sheet_state", "visible") or "visible")
            _read_values(wb_v, ws, sd)
            try:
                sd.merged, sd.hidden_rows, sd.hidden_cols = _scan_sheet_xml(wb_v._archive, ws._worksheet_path)
            except Exception:  # noqa: BLE001 - merges are a nice-to-have
                log.warning("merge scan failed for %s", ws.title, exc_info=True)
            doc.sheets.append(sd)
    finally:
        wb_v.close()

    wb_f = load_workbook(path, read_only=True, data_only=False)
    try:
        for ws, sd in zip(wb_f.worksheets, doc.sheets):
            ws.reset_dimensions()
            for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if r_i > MAX_ROWS:
                    break
                for c_i, v in enumerate(row, start=1):
                    if c_i > MAX_COLS:
                        break
                    if v is None:
                        continue
                    if isinstance(v, str):
                        if v.startswith("="):
                            sd.formulas[(r_i, c_i)] = v
                    elif hasattr(v, "text") and isinstance(getattr(v, "text", None), str):   # ArrayFormula
                        t = v.text
                        sd.formulas[(r_i, c_i)] = t if t.startswith("=") else "=" + t
    finally:
        wb_f.close()
    _fill_uncached(doc)
    return doc


def _read_values(wb, ws, sd: SheetData) -> None:
    ws.reset_dimensions()
    style_cache: dict[int, tuple[bool, str]] = {}
    cell_styles = wb._cell_styles
    fonts = wb._fonts
    from openpyxl.styles.numbers import BUILTIN_FORMATS, BUILTIN_FORMATS_MAX_SIZE
    numfmts = wb._number_formats
    rows = sd.rows
    last_nonempty = 0
    for r_i, row in enumerate(ws.iter_rows(), start=1):
        if r_i > MAX_ROWS:
            break
        vals: list[Any] = []
        last = 0
        for c_i, cell in enumerate(row, start=1):
            if c_i > MAX_COLS:
                break
            v = getattr(cell, "value", None)
            v = _clean_value(v)
            vals.append(v)
            if v is None:
                continue
            last = c_i
            sid = getattr(cell, "_style_id", 0) or 0
            st = style_cache.get(sid)
            if st is None:
                bold, nf = False, "General"
                try:
                    sa = cell_styles[sid]
                    font = fonts[sa.fontId]
                    bold = bool(font.b)
                    fid = sa.numFmtId
                    nf = BUILTIN_FORMATS.get(fid, "General") if fid < BUILTIN_FORMATS_MAX_SIZE \
                        else numfmts[fid - BUILTIN_FORMATS_MAX_SIZE]
                except Exception:  # noqa: BLE001
                    pass
                st = (bold, nf)
                style_cache[sid] = st
            if st[0]:
                sd.bold.add((r_i, c_i))
            if st[1] != "General":
                sd.number_formats[(r_i, c_i)] = st[1]
        del vals[last:]
        rows.append(vals)
        if last:
            last_nonempty = r_i
    del rows[last_nonempty:]


def _fill_uncached(doc: DocumentData) -> None:
    """Compute formula cells whose cached value is missing."""
    by_name = {s.name: s for s in doc.sheets}
    for sd in doc.sheets:
        missing = [(rc, f) for rc, f in sd.formulas.items() if sd.get(*rc) is None]
        if not missing:
            continue
        if len(missing) > 100_000:
            continue
        busy: set[tuple[str, int, int]] = set()

        def getter(sheet: str | None, col: int, row: int, _home=sd) -> Any:
            target = by_name.get(sheet, _home) if sheet else _home
            v = target.get(row, col)
            if v is None and (row, col) in target.formulas:
                key = (target.name, row, col)
                if key in busy or len(busy) > 400:
                    return None
                busy.add(key)
                try:
                    v = fx.evaluate(target.formulas[(row, col)], getter)
                finally:
                    busy.discard(key)
                if v is not None:
                    _set(target, row, col, v)
                    target.computed.add((row, col))
            return v

        for (r, c), f in missing:
            if sd.get(r, c) is None:
                v = fx.evaluate(f, getter)
                if v is not None:
                    _set(sd, r, c, v)
                    sd.computed.add((r, c))


def _set(sd: SheetData, row: int, col: int, v: Any) -> None:
    while len(sd.rows) < row:
        sd.rows.append([])
    r = sd.rows[row - 1]
    if len(r) < col:
        r.extend([None] * (col - len(r)))
    r[col - 1] = v


# ------------------------------------------------------------------ xls -> xlsx

def convert_to_xlsx(src: str, dest: str, *, timeout: int = 120) -> str:
    """Convert .xls (or anything LibreOffice reads) to .xlsx at `dest`. The original is left untouched.
    Each call gets its own LibreOffice profile dir so concurrent conversions don't clash."""
    if not os.path.exists(SOFFICE):
        raise IngestError("Old .xls files need LibreOffice for conversion, which is not installed on the server.")
    tmp_out = tempfile.mkdtemp(prefix="xls2xlsx_")
    profile = os.path.join(tempfile.gettempdir(), f"lo_{uuid.uuid4().hex}")
    try:
        cmd = [SOFFICE, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--nologo",
               "--convert-to", "xlsx", "--outdir", tmp_out, src]
        try:
            proc = subprocess.run(cmd, capture_output=True, timeout=timeout, check=False,
                                  env={**os.environ, "HOME": profile})
        except subprocess.TimeoutExpired:
            raise IngestError("Converting the .xls file timed out (over 2 minutes).") from None
        produced = os.path.join(tmp_out, os.path.splitext(os.path.basename(src))[0] + ".xlsx")
        if not os.path.exists(produced):
            files = [f for f in os.listdir(tmp_out) if f.endswith(".xlsx")]
            if not files:
                err = (proc.stderr or b"").decode(errors="replace")[-300:]
                raise IngestError(f"The .xls file could not be converted (LibreOffice: {err.strip() or 'no output'}).")
            produced = os.path.join(tmp_out, files[0])
        os.makedirs(os.path.dirname(os.path.abspath(dest)), exist_ok=True)
        shutil.move(produced, dest)
        return dest
    finally:
        shutil.rmtree(tmp_out, ignore_errors=True)
        shutil.rmtree(profile, ignore_errors=True)


# ------------------------------------------------------------------ docx

def read_docx(path: str) -> DocumentData:
    try:
        import docx  # python-docx
        d = docx.Document(path)
    except Exception as e:  # noqa: BLE001
        raise IngestError(f"The Word document could not be opened ({type(e).__name__}).") from e
    doc = DocumentData(path=path, kind="docx")
    doc.paragraphs = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
    for t_i, table in enumerate(d.tables):
        sd = SheetData(name=f"Table {t_i + 1}", idx=t_i, source="docx-table")
        for r_i, row in enumerate(table.rows, start=1):
            vals: list[Any] = []
            prev_tc = None
            start_c = None
            for c_i, cell in enumerate(row.cells, start=1):
                tc = cell._tc
                if tc is prev_tc:
                    vals.append(None)
                    continue
                if prev_tc is not None and start_c is not None and c_i - 1 > start_c:
                    sd.merged.append((r_i, start_c, r_i, c_i - 1))
                prev_tc, start_c = tc, c_i
                text = cell.text.strip()
                vals.append(text or None)
                if text and any(run.bold for p in cell.paragraphs for run in p.runs if run.text.strip()):
                    sd.bold.add((r_i, c_i))
            if start_c is not None and len(vals) > start_c:
                sd.merged.append((r_i, start_c, r_i, len(vals)))
            while vals and vals[-1] is None:
                vals.pop()
            sd.rows.append(vals)
        while sd.rows and not sd.rows[-1]:
            sd.rows.pop()
        if sd.rows:
            doc.sheets.append(sd)
    for i, s in enumerate(doc.sheets):
        s.idx = i
    return doc


# ------------------------------------------------------------------ pdf

def _is_password_error(e: BaseException) -> bool:
    seen = 0
    cur: BaseException | None = e
    while cur is not None and seen < 6:
        name = type(cur).__name__.lower()
        text = str(cur).lower()
        if "password" in name or "encrypt" in name or "password" in text or "encrypt" in text:
            return True
        for a in getattr(cur, "args", ()):
            if isinstance(a, BaseException) and _is_password_error(a):
                return True
        cur = cur.__cause__ or cur.__context__
        seen += 1
    return False


def read_pdf(path: str, *, max_pages: int = 300) -> DocumentData:
    import pdfplumber
    try:
        pdf = pdfplumber.open(path)
    except Exception as e:  # noqa: BLE001
        if _is_password_error(e):
            raise IngestError("PDF is encrypted (password-protected). Remove the password and upload again.") from e
        raise IngestError(f"The PDF could not be opened ({type(e).__name__}).") from e
    doc = DocumentData(path=path, kind="pdf")
    try:
        try:
            pages = pdf.pages
        except Exception as e:  # noqa: BLE001
            if _is_password_error(e):
                raise IngestError("PDF is encrypted (password-protected). Remove the password and upload again.") \
                    from e
            raise
        doc.pages = len(pages)
        total_chars = 0
        t_idx = 0
        for p_i, page in enumerate(pages[:max_pages], start=1):
            try:
                text = page.extract_text() or ""
            except Exception as e:  # noqa: BLE001
                if _is_password_error(e):
                    raise IngestError("PDF is encrypted (password-protected). Remove the password and upload again.") \
                        from e
                text = ""
            total_chars += len(text.strip())
            doc.paragraphs.extend(line.strip() for line in text.splitlines() if line.strip())
            try:
                tables = page.extract_tables() or []
                if not tables and text.strip():
                    tables = page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"}) or []
                    # a text-strategy "table" with one column is just prose
                    tables = [t for t in tables if t and max(len(r) for r in t) >= 3 and len(t) >= 3]
            except Exception:  # noqa: BLE001
                tables = []
            for tbl in tables:
                t_idx += 1
                sd = SheetData(name=f"Page {p_i} table {t_idx}", idx=len(doc.sheets), source="pdf-table")
                for row in tbl:
                    vals = [(str(c).replace("\n", " ").strip() or None) if c is not None else None for c in row]
                    while vals and vals[-1] is None:
                        vals.pop()
                    sd.rows.append(vals)
                if any(sd.rows):
                    doc.sheets.append(sd)
            page.flush_cache() if hasattr(page, "flush_cache") else None
        n = max(1, min(len(pages), max_pages))
        if total_chars < 25 * n and not doc.sheets:
            raise IngestError("Scanned PDF: no text layer found. Upload the original spreadsheet or a text-based PDF.")
    finally:
        pdf.close()
    return doc


def col_letter(c: int) -> str:
    return get_column_letter(c)
